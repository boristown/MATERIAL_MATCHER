#!/usr/bin/env python3
"""生成随安装介质交付的最小业务 smoke 数据（两个 .xlsx），完全不依赖网络。"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook

SOURCE_ROWS = [
    ("M9001", "三相异步电动机", "YE3-132S-4 5.5kW", "西门子", "台"),
    ("M9002", "深沟球轴承", "6205-2RS", "SKF", "套"),
    ("M9003", "闸阀", "Z45H-10Q DN100", "津沃", "台"),
    ("M9004", "控制电缆", "KVVP 4x1.5", "远东", "米"),
    ("M9005", "安全帽", "V型 ABS 红色", "汉威", "顶"),
    ("M9006", "液压油", "L-HM46 抗磨", "长城", "桶"),
    ("M9007", "螺栓", "M12x60 8.8级 镀锌", "晋亿", "套"),
    ("M9008", "工业风机", "4-72 No.4A", "九洲", "台"),
]

TARGET_ROWS = [
    ("JT00000000901", "三相异步电动机", "YE3-132S-4 5.5KW 380V", "西门子", "台"),
    ("JT00000000902", "深沟球轴承", "6205-2RS1", "SKF", "个"),
    ("JT00000000903", "明杆弹性座封闸阀", "Z45H-10Q DN100", "津沃", "台"),
    ("JT00000000904", "铜芯聚氯乙烯绝缘屏蔽控制电缆", "KVVP 450/750V 4×1.5", "远东", "米"),
    ("JT00000000905", "安全帽", "V型 ABS 红色", "汉威", "顶"),
    ("JT00000000906", "抗磨液压油", "L-HM46", "长城", "桶"),
    ("JT00000000907", "六角头螺栓", "M12×60 8.8 级镀锌", "晋亿", "套"),
    ("JT00000000908", "离心通风机", "4-72 No.4A", "九洲", "台"),
]


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[tuple]) -> None:
    sheet = workbook.active if title == workbook.active.title else workbook.create_sheet()
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))


def generate(output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    source_path = output_dir / "smoke-待匹配数据.xlsx"
    workbook = Workbook()
    _sheet(workbook, "待匹配物料", ["物料编码", "物料名称", "型号", "品牌", "计量单位"], SOURCE_ROWS)
    workbook.save(source_path)

    target_path = output_dir / "smoke-集团标准数据.xlsx"
    workbook = Workbook()
    _sheet(workbook, "集团标准物料", ["集团码", "物料名称", "型号", "生产厂家", "计量单位"], TARGET_ROWS)
    workbook.save(target_path)
    return [source_path, target_path]


def main() -> int:
    parser = argparse.ArgumentParser(description="生成安装介质最小 smoke Excel 数据")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    for path in generate(args.output_dir):
        print(path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
