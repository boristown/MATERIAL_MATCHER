from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

SEED = 20260917
DEFAULT_SOURCE_ROWS = 1000
DEFAULT_MATCHED_ROWS = 900
SUPPORTED_TYPES = ("Z001", "Z006")
TYPE_ALIASES = {"A001": "Z001", "Z001": "Z001", "A006": "Z006", "Z006": "Z006"}
TYPE_KEYWORDS = {"元器件": "Z001", "复合材料": "Z006"}
PAIR_THRESHOLD = 0.62

SOURCE_HEADERS = [
    "物料编码", "物料类型", "一层分类描述", "二层分类描述", "三层分类描述",
    "四层分类描述", "五层分类描述", "物料名称", "型号", "型号规格", "质量等级",
    "国产/进口", "总规范", "封装形式", "详细规范", "外形尺寸", "生产厂家", "特殊说明",
    "采购标准", "规格", "技术标准", "标准号", "型号(牌号)", "计量单位",
]
TARGET_HEADERS = [
    "集团码", "物料类型", "二层分类", "三层分类", "四层分类", "五层分类", "物料名称",
    "型号", "型号规格", "质量等级", "国产/进口", "封装形式", "生产厂家", "采用标准",
    "通用规范", "详细规范", "外形尺寸", "牌号", "规格", "计量单位", "特殊说明",
]
TRUTH_HEADERS = [
    "物料编码", "预期是否可匹配", "预期集团码", "物料类型", "场景",
    "种子源文件", "种子源行号", "种子目标文件", "种子目标行号", "种子配对分", "说明",
]

SOURCE_ALIASES: dict[str, list[str]] = {
    "物料编码": ["物料编码", "物料号", "物料", "MATNR"],
    "物料类型": ["物料类型", "一层分类", "ZWLYX"],
    "一层分类描述": ["一层分类描述", "ZWLYXT"],
    "二层分类描述": ["二层分类描述", "ZECFLT"],
    "三层分类描述": ["三层分类描述", "ZSCFLT"],
    "四层分类描述": ["四层分类描述", "ZSCFL2T"],
    "五层分类描述": ["五层分类描述", "ZWCFLT"],
    "物料名称": ["物料名称", "名称", "ZWLMC"],
    "型号": ["型号", "ZXH"],
    "型号规格": ["型号规格", "ZXHGG"],
    "质量等级": ["质量等级", "ZZLDJ"],
    "国产/进口": ["国产/进口", "国产_进口", "ZGCJK"],
    "总规范": ["总规范", "ZZGF"],
    "封装形式": ["封装形式", "ZFZXS"],
    "详细规范": ["详细规范", "ZXXGF"],
    "外形尺寸": ["外形尺寸", "ZWGCC"],
    "生产厂家": ["生产厂家", "厂家", "ZSCCJ"],
    "特殊说明": ["特殊说明", "ZTSSM"],
    "采购标准": ["采购标准", "ZCGBZ"],
    "规格": ["规格", "ZGG"],
    "技术标准": ["技术标准", "ZJSBZ"],
    "标准号": ["标准号", "ZBZH"],
    "型号(牌号)": ["型号(牌号)", "型号（牌号）", "ZXHPH"],
    "计量单位": ["计量单位", "单位", "MEINS"],
}

GROUP_CODE_ALIASES = ["集团编码", "集团物资编码", "集团物料编码", "集团码", "物资编码"]

TARGET_ALIASES: dict[str, dict[str, list[str]]] = {
    "Z001": {
        "集团码": GROUP_CODE_ALIASES,
        "二层分类": ["清洗后大类", "大类", "二层分类"],
        "三层分类": ["清洗后中类", "中类", "三层分类"],
        "四层分类": ["清洗后小类", "小类", "四层分类"],
        "五层分类": ["清洗后细类", "细类", "五层分类"],
        "物料名称": ["清洗后产品名称", "产品名称", "名称", "物料名称"],
        "型号": ["清洗后型号或系列", "型号或系列", "型号"],
        "型号规格": ["清洗后详细型号规格", "详细型号规格", "型号规格"],
        "质量等级": ["清洗后质量等级", "质量等级"],
        "国产/进口": ["国产_进口", "国产/进口"],
        "封装形式": ["清洗后封装形式", "封装形式"],
        "生产厂家": ["清洗后生产厂家", "生产厂家", "厂家"],
        "采用标准": ["清洗后详细规范号", "详细规范号", "采用标准"],
        "通用规范": ["清洗后通用规范", "通用规范"],
        "详细规范": ["清洗后详细规范号", "详细规范号", "详细规范"],
        "外形尺寸": ["清洗后外形尺寸", "外形尺寸"],
        "牌号": ["清洗后型号或系列", "型号或系列", "牌号"],
        "规格": ["清洗后详细型号规格", "详细型号规格", "规格"],
        "计量单位": ["计量单位", "单位"],
        "特殊说明": ["特殊说明"],
    },
    "Z006": {
        "集团码": GROUP_CODE_ALIASES,
        "二层分类": ["分类Ⅰ级", "分类I级", "分类1级", "二层分类"],
        "三层分类": ["分类Ⅱ级", "分类II级", "分类2级", "三层分类"],
        "四层分类": ["分类Ⅲ级", "分类III级", "分类3级", "四层分类"],
        "五层分类": ["分类Ⅳ级", "分类IV级", "分类4级", "五层分类"],
        "物料名称": ["名称", "物料名称", "产品名称"],
        "型号": ["牌号", "型号"],
        "型号规格": ["规格", "型号规格"],
        "质量等级": ["质量等级"],
        "国产/进口": ["国产_进口", "国产/进口"],
        "封装形式": ["封装形式"],
        "生产厂家": ["生产厂家", "厂家"],
        "采用标准": ["采用标准", "标准"],
        "通用规范": ["采用标准", "通用规范"],
        "详细规范": ["采用标准", "详细规范"],
        "外形尺寸": ["规格", "外形尺寸"],
        "牌号": ["牌号", "型号"],
        "规格": ["规格", "型号规格"],
        "计量单位": ["计量单位", "单位"],
        "特殊说明": ["特殊说明"],
    },
}

PAIR_RULES: dict[str, list[tuple[str, str, float]]] = {
    "Z001": [
        ("物料名称", "物料名称", 0.10), ("型号", "型号", 0.28), ("型号规格", "型号规格", 0.22),
        ("封装形式", "封装形式", 0.10), ("生产厂家", "生产厂家", 0.10),
        ("质量等级", "质量等级", 0.06), ("二层分类描述", "二层分类", 0.04),
        ("三层分类描述", "三层分类", 0.03), ("四层分类描述", "四层分类", 0.03),
        ("五层分类描述", "五层分类", 0.02), ("总规范", "通用规范", 0.02),
    ],
    "Z006": [
        ("物料名称", "物料名称", 0.13), ("型号(牌号)", "牌号", 0.23), ("型号", "牌号", 0.08),
        ("型号规格", "规格", 0.24), ("规格", "规格", 0.08), ("采购标准", "采用标准", 0.12),
        ("二层分类描述", "二层分类", 0.04), ("三层分类描述", "三层分类", 0.03),
        ("四层分类描述", "四层分类", 0.025), ("五层分类描述", "五层分类", 0.025),
    ],
}


@dataclass(frozen=True)
class SeedRow:
    kind: str
    material_type: str
    file_name: str
    sheet_name: str
    row_number: int
    payload: dict[str, str]


@dataclass(frozen=True)
class SeedPair:
    source: SeedRow
    target: SeedRow
    score: float
    gap: float


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).upper()
    return re.sub(r"[\s_\-—–/\\()（）\[\]【】.:：]+", "", text)


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _text(value)).upper()
    if text in {"", "88", "<NULL>", "-", "/"}:
        return ""
    return re.sub(r"[\s_\-—–/\\()（）\[\]【】,，.;；:：]+", "", text)


def _canonical_type(value: str, file_name: str = "") -> str:
    compact = _norm(value)
    for alias, canonical in TYPE_ALIASES.items():
        if alias in compact:
            return canonical
    upper_name = file_name.upper()
    for alias, canonical in TYPE_ALIASES.items():
        if alias in upper_name:
            return canonical
    for keyword, canonical in TYPE_KEYWORDS.items():
        if keyword in file_name:
            return canonical
    return ""


def _lookup(payload: dict[str, str], aliases: Iterable[str]) -> str:
    normalized = {_header_key(key): _text(value) for key, value in payload.items()}
    for alias in aliases:
        value = normalized.get(_header_key(alias), "")
        if value != "":
            return value
    return ""


def _target_type_from_headers(headers: Iterable[str]) -> str:
    keys = {_header_key(item) for item in headers if item}
    z001_markers = [
        "清洗后大类", "清洗后中类", "清洗后小类", "清洗后细类", "清洗后产品名称",
        "清洗后型号或系列", "清洗后详细型号规格", "清洗后封装形式", "清洗后生产厂家",
    ]
    z006_markers = ["分类Ⅰ级", "分类I级", "分类1级", "分类Ⅱ级", "分类2级", "分类Ⅲ级", "分类3级", "分类Ⅳ级", "分类4级", "牌号", "采用标准"]
    s1 = sum(_header_key(item) in keys for item in z001_markers)
    s6 = sum(_header_key(item) in keys for item in z006_markers)
    if s1 >= 2 and s1 > s6:
        return "Z001"
    if s6 >= 2 and s6 > s1:
        return "Z006"
    return ""


def _find_header_row(sheet) -> tuple[int, list[str]] | None:
    best: tuple[int, int, list[str]] | None = None
    best_row_no = 0
    for row_no, cells in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True), start=1):
        headers = [_text(cell) for cell in cells]
        non_empty = [item for item in headers if item]
        if len(non_empty) < 3:
            continue
        keys = {_header_key(item) for item in non_empty}
        score = 0
        if any(_header_key(alias) in keys for aliases in SOURCE_ALIASES.values() for alias in aliases):
            score += 2
        if any(_header_key(alias) in keys for alias in GROUP_CODE_ALIASES):
            score += 5
        if _target_type_from_headers(non_empty):
            score += 3
        if any(item.upper() == "MATNR" or item in {"物料", "物料号", "物料编码"} for item in non_empty):
            score += 4
        candidate = (score, len(non_empty), headers)
        if best is None or candidate[:2] > best[:2]:
            best = candidate
            best_row_no = row_no
    if best is None or best[0] < 2:
        return None
    return best_row_no, best[2]


def _looks_like_jt_path(relative_name: str) -> bool:
    normalized = relative_name.replace("\\", "/").upper()
    return "/JT/" in f"/{normalized}/" or normalized.startswith("JT/")


def _rows_from_workbook(path: Path, relative_name: str, diagnostics: list[dict[str, Any]]) -> list[SeedRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    output: list[SeedRow] = []
    try:
        for sheet in workbook.worksheets:
            header_info = _find_header_row(sheet)
            if not header_info:
                diagnostics.append({"file": relative_name, "sheet": sheet.title, "kind": "ignored", "reason": "header_not_found"})
                continue
            header_row, headers = header_info
            header_keys = {_header_key(item) for item in headers if item}
            has_group_code = any(_header_key(alias) in header_keys for alias in GROUP_CODE_ALIASES)
            target_type = _target_type_from_headers(headers)
            is_target = has_group_code or (_looks_like_jt_path(relative_name) and bool(target_type))
            kind = "target" if is_target else "source"
            file_type = _canonical_type("", relative_name + "/" + sheet.title)
            sheet_type = target_type if kind == "target" else file_type
            accepted = 0
            for row_no, cells in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                payload = {headers[idx]: _text(value) for idx, value in enumerate(cells) if idx < len(headers) and headers[idx]}
                if not any(payload.values()):
                    continue
                if kind == "source":
                    raw_type = _lookup(payload, SOURCE_ALIASES["物料类型"])
                    material_type = _canonical_type(raw_type, relative_name) or file_type
                    source_id = _lookup(payload, SOURCE_ALIASES["物料编码"])
                    if not source_id or material_type not in SUPPORTED_TYPES:
                        continue
                else:
                    material_type = sheet_type or file_type
                    if material_type not in SUPPORTED_TYPES:
                        continue
                    code = _lookup(payload, TARGET_ALIASES[material_type]["集团码"])
                    if not code:
                        continue
                output.append(SeedRow(kind, material_type, relative_name, sheet.title, row_no, payload))
                accepted += 1
            diagnostics.append({
                "file": relative_name,
                "sheet": sheet.title,
                "kind": kind,
                "detected_type": sheet_type or file_type,
                "header_row": header_row,
                "headers": [item for item in headers if item][:30],
                "accepted_rows": accepted,
            })
    finally:
        workbook.close()
    return output


def _rows_from_csv(path: Path, relative_name: str, diagnostics: list[dict[str, Any]]) -> list[SeedRow]:
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            pass
    if text is None or not text.strip():
        return []
    lines = text.splitlines()
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:5]))
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(lines, dialect=dialect)
    headers = reader.fieldnames or []
    header_keys = {_header_key(item) for item in headers if item}
    target_type = _target_type_from_headers(headers)
    has_group_code = any(_header_key(alias) in header_keys for alias in GROUP_CODE_ALIASES)
    kind = "target" if has_group_code or (_looks_like_jt_path(relative_name) and bool(target_type)) else "source"
    file_type = _canonical_type("", relative_name)
    sheet_type = target_type if kind == "target" else file_type
    output: list[SeedRow] = []
    for row_no, raw in enumerate(reader, start=2):
        payload = {str(k): _text(v) for k, v in raw.items() if k}
        if kind == "source":
            material_type = _canonical_type(_lookup(payload, SOURCE_ALIASES["物料类型"]), relative_name) or file_type
            if not _lookup(payload, SOURCE_ALIASES["物料编码"]) or material_type not in SUPPORTED_TYPES:
                continue
        else:
            material_type = sheet_type or file_type
            if material_type not in SUPPORTED_TYPES or not _lookup(payload, TARGET_ALIASES[material_type]["集团码"]):
                continue
        output.append(SeedRow(kind, material_type, relative_name, "CSV", row_no, payload))
    diagnostics.append({"file": relative_name, "sheet": "CSV", "kind": kind, "detected_type": sheet_type or file_type, "headers": headers[:30], "accepted_rows": len(output)})
    return output


def discover_seed_rows(seed_zip: Path) -> tuple[list[SeedRow], list[SeedRow], dict[str, Any]]:
    if not seed_zip.exists():
        raise FileNotFoundError(f"seed zip not found: {seed_zip}")
    sources: list[SeedRow] = []
    targets: list[SeedRow] = []
    scanned_files: list[str] = []
    diagnostics: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="material-seed-") as temp_dir:
        root = Path(temp_dir)
        with zipfile.ZipFile(seed_zip) as archive:
            archive.extractall(root)
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.name.startswith("~$"):
                continue
            relative_name = path.relative_to(root).as_posix()
            suffix = path.suffix.lower()
            if suffix not in {".xlsx", ".xlsm", ".csv"}:
                continue
            scanned_files.append(relative_name)
            rows = _rows_from_csv(path, relative_name, diagnostics) if suffix == ".csv" else _rows_from_workbook(path, relative_name, diagnostics)
            sources.extend(row for row in rows if row.kind == "source")
            targets.extend(row for row in rows if row.kind == "target")
    inventory = {
        "seed_zip": seed_zip.name,
        "seed_zip_sha256": hashlib.sha256(seed_zip.read_bytes()).hexdigest(),
        "seed_files_scanned": scanned_files,
        "source_seed_rows": len(sources),
        "target_seed_rows": len(targets),
        "source_seed_by_type": {t: sum(row.material_type == t for row in sources) for t in SUPPORTED_TYPES},
        "target_seed_by_type": {t: sum(row.material_type == t for row in targets) for t in SUPPORTED_TYPES},
        "file_diagnostics": diagnostics,
    }
    return sources, targets, inventory


def canonical_source(seed: SeedRow) -> dict[str, str]:
    row = {header: "" for header in SOURCE_HEADERS}
    for header in SOURCE_HEADERS:
        row[header] = _lookup(seed.payload, SOURCE_ALIASES.get(header, [header]))
    row["物料类型"] = seed.material_type
    if not row["一层分类描述"]:
        row["一层分类描述"] = "元器件" if seed.material_type == "Z001" else "复合材料"
    return row


def canonical_target(seed: SeedRow) -> dict[str, str]:
    row = {header: "" for header in TARGET_HEADERS}
    aliases = TARGET_ALIASES[seed.material_type]
    for header in TARGET_HEADERS:
        row[header] = seed.material_type if header == "物料类型" else _lookup(seed.payload, aliases.get(header, [header]))
    return row


def _similarity(left: str, right: str) -> float:
    a, b = _norm(left), _norm(right)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.86
    return SequenceMatcher(None, a, b).ratio()


def pair_score(source: dict[str, str], target: dict[str, str], material_type: str) -> float:
    total = 0.0
    used = 0.0
    for source_field, target_field, weight in PAIR_RULES[material_type]:
        left, right = source.get(source_field, ""), target.get(target_field, "")
        if not _norm(left) or not _norm(right):
            continue
        used += weight
        total += weight * _similarity(left, right)
    return total / used if used else 0.0


def build_seed_pairs(sources: list[SeedRow], targets: list[SeedRow]) -> list[SeedPair]:
    target_by_type = {t: [row for row in targets if row.material_type == t] for t in SUPPORTED_TYPES}
    target_canonical = {id(seed): canonical_target(seed) for seed in targets}
    pairs: list[SeedPair] = []
    for source_seed in sources:
        source = canonical_source(source_seed)
        scored = sorted(
            ((pair_score(source, target_canonical[id(target)], source_seed.material_type), target) for target in target_by_type[source_seed.material_type]),
            key=lambda item: item[0], reverse=True,
        )
        if not scored:
            continue
        best_score, best_target = scored[0]
        second = scored[1][0] if len(scored) > 1 else 0.0
        gap = best_score - second
        if best_score >= PAIR_THRESHOLD and (gap >= 0.02 or best_score >= 0.86):
            pairs.append(SeedPair(source_seed, best_target, best_score, gap))
    return pairs


def _stable_material_code(index: int) -> str:
    return f"{900000000000000000 + index:018d}"


def _format_noise(row: dict[str, str], scenario: str, rng: random.Random) -> None:
    if scenario == "exact_seed":
        return
    if scenario in {"space_punctuation", "combined"}:
        for field in ("型号", "型号规格", "规格", "外形尺寸"):
            value = row.get(field, "")
            if value:
                value = value.replace("×", " x ").replace("Φ", "φ")
                if field == "型号" and len(value) > 3:
                    value = f" {value} "
                row[field] = value
    if scenario in {"standard_format", "combined"}:
        for field in ("总规范", "详细规范", "采购标准", "技术标准", "标准号"):
            row[field] = row.get(field, "").replace(" ", "")
    if scenario in {"missing_secondary", "combined"}:
        candidates = [field for field in ("质量等级", "详细规范", "总规范", "特殊说明", "外形尺寸") if row.get(field)]
        for field in rng.sample(candidates, k=min(2, len(candidates))):
            row[field] = "88"


def _make_unmatched(row: dict[str, str], material_type: str, index: int) -> None:
    marker = f"NM{index:04d}"
    if material_type == "Z001":
        if row.get("型号"):
            row["型号"] = f"{row['型号']}-{marker}"
        else:
            row["型号"] = marker
        if row.get("型号规格"):
            row["型号规格"] = f"{row['型号规格']} / 目录外变体{index}"
        if row.get("封装形式"):
            row["封装形式"] = f"{row['封装形式']}-X{index % 7 + 1}"
    else:
        base = row.get("型号(牌号)") or row.get("型号") or "SEED"
        row["型号(牌号)"] = f"{base}-{marker}"
        row["型号"] = row["型号(牌号)"]
        spec = row.get("型号规格") or row.get("规格") or ""
        row["型号规格"] = f"{spec} / 目录外规格{index}".strip()
        row["规格"] = row["型号规格"]
    row["特殊说明"] = "基于真实种子生成的目录外测试变体；预期人工判定均不匹配"


def _type_plan(total: int, z001_ratio: float = 0.70) -> dict[str, int]:
    z001 = round(total * z001_ratio)
    return {"Z001": z001, "Z006": total - z001}


def generate_rows(seed_zip: Path, *, source_rows: int = DEFAULT_SOURCE_ROWS, matched_rows: int = DEFAULT_MATCHED_ROWS, seed: int = SEED) -> tuple[list[dict[str, str]], list[dict[str, str]], list[dict[str, str]], dict[str, Any]]:
    if not (0 < matched_rows <= source_rows):
        raise ValueError("matched_rows must be within 1..source_rows")
    rng = random.Random(seed)
    source_seeds, target_seeds, inventory = discover_seed_rows(seed_zip)
    pairs = build_seed_pairs(source_seeds, target_seeds)
    by_type = {t: [pair for pair in pairs if pair.source.material_type == t] for t in SUPPORTED_TYPES}
    missing = [t for t in SUPPORTED_TYPES if not by_type[t]]
    if missing:
        summary = {t: {"sources": inventory["source_seed_by_type"][t], "targets": inventory["target_seed_by_type"][t]} for t in SUPPORTED_TYPES}
        raise RuntimeError(f"origin_data.zip lacks high-confidence source/target seed pairs for: {', '.join(missing)}; inventory={summary}; diagnostics={inventory['file_diagnostics']}")

    matched_plan = _type_plan(matched_rows)
    unmatched_plan = _type_plan(source_rows - matched_rows)
    scenarios = ["exact_seed", "space_punctuation", "standard_format", "missing_secondary", "combined"]
    scenario_weights = [30, 20, 15, 15, 20]

    sources: list[dict[str, str]] = []
    truth: list[dict[str, str]] = []
    serial = 1
    for material_type in SUPPORTED_TYPES:
        type_pairs = by_type[material_type]
        for i in range(matched_plan[material_type]):
            pair = type_pairs[i % len(type_pairs)]
            row = canonical_source(pair.source)
            row["物料编码"] = _stable_material_code(serial)
            scenario = rng.choices(scenarios, weights=scenario_weights, k=1)[0]
            _format_noise(row, scenario, rng)
            target = canonical_target(pair.target)
            sources.append(row)
            truth.append({
                "物料编码": row["物料编码"], "预期是否可匹配": "Y", "预期集团码": target["集团码"],
                "物料类型": material_type, "场景": scenario, "种子源文件": pair.source.file_name,
                "种子源行号": str(pair.source.row_number), "种子目标文件": pair.target.file_name,
                "种子目标行号": str(pair.target.row_number), "种子配对分": f"{pair.score:.4f}",
                "说明": "真实 origin_data.zip 高置信源/目标配对扩增",
            })
            serial += 1
        for i in range(unmatched_plan[material_type]):
            pair = type_pairs[i % len(type_pairs)]
            row = canonical_source(pair.source)
            row["物料编码"] = _stable_material_code(serial)
            _make_unmatched(row, material_type, i + 1)
            sources.append(row)
            truth.append({
                "物料编码": row["物料编码"], "预期是否可匹配": "N", "预期集团码": "",
                "物料类型": material_type, "场景": "unmatched_variant", "种子源文件": pair.source.file_name,
                "种子源行号": str(pair.source.row_number), "种子目标文件": pair.target.file_name,
                "种子目标行号": str(pair.target.row_number), "种子配对分": f"{pair.score:.4f}",
                "说明": "由真实种子修改关键身份/规格得到目录外变体，预期人工选择均不匹配",
            })
            serial += 1

    combined = list(zip(sources, truth))
    rng.shuffle(combined)
    sources = [item[0] for item in combined]
    truth = [item[1] for item in combined]

    target_by_code: dict[str, dict[str, str]] = {}
    for seed_row in target_seeds:
        canonical = canonical_target(seed_row)
        code = canonical["集团码"]
        if code and code not in target_by_code:
            target_by_code[code] = canonical
    targets = list(target_by_code.values())

    manifest = {
        "generator": "origin_data.zip seed-driven CSV fixture",
        "seed": seed,
        "source_rows": len(sources),
        "target_rows": len(targets),
        "expected_matched": matched_rows,
        "expected_unmatched": source_rows - matched_rows,
        "match_ratio": round(matched_rows / source_rows, 4),
        "source_type_distribution": {t: sum(row["物料类型"] == t for row in sources) for t in SUPPORTED_TYPES},
        "target_type_distribution": {t: sum(row["物料类型"] == t for row in targets) for t in SUPPORTED_TYPES},
        "high_confidence_seed_pairs": len(pairs),
        "high_confidence_seed_pairs_by_type": {t: len(by_type[t]) for t in SUPPORTED_TYPES},
        "pair_threshold": PAIR_THRESHOLD,
        "scenario_distribution": {name: sum(row["场景"] == name for row in truth) for name in scenarios + ["unmatched_variant"]},
        **inventory,
    }
    return sources, targets, truth, manifest


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def generate_dataset(seed_zip: Path, output_dir: Path, *, source_rows: int = DEFAULT_SOURCE_ROWS, matched_rows: int = DEFAULT_MATCHED_ROWS, seed: int = SEED) -> dict[str, Path]:
    sources, targets, truth, manifest = generate_rows(seed_zip, source_rows=source_rows, matched_rows=matched_rows, seed=seed)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "source": output_dir / "source_materials_1000.csv",
        "target": output_dir / "target_group_codes_from_origin.csv",
        "truth": output_dir / "ground_truth.csv",
        "manifest": output_dir / "manifest.json",
    }
    _write_csv(paths["source"], SOURCE_HEADERS, sources)
    _write_csv(paths["target"], TARGET_HEADERS, targets)
    _write_csv(paths["truth"], TRUTH_HEADERS, truth)
    paths["manifest"].write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return paths


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a 1000-row material matching fixture from origin_data.zip")
    parser.add_argument("--seed-zip", type=Path, default=Path(__file__).resolve().parents[1] / "origin_data.zip")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "realistic_materials" / "generated")
    parser.add_argument("--source-rows", type=int, default=DEFAULT_SOURCE_ROWS)
    parser.add_argument("--matched-rows", type=int, default=DEFAULT_MATCHED_ROWS)
    parser.add_argument("--seed", type=int, default=SEED)
    args = parser.parse_args()
    paths = generate_dataset(args.seed_zip, args.output_dir, source_rows=args.source_rows, matched_rows=args.matched_rows, seed=args.seed)
    print(json.dumps({key: str(value) for key, value in paths.items()}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
