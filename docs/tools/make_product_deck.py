import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

NAVY = RGBColor(0x1F, 0x4E, 0x79); BLUE = RGBColor(0x2E, 0x75, 0xB6); LIGHT = RGBColor(0xDD, 0xEB, 0xF7)
GREEN = RGBColor(0x54, 0x82, 0x35); LGREEN = RGBColor(0xE2, 0xEF, 0xDA); AMBER = RGBColor(0xBF, 0x8F, 0x00)
YELL = RGBColor(0xFF, 0xD9, 0x66); BG = RGBColor(0xF4, 0xF6, 0xFA); DARK = RGBColor(0x33, 0x33, 0x33); GRAY = RGBColor(0x88, 0x88, 0x88)
FONT = "微软雅黑"
A = Path(os.environ.get("DECK_ASSETS", "assets")); EV = Path(os.environ.get("DECK_EVIDENCE", "docs/deploy-evidence/2026-09-25-matbjlzc-demo"))

prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5
total = [0]

def slide():
    s = prs.slides.add_slide(BLANK); total[0] += 1
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = BG; r.line.fill.background(); r.shadow.inherit = False
    return s

def box(s, x, y, w, h, fill=None, line=None, round_=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try: shp.adjustments[0] = 0.08
        except Exception: pass
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line; shp.line.width = Pt(1.5)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, size=18, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, leading=1.15):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    if isinstance(runs, str): runs = [runs]
    for i, ln in enumerate(runs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align; para.line_spacing = leading
        parts = ln if isinstance(ln, list) else [(ln, {})]
        for t, st in parts:
            r = para.add_run(); r.text = t
            r.font.name = FONT; r.font.size = Pt(st.get("size", size))
            r.font.bold = st.get("bold", bold); r.font.color.rgb = st.get("color", color)
    return tb

def header(s, title, sub=None, chapter=None):
    box(s, 0, 0, W, 1.05, fill=NAVY)
    box(s, 0, 1.05, W, 0.06, fill=YELL)
    text(s, 0.55, 0.16, 9.6, 0.8, title, size=27, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True, anchor=MSO_ANCHOR.MIDDLE)
    if chapter:
        text(s, 10.2, 0.16, 2.7, 0.8, chapter, size=13, color=LIGHT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(s, 0.55, 1.25, 12.2, 0.5, sub, size=14, color=GRAY)
    text(s, 0.55, 7.05, 6, 0.4, "小罡 AI · 物料集团码智能匹配平台", size=10, color=GRAY)
    text(s, 12.2, 7.05, 0.8, 0.4, str(total[0]), size=10, color=GRAY, align=PP_ALIGN.RIGHT)

def pic_fit(s, path, x, y, w, h):
    from PIL import Image
    im = Image.open(path); iw, ih = im.size
    scale = min(w / iw * 914400 * 0, 1) if False else min(w / (iw), h / (ih))
    ar = iw / ih
    bw, bh = (w, w / ar) if w / ar <= h else (h * ar, h)
    s.shapes.add_picture(str(path), Inches(x + (w - bw) / 2), Inches(y + (h - bh) / 2), Inches(bw), Inches(bh))

def bullets(s, x, y, w, items, size=17, gap=0.52, mark_color=BLUE):
    for i, it in enumerate(items):
        yy = y + i * gap
        box(s, x, yy + 0.10, 0.09, 0.09, fill=mark_color)
        if isinstance(it, tuple):
            head, desc = it
            text(s, x + 0.25, yy, w - 0.25, gap, [[(head + "  ", {"bold": True, "color": NAVY}), (desc, {})]], size=size)
        else:
            text(s, x + 0.25, yy, w - 0.25, gap, it, size=size)

def kpi(s, x, y, w, h, num, label, color=NAVY):
    box(s, x, y, w, h, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT, round_=True)
    text(s, x, y + 0.18, w, 0.85, num, size=34, color=color, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, y + 1.05, w, h - 1.1, label, size=13, color=DARK, align=PP_ALIGN.CENTER)

def placeholder(s, x, y, w, h, note="此处可替换：客户现场照片 / 系统截图"):
    box(s, x, y, w, h, fill=RGBColor(0xFF, 0xFF, 0xFF), line=GRAY, round_=True)
    text(s, x, y + h / 2 - 0.3, w, 0.6, note, size=14, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def divider(no, title, sub):
    s = slide()
    box(s, 0, 0, W, H, fill=NAVY)
    box(s, 0, 4.62, W, 0.06, fill=YELL)
    text(s, 0.9, 2.1, 3, 1.6, no, size=88, color=BLUE, bold=True)
    text(s, 0.95, 3.6, 11.4, 1.0, title, size=40, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    text(s, 0.98, 4.85, 11.4, 0.7, sub, size=17, color=LIGHT)
    text(s, 12.2, 7.05, 0.8, 0.4, str(total[0]), size=10, color=LIGHT, align=PP_ALIGN.RIGHT)

def img_slide(title, img, caption, chapter, sub=None):
    s = slide(); header(s, title, sub, chapter)
    pic_fit(s, img, 0.55, 1.55, 12.25, 4.9)
    text(s, 0.55, 6.55, 12.25, 0.5, caption, size=13, color=GRAY, align=PP_ALIGN.CENTER)

def table_slide(title, rows, chapter, col_w=None, y0=1.7, sub=None):
    s = slide(); header(s, title, sub, chapter)
    x0 = 0.7
    col_w = col_w or [3.2, 9.4]
    from pptx.util import Inches as I
    gtab = s.shapes.add_table(len(rows), len(rows[0]), I(x0), I(y0), I(sum(col_w)), I(0.52 * len(rows))).table
    for ci, cw in enumerate(col_w): gtab.columns[ci].width = I(cw)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gtab.cell(ri, ci); cell.text = str(val)
            for pr in cell.text_frame.paragraphs:
                for r in pr.runs:
                    r.font.name = FONT; r.font.size = Pt(14 if ri else 15)
                    r.font.bold = (ri == 0); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else DARK
            cell.fill.solid(); cell.fill.fore_color.rgb = NAVY if ri == 0 else (RGBColor(0xFF, 0xFF, 0xFF) if ri % 2 else LIGHT)
            cell.margin_left = I(0.12); cell.margin_top = I(0.05); cell.margin_bottom = I(0.05)
    return s

# ================= 内容 =================
# 1 封面
s = slide()
box(s, 0, 0, W, H, fill=NAVY)
box(s, 0, 4.9, W, 0.07, fill=YELL)
text(s, 0.9, 1.7, 11.5, 1.2, "物料集团码智能匹配平台", size=52, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
text(s, 0.95, 3.0, 11.5, 0.7, "SAP 物料与集团编码的批量匹配 · 人工复核 · 全程留痕", size=24, color=YELL)
text(s, 0.95, 3.85, 11.5, 0.6, "产品介绍 · 功能说明 · 使用向导", size=17, color=LIGHT)
placeholder(s, 8.6, 5.35, 4.0, 1.5, "此处可替换：单位 LOGO / 主视觉")
text(s, 0.95, 5.5, 7, 0.5, "小罡 AI · 版本 v1.3.21", size=15, color=LIGHT)
text(s, 0.95, 6.1, 7, 0.5, "离线内网部署 · 数据不出域", size=15, color=LIGHT)

# 2 目录
s = slide(); header(s, "目录", None, None)
items = [("01", "产品概述", "背景痛点 · 定位 · 核心价值 · 适用对象"),
         ("02", "系统架构", "离线部署 · 数据安全 · 交付与升级"),
         ("03", "核心功能", "匹配方案 · 引擎 · 人工复核 · 结果输出"),
         ("04", "使用向导", "六步上手 · 系统管理 · 常见问题")]
for i, (no, t, d) in enumerate(items):
    y = 1.7 + i * 1.32
    box(s, 0.8, y, 11.7, 1.1, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT, round_=True)
    text(s, 1.1, y + 0.12, 1.2, 0.8, no, size=30, color=BLUE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.5, y + 0.10, 4, 0.55, t, size=20, color=NAVY, bold=True)
    text(s, 2.5, y + 0.60, 9.5, 0.45, d, size=13, color=GRAY)

# 章节一
divider("01", "产品概述", "内网环境下的物料编码映射：批量初筛 · 人工复核 · 全程留痕")
# 4 背景痛点
s = slide(); header(s, "背景与痛点", None, "01 产品概述")
bullets(s, 0.8, 1.7, 11.7, [
    ("两套编码体系长期并存", "SAP 物料主数据与集团统一物料编码相互独立，存量映射关系长期未补齐"),
    ("人工对照周期长", "数万行物料对百万级集团码库，逐行比对周期长，且不同人员口径难以统一"),
    ("表格离线传递", "Excel 在多人间往返流转，版本易混乱、过程难追溯、成果难复用"),
    ("数据安全红线", "物料数据涉及企业核心资产，不允许上公网、不允许调用外部云服务"),
], size=17, gap=1.18)
box(s, 0.8, 6.15, 11.7, 0.75, fill=LGREEN, round_=True)
text(s, 1.1, 6.28, 11.2, 0.5, "应对方式：机器完成批量初筛与分流，人工聚焦不确定记录；全过程留痕，数据不出内网。", size=16, color=GREEN, bold=True)
# 5 定位
s = slide(); header(s, "产品定位", None, "01 产品概述")
text(s, 0.8, 1.8, 11.7, 1.0, [[("面向集团客户的", {}), ("内网离线部署", {"bold": True, "color": NAVY}), ("的物料编码智能映射与人工复核一体化平台", {})]], size=26)
for i, (t, d) in enumerate([("语义匹配引擎", "本地大模型 BGE 向量化 + ANN 召回 + 逐字段加权精排"),
                            ("人机协同闭环", "系统完成初筛与三分流，人工聚焦待复核记录；复核结论可沉淀为值映射与同义词"),
                            ("交付与运维", "介质离线安装；增量升级具备预检、备份与回滚；无外网、无云依赖")]):
    x = 0.8 + i * 4.0
    box(s, x, 3.3, 3.7, 2.9, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT, round_=True)
    box(s, x, 3.3, 3.7, 0.16, fill=[BLUE, GREEN, AMBER][i])
    text(s, x + 0.25, 3.7, 3.2, 0.6, t, size=19, color=NAVY, bold=True)
    text(s, x + 0.25, 4.4, 3.2, 1.6, d, size=14, color=DARK)
# 6 核心价值
s = slide(); header(s, "核心价值", None, "01 产品概述")
kpi(s, 0.8, 1.8, 2.85, 1.9, "批量初筛", "系统完成全量初筛与分流\n人工聚焦待复核记录")
kpi(s, 3.85, 1.8, 2.85, 1.9, "百万级", "目标库支撑规模\n（同构环境实测见后页）", BLUE)
kpi(s, 6.9, 1.8, 2.85, 1.9, "全程留痕", "确认/改判/导入均记录\n操作人 · 时间 · 依据", GREEN)
kpi(s, 9.95, 1.8, 2.55, 1.9, "无外网依赖", "模型与依赖离线分发\n数据不出内网", AMBER)
box(s, 0.8, 4.1, 11.7, 2.6, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT, round_=True)
text(s, 1.1, 4.3, 11, 0.5, "对业务人员的意义", size=17, color=NAVY, bold=True)
bullets(s, 1.1, 4.9, 11.2, [
    "结果 Excel 便于核对：映射字段成对并排、一致性四色标注，减少逐列翻找；",
    "阈值可调：高置信区间可批量确认，支持在线复核与离线 Excel 协作两种方式；",
    "复核结论可沉淀：值映射、同义词一次配置，后续任务自动生效。",
], size=15, gap=0.55)
# 7 适用对象
table_slide_rows = [
    ["角色", "使用方式"],
    ["物资/主数据管理员", "维护匹配方案、同义词与值映射；发起任务、审定结果"],
    ["各专业工程师", "第三步“人工调整”处理本领域待复核记录；离线 Excel 协作"],
    ["信息化/运维", "一台内网服务器安装即用；增量升级、备份回滚、账号管理"],
    ["集团报送接口人", "下载正式结果 Excel，直接对接集团码报送口径"],
]
table_slide("适用对象与典型场景", table_slide_rows, "01 产品概述", col_w=[3.4, 8.4])
# 8 数据说话
rows = [["场景", "旧方式痛点", "本平台实测"],
        ["90 万行集团码上传", "上传期间界面无响应", "约 22 秒完成，期间界面可持续操作"],
        ["发起匹配", "长时间等待且无状态反馈", "即时返回，任务状态（冻结/排队/运行）实时可见"],
        ["4 万×90 万匹配（冷跑）", "—", "约 1 小时（含一次性索引构建）；同库热跑约 40-60 分钟"],
        ["生成正式结果", "导出期间界面阻塞", "转后台导出（实测 46 秒），期间可离开页面"],
        ["结果 Excel（4 万行）", "生成缓慢", "36 秒生成 28.5MB，成对四色版式"]]
table_slide("系统处理时间实测记录", rows, "01 产品概述", col_w=[3.4, 3.9, 4.9], sub="交付版本在同构内网环境的实测；为系统处理时间，不含人工复核耗时，整体周期取决于数据规模与待复核量")

# 章节二
divider("02", "系统架构", "离线介质部署 · 浏览器访问 · 无外网依赖")
img_slide("离线内网部署架构", A / "arch.png", "介质一次拷贝上机 → 一键安装（Docker 推荐 / systemd 可选）→ 浏览器访问", "02 系统架构")
s = slide(); header(s, "数据安全与合规", None, "02 系统架构")
bullets(s, 0.8, 1.7, 11.7, [
    ("网络环境", "系统部署于客户内网，无外部网络调用；模型、依赖、升级包均离线分发"),
    ("数据不出域", "物料数据、集团码库、匹配结果仅存于本地 SQLite 与文件目录"),
    ("账号与审计", "登录鉴权、角色控制；关键操作（定稿/导入/升级）全部留痕可查"),
    ("可验证交付", "介质与增量包均带 SHA256 校验；升级前自动备份，可一键回滚"),
], size=17, gap=1.05)
s = slide(); header(s, "交付与升级体系", None, "02 系统架构")
for i, (t, d, c) in enumerate([
    ("全量介质", "DVD/U 盘：Docker 镜像+模型+依赖+安装脚本；新装一条命令", BLUE),
    ("增量升级包", "小版本升级专用：预检版本 → 自动备份 → 替换 → 健康检查 → 可回滚", GREEN),
    ("版本可追溯", "页面左下角实时显示版本号；每个包锚定代码提交与验收清单", AMBER)]):
    x = 0.8 + i * 4.0
    box(s, x, 1.8, 3.7, 3.4, fill=RGBColor(0xFF, 0xFF, 0xFF), line=LIGHT, round_=True)
    box(s, x, 1.8, 3.7, 0.16, fill=c)
    text(s, x + 0.25, 2.15, 3.2, 0.6, t, size=19, color=NAVY, bold=True)
    text(s, x + 0.25, 2.85, 3.2, 2.2, d, size=14)
box(s, 0.8, 5.5, 11.7, 1.2, fill=LIGHT, round_=True)
text(s, 1.1, 5.7, 11.2, 0.9, "当前交付版本：v1.3.21（含现场反馈的九项体验与性能修复、结果版式优化、大数据量稳定性加固）", size=16, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# 章节三
divider("03", "核心功能", "方案可复用 · 批量计算 · 人工闭环 · 结果可溯")
img_slide("智能匹配全流程", A / "pipeline.png", "两份 Excel 进，五表成对结果出；人工判断沉淀回方案", "03 核心功能")
img_slide("匹配方案管理", A / "03-profiles.png", "内置 6 个领域方案（元器件/紧固件/金属材料/非金属/复合材料/跨类目），支持复制、发布、版本化", "03 核心功能")
img_slide("字段映射与权重", A / "04-profile-editor.png", "源/目标字段对应关系、匹配器（精确/模糊/语义/包含）、权重合计 100%，支持固定值与多字段拼接", "03 核心功能")
rows = [["源字段值（SAP）", "映射后", "集团码库值", "结果"],
        ["10", "国产", "国产", "一致 ✓"],
        ["11", "进口", "进口", "一致 ✓"],
        ["（空）", "—", "国产", "缺失（灰）"]]
vs = table_slide("值映射：让编码口径自动对齐", rows, "03 核心功能", col_w=[3.2, 2.4, 2.8, 2.6], y0=1.9, sub="示例取自内置方案 A001（国产/进口：10→国产、11→进口）")
text(vs, 0.8, 4.5, 11.7, 0.5, "配置入口：方案/任务第一步 → 字段行“值映射”（支持连续录入、候选值下拉）", size=15, color=DARK)
text(vs, 0.8, 5.05, 11.7, 1.3, ["值映射只作用于打分对比，不篡改任何原始数据；", "映射关系同步体现在结果 Excel 的成对列与四色标注中。"], size=15, color=GRAY)
s = slide(); header(s, "同义词与文本清洗", None, "03 核心功能")
bullets(s, 0.8, 1.7, 11.7, [
    ("同义词表", "内置行业同义词（如 六角螺栓↔外六角螺栓），支持在线维护，即时生效于语义打分"),
    ("清洗管线", "全半角、繁简、空格标点归一；字段可挂多步清洗算子"),
    ("固定值/拼接", "一侧可用固定值或多字段拼接参与比对，兼容不同建库习惯"),
], size=17, gap=1.1)
s = slide(); header(s, "跨类目组合方案", "一份 SAP 源文件 × 多个集团码目标文件，一次任务覆盖多类目", "03 核心功能")
bullets(s, 0.8, 1.7, 11.7, [
    ("适用场景", "物资类其他（A007）等横跨多个专业领域的物料，逐领域分别匹配"),
    ("操作方式", "上传一份待匹配源数据 + 各子方案对应的集团码文件，系统按子方案自动路由"),
    ("结果统一", "五表版式不变，候选列标注来源子方案，复核与定稿一次完成"),
], size=17, gap=1.05)
placeholder(s, 0.8, 4.6, 11.7, 1.9, "此处可替换：跨类目绑定界面截图")
img_slide("数据管理", A / "05-data.png", "上传文件、集团码目录版本、匹配方案、同义词集中管理，来源可溯", "03 核心功能")
img_slide("任务四步向导", A / "07-task-progress.png", "第一步上传 → 第二步进度 → 第三步人工调整 → 第四步输出，全程状态实时（冻结中/排队中/运行中）", "03 核心功能")
s = slide(); header(s, "进度监控与大数据性能", None, "03 核心功能")
rows = [["能力", "说明"],
        ["实时看板", "已处理行数 / 总数、阶段（索引·召回·精排·落库）、每分钟速率、预计完成时间"],
        ["断点恢复", "服务重启不丢任务，运行中任务自动恢复续算"],
        ["并发友好", "重活全部后台化：上传、冻结、匹配、导出期间页面始终可操作"],
        ["实测规模", "90 万行目标上传约 22 秒；4 万×90 万冷跑约 1 小时（含一次性索引）、热跑约 40-60 分钟（同构环境）"]]
table_slide("", rows, "03 核心功能", col_w=[2.8, 9.0])
img_slide("第三步 · 人工调整（卡片式复核）", A / "08-review.png", "左侧候选列表 + 右侧逐字段对比（成对四色），确认 / 改选 / 标记均不匹配，支持备注", "03 核心功能")
s = slide(); header(s, "批量操作与阈值重判", None, "03 核心功能")
bullets(s, 0.8, 1.7, 11.7, [
    ("批量确认 Top1", "对高置信区间批量确认，人工聚焦待复核记录"),
    ("阈值重判", "调整自动阈值即时重算三分流，结果版本随之演进、旧版本可查"),
    ("检索与筛选", "按状态/关键字/来源行号定位待办，处理进度实时汇总"),
    ("操作留痕", "每条决策记录操作人、时间、依据，进入结果“人工操作记录”表"),
], size=17, gap=1.05)
img_slide("离线协作 · 人工匹配 Excel", EV / "shot-manual-main.png", "“人工选择”列在最左带下拉（候选1-5/均不匹配）；源与候选字段成对并排四色标注", "03 核心功能")
s = slide(); header(s, "多人协作与冲突保护", None, "03 核心功能")
bullets(s, 0.8, 1.7, 11.7, [
    ("分工灵活", "结果 Excel 可复制拆分为多份，分头处理、分别回传"),
    ("幂等合并", "重复上传相同结论自动跳过，不产生重复记录"),
    ("冲突不覆盖", "同一行给出不同结论时拒绝覆盖并明确提示，由管理员裁决"),
    ("即传即生效", "导入后工作台与结果页即时更新，全程无需重跑匹配"),
], size=17, gap=1.05)
img_slide("正式结果 · 匹配摘要页", EV / "shot-result-summary.png", "方案名称、任务信息、四类计数、输入文件与图例，一页看全", "03 核心功能")
img_slide("正式结果 · 成对四色版式", EV / "shot-result-final.png", "匹配结果区在前；源·X 与 候选·X 按权重成对并排；绿=一致 黄=部分 橙=不一致 灰=缺失；未参与字段沉入完整数据专区", "03 核心功能")
img_slide("正式结果 · Top5 候选", EV / "shot-result-top5.png", "每个源行的前五候选完整呈现（排名/集团码/相似度/成对字段），落选依据可复核", "03 核心功能")
rows = [["工作表", "内容"],
        ["匹配摘要", "任务与方案快照、统计、图例"],
        ["最终匹配结果", "逐行结论 + 成对字段四色 + 操作人/时间"],
        ["Top5候选", "全部候选及相似度，支撑抽查"],
        ["人工操作记录", "谁在何时改了什么、备注原因"],
        ["未匹配清单", "无结论行集中复核清单"]]
table_slide("结果追溯与审计", rows, "03 核心功能", col_w=[3.2, 8.6])

# 章节四
divider("04", "使用向导", "六步完成一次匹配任务")
img_slide("第 0 步 · 登录", A / "01-login.png", "浏览器打开 http://服务器地址:18080，使用管理员分配的账号登录", "04 使用向导")
s = slide(); header(s, "第 1 步 · 准备匹配方案", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "新建方案或复制内置领域方案，选择源/目标字段建立对应关系",
    "为每条映射设置匹配器与权重（合计 100%），必要时开启关键冲突",
    "配置值映射与同义词，发布方案供任务引用",
], size=17, gap=1.0)
placeholder(s, 0.8, 4.6, 11.7, 1.9, "此处可替换：方案配置界面大图（演示环境截图）")
s = slide(); header(s, "第 2 步 · 上传数据并发起任务", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "第一步上传：待匹配 SAP 文件 + 集团码标准文件（跨类目方案则一源多目标）",
    "确认字段解析与行数估算后点击“开始匹配”，任务即时进入队列",
    "大文件（数十万行）上传期间界面可持续操作，可并行处理其它任务",
], size=17, gap=1.0)
placeholder(s, 0.8, 4.6, 11.7, 1.9, "此处可替换：上传页截图（拖拽上传状态）")
s = slide(); header(s, "第 3 步 · 监控进度", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "第二步实时看板：阶段、速率、预计完成时间持续刷新",
    "已匹配结果滚动可见，无需等待任务结束",
    "服务重启后，运行中任务自动恢复续算，无需人工干预",
], size=17, gap=1.0)
placeholder(s, 0.8, 4.6, 11.7, 1.9, "此处可替换：运行中任务看板截图")
s = slide(); header(s, "第 4 步 · 人工调整", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "在线：卡片逐条处理（确认/改选/标记不匹配），或批量确认高置信 Top1",
    "离线：下载“人工匹配 Excel”，填最左侧下拉列，回传即合并（冲突保护）",
    "暂无法判定的记录可先跳过；定稿时系统会明确提示未处理行数及其结果",
], size=17, gap=1.0)
s = slide(); header(s, "第 5 步 · 定稿与下载", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "点击“生成最终结果”：后台导出，页面可自由离开",
    "完成后第四步下载正式结果 Excel（五表成对四色版式）",
    "阈值调整、追加复核后重新定稿，自动形成新版本，历史可查",
], size=17, gap=1.0)
s = slide(); header(s, "第 6 步 · 沉淀与复用", None, "04 使用向导")
bullets(s, 0.8, 1.7, 11.7, [
    "把复核中发现的口径差异固化为值映射/同义词——下次任务自动生效",
    "方案支持复制改版，跨类目组合方案覆盖杂项物资",
    "任务、结果、操作记录长期保留，支撑审计与复盘",
], size=17, gap=1.0)
img_slide("系统管理", A / "06-system.png", "账号、参数与运行状态集中管理（管理员）", "04 使用向导")
rows = [["问题", "回答"],
        ["匹配结果不准怎么办？", "调权重/阈值即时重判；把差异配成值映射；关键冲突字段单独设防"],
        ["大数据量是否影响操作？", "上传、匹配、导出均在后台执行，实测 90 万行上传期间界面可持续操作"],
        ["能否完全离线使用？", "可以，系统面向纯内网环境设计，无外部网络依赖"],
        ["如何升级？", "增量包一条命令升级，自动备份、失败自动回滚，数据零丢失"],
        ["结果能直接报送集团吗？", "正式结果 Excel 含集团码/相似度/依据字段，可直接对接报送口径"]]
table_slide("常见问题", rows, "04 使用向导", col_w=[4.2, 7.6])
# 结尾
s = slide()
box(s, 0, 0, W, H, fill=NAVY)
box(s, 0, 4.3, W, 0.06, fill=YELL)
text(s, 0.9, 2.5, 11.5, 1.2, "批量计算 · 人工复核 · 全程留痕 · 数据不出域", size=36, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
text(s, 0.92, 4.7, 11.5, 0.6, "小罡 AI · 物料集团码智能匹配平台 v1.3.21", size=18, color=LIGHT)
placeholder(s, 8.9, 5.4, 3.7, 1.3, "此处可替换：联系方式 / 二维码")
text(s, 0.92, 5.5, 7.5, 0.5, "谢谢", size=22, color=YELL, bold=True)

out = Path("/oracle/codex/work/MATERIAL/media/物料集团码智能匹配平台-产品介绍-v1.3.21.pptx")
prs.save(out)
print("SLIDES:", total[0], "->", out)
