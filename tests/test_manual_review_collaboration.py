from __future__ import annotations

from io import BytesIO
import json
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _seed_review_task(client: TestClient, task_id: str = "manual-task") -> str:
    meta = client.app.state.meta
    now = "2026-09-16T15:20:08+08:00"
    with meta.connect() as connection:
        connection.execute(
            """INSERT INTO tasks(
                task_id,name,source_file_id,catalog_version_id,profile_id,profile_version,
                config_snapshot,config_sha256,stage,status,progress,processed_rows,total_rows,
                created_at,started_at,finished_at,error_code,error_message,result_file_id,
                created_by,started_by
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                task_id,
                "人工协作测试",
                "source-file",
                "catalog-version",
                None,
                None,
                "{}",
                "sha",
                "REVIEW",
                "COMPLETED",
                100.0,
                3,
                3,
                now,
                now,
                now,
                None,
                None,
                None,
                "admin",
                "admin",
            ),
        )
        for index, original_status in ((1, "REVIEW"), (2, "REVIEW"), (3, "MATCHED")):
            source_row_id = str(index)
            source_payload = {"物料编码": f"S{index:03d}", "物料描述": f"源物料{index}", "单位": "EA", "规格": f"M{index}"}
            final_group = f"G{index}1" if original_status == "MATCHED" else None
            connection.execute(
                """INSERT INTO match_items(
                    task_id,source_row_id,source_row_number,source_id,source_payload,
                    original_status,current_status,top1_group_code,top1_score,second_score,
                    score_gap,critical_conflict,final_group_code,created_at,updated_at
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    task_id,
                    source_row_id,
                    index + 10,
                    f"S{index:03d}",
                    json.dumps(source_payload, ensure_ascii=False),
                    original_status,
                    original_status,
                    f"G{index}1",
                    96.0 - index,
                    90.0 - index,
                    6.0,
                    0,
                    final_group,
                    now,
                    now,
                ),
            )
            for rank in range(1, 6):
                target_payload = {"集团码": f"G{index}{rank}", "物料描述": f"目标物料{index}-{rank}", "单位": "EA", "规格": f"T{rank}"}
                connection.execute(
                    """INSERT INTO match_candidates(
                        task_id,source_row_id,rank,target_row_number,target_group_code,
                        target_payload,score,field_scores,critical_conflict
                    ) VALUES(?,?,?,?,?,?,?,?,?)""",
                    (
                        task_id,
                        source_row_id,
                        rank,
                        1000 + index * 10 + rank,
                        f"G{index}{rank}",
                        json.dumps(target_payload, ensure_ascii=False),
                        99.0 - rank,
                        "{}",
                        0,
                    ),
                )
    return task_id


def _selection_workbook(payload: bytes, selections: dict[str, str]) -> bytes:
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["人工匹配"]
    headers = {str(cell.value): cell.column for cell in sheet[1] if cell.value is not None}
    for row in range(2, sheet.max_row + 1):
        source_row_id = str(sheet.cell(row, headers["__source_row_id"]).value)
        sheet.cell(row, headers["人工选择"]).value = selections.get(source_row_id)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_manual_review_excel_contains_full_business_context_and_hidden_ids(authed: TestClient) -> None:
    task_id = _seed_review_task(authed)
    response = authed.get(f"/api/tasks/{task_id}/manual-review.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert {"人工匹配", "填写说明"} <= set(workbook.sheetnames)
    sheet = workbook["人工匹配"]
    headers = [str(cell.value) for cell in sheet[1]]
    assert "源表原始行号" in headers
    assert "源.物料编码" in headers
    assert "源.物料描述" in headers
    assert "候选1.目标表原始行号" in headers
    assert "候选1.集团码" in headers
    assert "候选1.相似度" in headers
    assert "候选5.目标表原始行号" in headers
    assert "候选5.物料描述" in headers
    assert "人工选择" in headers
    task_col = headers.index("__task_id") + 1
    source_col = headers.index("__source_row_id") + 1
    assert sheet.column_dimensions[sheet.cell(1, task_col).column_letter].hidden is True
    assert sheet.column_dimensions[sheet.cell(1, source_col).column_letter].hidden is True
    assert sheet.max_row == 4
    workbook.close()


def test_multiple_manual_excel_uploads_merge_idempotently_and_report_conflicts(authed: TestClient) -> None:
    task_id = _seed_review_task(authed)
    template = authed.get(f"/api/tasks/{task_id}/manual-review.xlsx").content

    first = _selection_workbook(template, {"1": "候选2"})
    applied = authed.post(
        f"/api/tasks/{task_id}/manual-review/import",
        files={"file": ("worker-a.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert applied.status_code == 200
    assert applied.json()["success_count"] == 1
    assert applied.json()["conflict_count"] == 0

    repeated = authed.post(
        f"/api/tasks/{task_id}/manual-review/import",
        files={"file": ("worker-a-repeat.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert repeated.status_code == 200
    assert repeated.json()["success_count"] == 0
    assert repeated.json()["conflict_count"] == 0
    assert repeated.json()["skipped_count"] >= 1

    conflicting = _selection_workbook(template, {"1": "候选3"})
    conflict = authed.post(
        f"/api/tasks/{task_id}/manual-review/import",
        files={"file": ("worker-b.xlsx", conflicting, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert conflict.status_code == 200
    assert conflict.json()["conflict_count"] == 1
    assert conflict.json()["success_count"] == 0
    assert conflict.json()["details"][0]["code"] == "MANUAL_RESULT_CONFLICT"

    second_source = _selection_workbook(template, {"2": "均不匹配"})
    merged = authed.post(
        f"/api/tasks/{task_id}/manual-review/import",
        files={"file": ("worker-c.xlsx", second_source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert merged.status_code == 200
    assert merged.json()["success_count"] == 1

    logs = authed.get(f"/api/tasks/{task_id}/operations").json()
    assert logs["total"] == 2
    assert {item["operation_type"] for item in logs["items"]} == {"IMPORT_MATCH", "MARK_UNMATCHED"}
    assert {item["source"] for item in logs["items"]} == {"EXCEL"}
    assert {item["operator"] for item in logs["items"]} == {"admin"}

    with authed.app.state.meta.connect() as connection:
        row = connection.execute("SELECT current_status,final_group_code FROM match_items WHERE task_id=? AND source_row_id='1'", (task_id,)).fetchone()
    assert row["current_status"] == "CONFIRMED"
    assert row["final_group_code"] == "G12"


def test_web_manual_match_rematch_and_cancel_are_reversible_and_append_only(authed: TestClient) -> None:
    task_id = _seed_review_task(authed)

    confirmed = authed.post(f"/api/tasks/{task_id}/items/1/confirm", json={"target_id": "G11", "comment": "首次确认"})
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "CONFIRMED"

    rematched = authed.post(f"/api/tasks/{task_id}/items/1/rematch", json={"target_id": "G12", "comment": "改选候选2"})
    assert rematched.status_code == 200
    assert rematched.json()["final_group_code"] == "G12"

    cancelled = authed.post(f"/api/tasks/{task_id}/items/1/cancel-match", json={"comment": "撤销人工匹配"})
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "REVIEW"
    assert cancelled.json()["final_group_code"] is None

    rejected = authed.post(f"/api/tasks/{task_id}/items/1/mark-unmatched", json={"comment": "均不匹配"})
    assert rejected.status_code == 200
    assert rejected.json()["status"] == "UNMATCHED"

    cancel_unmatched = authed.post(f"/api/tasks/{task_id}/items/1/cancel-unmatched", json={"comment": "恢复待确认"})
    assert cancel_unmatched.status_code == 200
    assert cancel_unmatched.json()["status"] == "REVIEW"

    logs = authed.get(f"/api/tasks/{task_id}/items/1/operations").json()
    assert logs["total"] == 5
    assert [item["operation_type"] for item in reversed(logs["items"])] == [
        "MATCH",
        "REMATCH",
        "CANCEL_MATCH",
        "MARK_UNMATCHED",
        "CANCEL_UNMATCHED",
    ]
    assert all(item["source_row_number"] == 11 for item in logs["items"])
    assert all(item["operator"] == "admin" for item in logs["items"])
    assert all(item["source"] == "WEB" for item in logs["items"])

    with authed.app.state.meta.connect() as connection:
        review_count = int(connection.execute("SELECT COUNT(*) FROM reviews WHERE task_id=? AND source_row_id='1'", (task_id,)).fetchone()[0])
    assert review_count == 5


def test_task_start_records_logged_in_creator_and_starter(authed: TestClient) -> None:
    config = {
        "source_id_column": "物料号",
        "scope_mode": "GLOBAL",
        "rules": [
            {
                "id": "name",
                "source": {"fields": ["物料名称"]},
                "target": {"fields": ["物料名称"]},
                "matcher": "exact",
                "weight": 100,
            }
        ],
        "decision": {"success_threshold": 90, "review_enabled": True, "review_threshold": 70, "top_n": 5},
        "advanced": {},
    }
    now = "2026-09-16T15:20:08+08:00"
    with authed.app.state.meta.connect() as connection:
        connection.execute(
            "INSERT INTO task_drafts VALUES(?,?,?,?,?,?,?,?,?,?)",
            ("actor-draft", "发起人测试", "missing-source", "missing-catalog", None, None, json.dumps(config, ensure_ascii=False), 2, now, now),
        )
    response = authed.post("/api/task-drafts/actor-draft/start")
    assert response.status_code == 202
    task = response.json()
    assert task["created_by"] == "admin"
    assert task["started_by"] == "admin"
    assert task["created_at"]

    task_id = task["task_id"]
    for _ in range(100):
        current = authed.get(f"/api/tasks/{task_id}").json()
        if current.get("started_at"):
            break
        time.sleep(0.01)
    assert current["started_at"]
    assert current["created_by"] == "admin"
    assert current["started_by"] == "admin"
