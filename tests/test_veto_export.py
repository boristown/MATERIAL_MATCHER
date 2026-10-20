from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from material_matcher.matching.engine import CandidateResult, RowResult

from test_result_export_business import _row, _seed_task, _service


VETO_CONFIG = {
    "rules": [
        {"id": "name", "source": {"fields": ["名称"]}, "target": {"fields": ["名称"]}, "matcher": "exact", "weight": 40},
        {"id": "model", "source": {"fields": ["型号"]}, "target": {"fields": ["型号"]}, "matcher": "exact", "weight": 40, "veto_threshold": 60},
        {"id": "category", "source": {"fields": ["类别"]}, "target": {"fields": ["类别"]}, "matcher": "exact", "weight": 20},
    ],
    "decision": {"success_threshold": 80, "review_enabled": True, "review_threshold": 50, "top_n": 5},
    "retrieval": {"mode": "scan", "retrieval_top_k": 10, "oversample": 2},
    "advanced": {},
}


def _veto_candidate(*, model_score: float, model_before: float, name_score: float, category_score: float) -> CandidateResult:
    field_scores = [
        {"rule_id": "name", "score": name_score, "weight": 40, "source_value": "六角螺栓", "target_value": "六角螺栓"},
        {"rule_id": "model", "score": model_score, "weight": 40, "source_value": "M8X20", "target_value": "M8X22",
         "vetoed": True, "score_before_veto": model_before},
        {"rule_id": "category", "score": category_score, "weight": 20, "source_value": "A", "target_value": "A"},
    ]
    return CandidateResult(
        rank=1, group_code="0000456", score=0.0, raw_score=0.0, critical_conflict=False,
        target_payload={"集团码": "0000456", "名称": "六角螺栓", "型号": "M8X22", "类别": "A"},
        field_scores=field_scores, target_row_number=34567,
    )


def _veto_row(source_row_id: str, source_row_number: int, candidate: CandidateResult, status: str = "REVIEW") -> RowResult:
    return RowResult(
        source_row_id=source_row_id,
        source_id=f"00012{source_row_id}",
        source_payload={"物料号": f"00012{source_row_id}", "名称": "六角螺栓", "型号": "M8X20", "类别": "A"},
        status=status,
        final_group_code=None,
        first_score=0.0,
        second_score=0.0,
        score_gap=0.0,
        critical_conflict=False,
        candidates=[candidate],
        source_row_number=source_row_number,
    )


def _enable_veto_config(meta) -> None:
    with meta.connect() as connection:
        connection.execute(
            "UPDATE tasks SET config_snapshot=? WHERE task_id=?",
            (json.dumps(VETO_CONFIG, ensure_ascii=False), "business-result"),
        )


def test_veto_compare_sheet_keeps_pre_and_post_veto_rows(tmp_path: Path) -> None:
    service, meta, files = _service(tmp_path)
    _seed_task(meta)
    _enable_veto_config(meta)
    # A: 否决前 83.6 本可自动匹配却被否决；B: 否决前 64 转人工；C: 未被否决（不应出现在对照页）
    a = _veto_row("1", 8, _veto_candidate(model_score=0.0, model_before=0.59, name_score=1.0, category_score=1.0))
    b = _veto_row("2", 9, _veto_candidate(model_score=0.0, model_before=0.5, name_score=0.6, category_score=1.0))
    c = _row("3", 10, "MATCHED", "0000999", 34569, 88.0)
    service._persist_rows("business-result", [a, b, c])

    finalized = service.finalize("business-result", allow_unresolved_review=True)
    record = files.get(str(finalized["result_file_id"]))
    workbook = load_workbook(Path(str(record["stored_path"])), data_only=True)
    try:
        assert "否决对照" in workbook.sheetnames, workbook.sheetnames
        sheet = workbook["否决对照"]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        flat = [row for row in rows if row and row[0] is not None and str(row[0]).isdigit()]
        assert len(flat) == 2, flat
        first = flat[0]
        assert first[0] == 8 and str(first[1]) == "000121"
        assert first[2] == "0000456"
        assert abs(float(first[3]) - 83.6) < 0.1, first
        assert first[4] == "自动匹配", first
        assert first[5] == "待处理", first
        assert first[6] == "model"
        assert first[7] == "M8X20 → M8X22"
        assert str(first[8]) == "60"
        second = flat[1]
        assert abs(float(second[3]) - 64.0) < 0.1, second
        note = [row for row in rows if row and row[0] and "合计" in str(row[0])]
        assert note and "被否决 2 行" in str(note[0][0])
        assert "1 行将自动匹配" in str(note[0][0]) and "1 行进入人工处理" in str(note[0][0])
    finally:
        workbook.close()


def test_veto_sheet_absent_without_veto_rules(tmp_path: Path) -> None:
    service, meta, files = _service(tmp_path)
    _seed_task(meta)
    service._persist_rows("business-result", [_row("1", 8, "MATCHED", "0000456", 34567, 88.0)])
    finalized = service.finalize("business-result")
    record = files.get(str(finalized["result_file_id"]))
    workbook = load_workbook(Path(str(record["stored_path"])), data_only=True)
    try:
        assert "否决对照" not in workbook.sheetnames
    finally:
        workbook.close()