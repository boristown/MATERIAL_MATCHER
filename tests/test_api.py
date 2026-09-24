from __future__ import annotations
from hashlib import sha256
from io import BytesIO
import time
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


def workbook_bytes(rows=None) -> bytes:
    workbook=Workbook(); sheet=workbook.active; sheet.title='物料'; sheet.append(['说明']); sheet.append(['物料号','物料名称','型号'])
    for row in rows or [['000123','电阻','R10']]: sheet.append(row)
    output=BytesIO(); workbook.save(output); return output.getvalue()

def target_bytes(rows=None)->bytes:
    workbook=Workbook(); sheet=workbook.active; sheet.title='集团'; sheet.append(['集团码','物料名称','型号'])
    for row in rows or [['G001','电阻','R10']]: sheet.append(row)
    output=BytesIO(); workbook.save(output); return output.getvalue()

def setup_draft(authed: TestClient, source_rows=None, target_rows=None, success=90, review=70):
    source=authed.post('/api/files/upload',data={'role':'source'},files={'file':('source.xlsx',workbook_bytes(source_rows),'application/octet-stream')}).json()
    target=authed.post('/api/files/upload',data={'role':'target'},files={'file':('target.xlsx',target_bytes(target_rows),'application/octet-stream')}).json()
    catalog=authed.post('/api/catalogs',json={'name':'集团目录','source_file_id':target['file']['file_id'],'group_code_column':'集团码'}).json()
    draft=authed.post('/api/task-drafts',json={'name':'测试任务'}).json()
    authed.put(f"/api/task-drafts/{draft['draft_id']}/data",json={'source_file_id':source['file']['file_id'],'catalog_version_id':catalog['version_id']})
    rules={'source_id_column':'物料号','scope_mode':'GLOBAL','rules':[{'id':'name','source':{'fields':['物料名称']},'target':{'fields':['物料名称']},'matcher':'fuzzy','weight':80},{'id':'model','source':{'fields':['型号']},'target':{'fields':['型号']},'matcher':'exact','weight':100,'critical':True}], 'decision':{'success_threshold':success,'review_enabled':True,'review_threshold':review,'top_n':3},'advanced':{}}
    assert authed.put(f"/api/task-drafts/{draft['draft_id']}/rules",json=rules).status_code==200
    return draft, rules

def wait_task(authed: TestClient, task_id: str):
    for _ in range(200):
        task=authed.get(f'/api/tasks/{task_id}').json()
        if task['status'] in {'COMPLETED','FAILED'}: return task
        time.sleep(0.01)
    raise AssertionError('task timeout')

def test_auth_and_health(client: TestClient)->None:
    assert client.get('/api/health').status_code==200; assert client.get('/api/tasks').status_code==401
    bad=client.post('/api/auth/login',json={'username':'admin','password':'bad'}); assert bad.status_code==401; assert bad.json()['error']['code']=='AUTH_FAILED'

def test_draft_to_immutable_task_snapshot(authed: TestClient)->None:
    draft,rules=setup_draft(authed)
    started=authed.post(f"/api/task-drafts/{draft['draft_id']}/start"); assert started.status_code==202; task=started.json(); assert task['profile_id'] is None; assert task['config_snapshot']['rules'][0]['weight']==80; assert len(task['config_sha256'])==64
    rules['rules'][0]['weight']=20; authed.put(f"/api/task-drafts/{draft['draft_id']}/rules",json=rules)
    reloaded=authed.get(f"/api/tasks/{task['task_id']}").json(); assert reloaded['config_snapshot']['rules'][0]['weight']==80

def test_task_workspace_draft_survives_relogin_with_full_step1_state(authed: TestClient)->None:
    source=authed.post('/api/files/upload',data={'role':'source'},files={'file':('source.xlsx',workbook_bytes(),'application/octet-stream')}).json()
    target=authed.post('/api/files/upload',data={'role':'target'},files={'file':('target.xlsx',target_bytes(),'application/octet-stream')}).json()
    catalog=authed.post('/api/catalogs',json={'name':'恢复目录','source_file_id':target['file']['file_id'],'group_code_column':'集团码'}).json()
    draft=authed.post('/api/task-drafts',json={'name':'初始任务'}).json()
    document={
        'source_id_column':'物料号',
        'scope_mode':'GLOBAL',
        'rules':[{
            'id':'recover-rule',
            'source':{'fields':['物料名称','型号'],'combine':'concat','separator':' / ','pipeline':[{'op':'trim','options':{}}]},
            'target':{'fields':['物料名称'],'combine':'coalesce','separator':' ','pipeline':[]},
            'matcher':'hybrid','weight':73,'critical':True,'matcher_options':{'boost':1.1},
        }],
        'decision':{'success_threshold':91,'review_enabled':True,'review_threshold':76,'top_n':7},
        'retrieval':{'mode':'auto','provider':'onnx_local','model_id':'BAAI/bge-base-zh-v1.5','dimensions':768,'max_length':512,'precision':'fp32','retrieval_top_k':200,'oversample':4},
        'advanced':{'normalization':{'enabled':True}},
    }
    saved=authed.patch(f"/api/task-drafts/{draft['draft_id']}",json={
        'name':'可恢复任务',
        'source_file_id':source['file']['file_id'],
        'catalog_version_id':catalog['version_id'],
        'config_document':document,
    })
    assert saved.status_code==200
    # “任务名称” has been removed as a business concept: legacy name payloads are
    # accepted and ignored; the draft keeps its system-generated internal name.
    assert saved.json()['name']==draft['name'] and str(draft['name']).startswith('run-')
    saved_document=saved.json()['config_document']
    assert saved_document['decision']['success_threshold']==91
    assert saved_document['decision']['top_n']==7
    assert 'review_threshold' not in saved_document['decision']

    assert authed.post('/api/auth/logout').status_code==200
    assert authed.post('/api/auth/login',json={'username':'admin','password':'ChangedAdmin123'}).status_code==200
    restored=authed.get(f"/api/task-drafts/{draft['draft_id']}")
    assert restored.status_code==200
    body=restored.json()
    assert body['source_file_id']==source['file']['file_id']
    assert body['catalog_version_id']==catalog['version_id']
    assert body['config_document']['rules'][0]['source']['pipeline'][0]['op']=='trim'
    assert body['config_document']['rules'][0]['matcher']=='hybrid'
    assert body['config_document']['rules'][0]['weight']==73
    assert body['config_document']['rules'][0]['critical'] is True
    assert body['config_document']['decision']['success_threshold']==91
    assert body['config_document']['decision']['top_n']==7
    assert 'review_threshold' not in body['config_document']['decision']
    assert body['config_document']['retrieval']['max_length']==512

def test_chunk_upload_validates_hash_and_order(authed: TestClient)->None:
    payload=workbook_bytes(); digest=sha256(payload).hexdigest(); initialized=authed.post('/api/uploads/init',json={'role':'source','original_name':'chunked.xlsx','total_size':len(payload),'sha256':digest}).json(); upload_id=initialized['upload_id']
    wrong=authed.put(f'/api/uploads/{upload_id}/chunks/1',content=payload); assert wrong.status_code==409
    assert authed.put(f'/api/uploads/{upload_id}/chunks/0',content=payload).status_code==200; completed=authed.post(f'/api/uploads/{upload_id}/complete'); assert completed.status_code==200; assert completed.json()['file']['sha256']==digest

def test_old_xls_is_explicitly_rejected(authed: TestClient)->None:
    response=authed.post('/api/files/upload',data={'role':'source'},files={'file':('old.xls',b'fake','application/octet-stream')}); assert response.status_code==400; assert response.json()['error']['code']=='UNSUPPORTED_FILE'

def test_dry_run_returns_top2_gap_and_field_explanation(authed: TestClient)->None:
    draft,_=setup_draft(authed,source_rows=[['0001','电阻','R10']],target_rows=[['G1','电阻','R10'],['G2','电阻器','R11']])
    response=authed.post(f"/api/task-drafts/{draft['draft_id']}/dry-run",json={'sample_rows':100}); assert response.status_code==200
    body=response.json(); assert body['summary']['total']==1; row=body['rows'][0]; assert row['first_score']>=row['second_score']; assert row['score_gap']==round(row['first_score']-row['second_score'],4); assert len(row['candidates'][0]['field_scores'])==2

def test_persistent_worker_review_and_finalize_flow(authed: TestClient)->None:
    draft,rules=setup_draft(authed,source_rows=[['0001','电阻','R10'],['0002','电容','C10']],target_rows=[['G1','电阻','R10'],['G2','电容','C11']],success=95,review=40)
    started=authed.post(f"/api/task-drafts/{draft['draft_id']}/start").json(); task=wait_task(authed,started['task_id']); assert task['status']=='COMPLETED'; assert task['stage']=='REVIEW'
    summary=authed.get(f"/api/tasks/{task['task_id']}/workbench/summary").json(); assert summary['pending_review']>=1
    items=authed.get(f"/api/tasks/{task['task_id']}/workbench/items").json()['items']; assert items
    row_id=items[0]['source_row_id']; candidates=authed.get(f"/api/tasks/{task['task_id']}/items/{row_id}/candidates").json()['candidates']; assert candidates
    confirm=authed.post(f"/api/tasks/{task['task_id']}/items/{row_id}/confirm",json={'target_id':candidates[0]['target_group_code'],'comment':'测试确认'}); assert confirm.status_code==200
    premature=authed.post(f"/api/tasks/{task['task_id']}/finalize",json={'allow_unresolved_review':False})
    assert premature.status_code in (200, 202, 409)
    from conftest import finalize_wait
    finalize_wait(authed, task['task_id'])
    result=authed.get(f"/api/tasks/{task['task_id']}/result"); assert result.status_code==200; wb=load_workbook(BytesIO(result.content),read_only=True); assert {'匹配摘要','最终匹配结果','Top5候选','人工操作记录','未匹配清单'} <= set(wb.sheetnames); wb.close()


def test_finalize_after_review_bumps_result_revision(authed) -> None:
    draft, rules = setup_draft(authed, source_rows=[["0001", "电阻", "R10"], ["0002", "电容", "C10"], ["0003", "电感", "L5"]],
                               target_rows=[["G1", "电阻", "R10"], ["G2", "电容", "C11"], ["G3", "电感", "L5"]], success=95, review=40)
    started = authed.post(f"/api/task-drafts/{draft['draft_id']}/start").json()
    task = wait_task(authed, started["task_id"])
    from conftest import finalize_wait
    fin1 = finalize_wait(authed, task['task_id'])
    file1 = fin1["result_file_id"]
    items = authed.get(f"/api/tasks/{task['task_id']}/workbench/items?status=REVIEW&page_size=50").json()["items"]
    assert items, "预置待人工行用于二次确认"
    row = items[0]["source_row_id"]
    cands = authed.get(f"/api/tasks/{task['task_id']}/items/{row}/candidates").json()["candidates"]
    conf = authed.post(f"/api/tasks/{task['task_id']}/items/{row}/confirm", json={"target_id": cands[0]["target_group_code"], "comment": "定稿后追加确认"})
    assert conf.status_code == 200
    fin2 = finalize_wait(authed, task['task_id'])
    assert fin2["result_file_id"] != file1, "定稿后新增人工确认必须重新生成结果文件"
    fin3 = finalize_wait(authed, task['task_id'])
    assert fin3["result_file_id"] == fin2["result_file_id"] and fin3.get("reused") is True, "无新操作时保持幂等复用"
