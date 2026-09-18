from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
import sqlite3
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from material_matcher.services.task_service import format_duration_ms, task_time_fields
from material_matcher.storage.metadata import MetadataRepository


def _upload_csv(client: TestClient, *, role: str, name: str, content: str) -> dict[str, object]:
    response = client.post(
        "/api/files/upload",
        data={"role": role},
        files={"file": (name, content.encode("utf-8"), "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _start_task(authed: TestClient) -> tuple[str, dict[str, object]]:
    source = _upload_csv(
        authed,
        role="source",
        name="source.csv",
        content="物料编码,物料名称\nS001,精密电阻\nS002,陶瓷电容\n",
    )
    target = _upload_csv(
        authed,
        role="target",
        name="target.csv",
        content="集团码,物料名称\nG001,精密电阻\nG002,陶瓷电容\n",
    )
    catalog_response = authed.post(
        "/api/catalogs",
        json={
            "name": "时间语义测试目录",
            "source_file_id": target["file"]["file_id"],
            "group_code_column": "集团码",
        },
    )
    assert catalog_response.status_code == 200, catalog_response.text
    catalog = catalog_response.json()

    draft = authed.post("/api/task-drafts", json={}).json()
    saved_data = authed.put(
        f"/api/task-drafts/{draft['draft_id']}/data",
        json={
            "source_file_id": source["file"]["file_id"],
            "catalog_version_id": catalog["version_id"],
        },
    )
    assert saved_data.status_code == 200, saved_data.text
    config = {
        "source_id_column": "物料编码",
        "scope_mode": "GLOBAL",
        "rules": [
            {
                "id": "name",
                "source": {"fields": ["物料名称"]},
                "target": {"fields": ["物料名称"]},
                "matcher": "exact",
                "weight": 100,
            }
        ],
        "decision": {
            "success_threshold": 90,
            "review_enabled": True,
            "review_threshold": 50,
            "top_n": 2,
        },
        "retrieval": {"mode": "scan", "retrieval_top_k": 10, "oversample": 2},
        "advanced": {"scheme_display_name": "任务时间语义测试"},
    }
    saved_rules = authed.put(f"/api/task-drafts/{draft['draft_id']}/rules", json=config)
    assert saved_rules.status_code == 200, saved_rules.text

    started_response = authed.post(f"/api/task-drafts/{draft['draft_id']}/start")
    assert started_response.status_code == 202, started_response.text
    started = started_response.json()
    # Business start is captured when the user starts the run, not when a worker
    # later claims it.
    assert started["started_at"] == started["created_at"]
    assert started["compute_started_at"] is None
    assert started["compute_completed_at"] is None
    assert started["compute_duration_ms"] is None

    task_id = str(started["task_id"])
    deadline = time.time() + 10
    task = started
    while task["status"] not in {"COMPLETED", "FAILED"} and time.time() < deadline:
        time.sleep(0.02)
        response = authed.get(f"/api/tasks/{task_id}")
        assert response.status_code == 200, response.text
        task = response.json()
    assert task["status"] == "COMPLETED", task
    return task_id, task


def _duration_from_iso(started_at: str, completed_at: str) -> int:
    started = datetime.fromisoformat(started_at.replace("Z", "+00:00"))
    completed = datetime.fromisoformat(completed_at.replace("Z", "+00:00"))
    return int(round((completed - started).total_seconds() * 1000))


def test_new_task_records_precise_compute_window_and_all_apis_share_it(authed: TestClient) -> None:
    task_id, task = _start_task(authed)

    assert task["compute_started_at"]
    assert task["compute_completed_at"]
    expected = _duration_from_iso(str(task["compute_started_at"]), str(task["compute_completed_at"]))
    assert task["compute_duration_ms"] == expected
    assert task["compute_elapsed_ms"] == expected

    progress = authed.get(f"/api/tasks/{task_id}/progress").json()
    summary = authed.get(f"/api/tasks/{task_id}/workbench/summary").json()
    history = {row["task_id"]: row for row in authed.get("/api/tasks/review-history").json()}
    listed = {row["task_id"]: row for row in authed.get("/api/tasks").json()}

    for payload in (progress, summary, history[task_id], listed[task_id]):
        assert payload["started_at"] == task["started_at"]
        assert payload["compute_duration_ms"] == expected

    # The duration is based on the server-side automatic lifecycle, not page load.
    time.sleep(0.03)
    assert authed.get(f"/api/tasks/{task_id}").json()["compute_duration_ms"] == expected


def test_manual_wait_redecide_excel_import_and_result_generation_never_change_compute_duration(authed: TestClient) -> None:
    task_id, _ = _start_task(authed)
    meta = authed.app.state.meta
    matches = authed.app.state.matches

    # Pin a deterministic automatic compute window. Everything below represents
    # later human activity and must leave these timestamps untouched.
    business_start = "2026-09-18T09:30:21+08:00"
    compute_start = "2026-09-18T09:30:21+08:00"
    compute_end = "2026-09-18T09:30:43+08:00"
    with meta.connect() as connection:
        connection.execute(
            "UPDATE tasks SET started_at=? WHERE task_id=?",
            (business_start, task_id),
        )
        connection.execute(
            "UPDATE task_compute_lifecycle SET compute_started_at=?,compute_completed_at=? WHERE task_id=?",
            (compute_start, compute_end, task_id),
        )
    baseline = authed.get(f"/api/tasks/{task_id}").json()
    assert baseline["compute_duration_ms"] == 22_000

    # Threshold re-decision uses persisted scores/candidates only.
    matches.re_decide(task_id, 90, 50, mode="apply")
    assert authed.get(f"/api/tasks/{task_id}").json()["compute_duration_ms"] == 22_000

    with meta.connect() as connection:
        rows = connection.execute(
            "SELECT source_row_id FROM match_items WHERE task_id=? ORDER BY CAST(source_row_id AS INTEGER)",
            (task_id,),
        ).fetchall()
        assert len(rows) == 2
        source_ids = [str(row["source_row_id"]) for row in rows]
        connection.execute(
            "UPDATE tasks SET stage='REVIEW' WHERE task_id=?",
            (task_id,),
        )
        connection.execute(
            "UPDATE match_items SET current_status='REVIEW',final_group_code=NULL WHERE task_id=?",
            (task_id,),
        )
        candidate = connection.execute(
            """SELECT target_group_code FROM match_candidates
               WHERE task_id=? AND source_row_id=? AND rank=1""",
            (task_id, source_ids[0]),
        ).fetchone()
        assert candidate is not None

    # Online STEP3 confirmation.
    matches.confirm(task_id, source_ids[0], str(candidate["target_group_code"]), "一小时后人工确认", operator="reviewer")
    with meta.connect() as connection:
        connection.execute(
            "UPDATE reviews SET created_at=? WHERE task_id=? AND source_row_id=?",
            ("2026-09-18T10:30:43+08:00", task_id, source_ids[0]),
        )
        connection.execute(
            "UPDATE match_items SET updated_at=? WHERE task_id=? AND source_row_id=?",
            ("2026-09-18T10:30:43+08:00", task_id, source_ids[0]),
        )
    assert authed.get(f"/api/tasks/{task_id}").json()["compute_duration_ms"] == 22_000

    # Offline Excel confirmation for another row.
    manual = authed.get(f"/api/tasks/{task_id}/manual-review.xlsx")
    assert manual.status_code == 200, manual.text
    workbook = load_workbook(BytesIO(manual.content))
    sheet = workbook["人工匹配"]
    headers = {str(cell.value): cell.column for cell in sheet[1] if cell.value is not None}
    for row_no in range(2, sheet.max_row + 1):
        if str(sheet.cell(row_no, headers["__source_row_id"]).value) == source_ids[1]:
            sheet.cell(row_no, headers["人工选择"]).value = "候选1"
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    imported = authed.post(
        f"/api/tasks/{task_id}/manual-review/import",
        files={
            "file": (
                "manual.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["success_count"] >= 1
    assert authed.get(f"/api/tasks/{task_id}").json()["compute_duration_ms"] == 22_000

    # STEP4 result generation/export is audit activity, not compute performance.
    finalized = matches.finalize(task_id)
    assert finalized["result_file_id"]
    after_finalize = authed.get(f"/api/tasks/{task_id}").json()
    assert after_finalize["compute_duration_ms"] == 22_000
    assert after_finalize["compute_started_at"] == compute_start
    assert after_finalize["compute_completed_at"] == compute_end

    result_response = authed.get(f"/api/tasks/{task_id}/result")
    assert result_response.status_code == 200
    exported = load_workbook(BytesIO(result_response.content), data_only=True)
    summary = exported["匹配摘要"]
    values: dict[str, object] = {}
    for row_no in range(3, summary.max_row + 1):
        for key_col, value_col in ((1, 2), (3, 4)):
            key = summary.cell(row_no, key_col).value
            if key:
                values[str(key)] = summary.cell(row_no, value_col).value
    exported.close()

    assert "任务开始时间" in values
    assert values["自动计算耗时"] == "22 秒"
    assert "自动计算完成时间" in values
    assert "结果生成时间" in values
    assert "任务耗时" not in values
    assert "创建时间" not in values


def test_legacy_migration_preserves_reliable_compute_pair_and_never_uses_updated_time(tmp_path: Path) -> None:
    db_path = tmp_path / "legacy.db"
    connection = sqlite3.connect(db_path)
    connection.execute(
        """CREATE TABLE tasks(
          task_id TEXT PRIMARY KEY, name TEXT NOT NULL, source_file_id TEXT NOT NULL,
          catalog_version_id TEXT NOT NULL, profile_id TEXT, profile_version INTEGER,
          config_snapshot TEXT NOT NULL, config_sha256 TEXT NOT NULL,
          stage TEXT NOT NULL, status TEXT NOT NULL, progress REAL NOT NULL,
          processed_rows INTEGER NOT NULL, total_rows INTEGER NOT NULL,
          created_at TEXT NOT NULL, started_at TEXT, finished_at TEXT,
          error_code TEXT, error_message TEXT, result_file_id TEXT
        )"""
    )
    connection.execute(
        "INSERT INTO tasks VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            "legacy-good", "legacy", "source", "catalog", None, None, "{}", "sha",
            "REVIEW", "COMPLETED", 100.0, 1, 1,
            "2026-09-18T09:30:21+08:00",
            "2026-09-18T09:30:22+08:00",
            "2026-09-18T09:30:43+08:00",
            None, None, None,
        ),
    )
    connection.execute(
        "INSERT INTO tasks VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            "legacy-missing", "legacy", "source", "catalog", None, None, "{}", "sha",
            "REVIEW", "COMPLETED", 100.0, 1, 1,
            "2026-09-18T11:00:00+08:00",
            None,
            "2026-09-18T14:00:00+08:00",
            None, None, None,
        ),
    )
    connection.commit()
    connection.close()

    repo = MetadataRepository(db_path)
    with repo.connect() as connection:
        good = dict(connection.execute("SELECT * FROM tasks WHERE task_id='legacy-good'").fetchone())
        missing = dict(connection.execute("SELECT * FROM tasks WHERE task_id='legacy-missing'").fetchone())
        lifecycle = connection.execute(
            "SELECT * FROM task_compute_lifecycle WHERE task_id='legacy-good'"
        ).fetchone()
        missing_lifecycle = connection.execute(
            "SELECT * FROM task_compute_lifecycle WHERE task_id='legacy-missing'"
        ).fetchone()

    # Business start is the original click-to-start timestamp.
    assert good["started_at"] == good["created_at"] == "2026-09-18T09:30:21+08:00"
    assert lifecycle is not None
    assert lifecycle["compute_started_at"] == "2026-09-18T09:30:22+08:00"
    assert lifecycle["compute_completed_at"] == "2026-09-18T09:30:43+08:00"
    assert task_time_fields(repo, good)["compute_duration_ms"] == 21_000

    # An incomplete historical pair is intentionally not guessed from
    # created_at/finished_at/updated_at-like values.
    assert missing_lifecycle is None
    missing["updated_at"] = "2026-09-19T14:00:00+08:00"
    timing = task_time_fields(repo, missing)
    assert timing["compute_duration_ms"] is None
    assert format_duration_ms(timing["compute_duration_ms"]) == "暂无准确记录"


def test_duration_formatter_keeps_seconds_for_business_readability() -> None:
    assert format_duration_ms(8_000) == "8 秒"
    assert format_duration_ms(92_000) == "1 分 32 秒"
    assert format_duration_ms(3_912_000) == "1 小时 5 分 12 秒"
    assert format_duration_ms(None) == "暂无准确记录"
