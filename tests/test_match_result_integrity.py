from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from material_matcher.domain.models import MatchingConfig
from material_matcher.matching.scorer import decide_status, score_candidate
from material_matcher.services.match_service import MatchService
from material_matcher.settings import Settings
from material_matcher.storage.files import FileRepository
from material_matcher.storage.metadata import MetadataRepository


def _service(tmp_path: Path) -> tuple[MatchService, MetadataRepository, FileRepository]:
    settings = Settings(
        data_dir=tmp_path / "data",
        config_dir=tmp_path / "etc",
        log_dir=tmp_path / "log",
        admin_password="x",
    )
    settings.ensure_dirs()
    meta = MetadataRepository(settings.data_dir / "meta" / "material_matcher.db")
    files = FileRepository(settings.data_dir, meta)
    return MatchService(meta, files, settings), meta, files


def _config() -> MatchingConfig:
    return MatchingConfig.model_validate(
        {
            "source_id_column": "物料号",
            "rules": [
                {
                    "id": "optional-model",
                    "source": {"fields": ["型号"]},
                    "target": {"fields": ["型号"]},
                    "matcher": "exact",
                    "weight": 70,
                },
                {
                    "id": "name",
                    "source": {"fields": ["名称"]},
                    "target": {"fields": ["名称"]},
                    "matcher": "fuzzy",
                    "weight": 30,
                },
            ],
            "decision": {
                "success_threshold": 78,
                "review_enabled": True,
                "review_threshold": 50,
                "top_n": 5,
            },
        }
    )


def _seed_task(meta: MetadataRepository, config: MatchingConfig, *, task_id: str = "t-score", rows: int = 150) -> None:
    now = "2026-09-16T00:00:00+08:00"
    with meta.connect() as connection:
        connection.execute(
            """INSERT INTO tasks(
                task_id,name,source_file_id,catalog_version_id,config_snapshot,config_sha256,
                stage,status,progress,processed_rows,total_rows,created_at,started_at,finished_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                task_id,
                "score-integrity",
                "source-file",
                "catalog-v1",
                json.dumps(config.model_dump(mode="json"), ensure_ascii=False),
                "audit-sha-kept-in-db",
                "RESULT",
                "COMPLETED",
                100.0,
                rows,
                rows,
                now,
                now,
                now,
            ),
        )


def test_original_excel_row_number_respects_non_first_header_and_blank_rows(tmp_path: Path) -> None:
    from material_matcher.ingestion.reader import iter_tabular_rows_with_position

    path = tmp_path / "positioned.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "说明"
    sheet["A4"] = "物料号"
    sheet["B4"] = "名称"
    sheet["A5"] = "S-001"
    sheet["B5"] = "电阻"
    # Row 6 is deliberately blank. The next business row must still be row 7.
    sheet["A7"] = "S-002"
    sheet["B7"] = "电容"
    workbook.save(path)

    positioned = list(iter_tabular_rows_with_position(path, header_row=4))
    assert [row_number for row_number, _ in positioned] == [5, 7]
    assert [row["物料号"] for _, row in positioned] == ["S-001", "S-002"]


def test_low_scores_stay_real_across_score_decision_persistence_and_excel(tmp_path: Path) -> None:
    service, meta, files = _service(tmp_path)
    config = _config()

    # The high-weight model rule is missing on the source side. It is not scored as
    # zero; the remaining comparable field is re-normalized. Therefore this low
    # score is a genuine fuzzy mismatch rather than a 0..1/0..100 scale defect.
    source = {"物料号": "S-1", "名称": "ABC", "型号": None}
    target = {"集团码": "G-1", "名称": "AXYZ", "型号": "M-1"}
    candidate_score = score_candidate(source, target, config)
    assert 10.0 < candidate_score.display_score < 40.0
    assert candidate_score.display_score == round(candidate_score.raw_score * 100.0, 4)
    assert len(candidate_score.field_scores) == 1
    assert candidate_score.display_score == round(candidate_score.field_scores[0].score * 100.0, 4)
    assert decide_status(candidate_score.display_score, config) == "UNMATCHED"

    from material_matcher.matching.engine import CandidateResult, RowResult

    rows: list[RowResult] = []
    for index in range(1, 151):
        candidate = CandidateResult(
            rank=1,
            group_code="G-1",
            score=candidate_score.display_score,
            raw_score=candidate_score.raw_score,
            critical_conflict=False,
            target_payload=target,
            field_scores=[item.__dict__ for item in candidate_score.field_scores],
            target_row_number=7,
        )
        rows.append(
            RowResult(
                source_row_id=str(index),
                source_id=f"S-{index}",
                source_payload={**source, "物料号": f"S-{index}"},
                status="UNMATCHED",
                final_group_code=None,
                first_score=candidate_score.display_score,
                second_score=0.0,
                score_gap=candidate_score.display_score,
                critical_conflict=False,
                candidates=[candidate],
                source_row_number=index + 4,
            )
        )

    _seed_task(meta, config)
    service._persist_rows("t-score", rows)
    assert service.summary("t-score")["unmatched"] == 150
    with meta.connect() as connection:
        stored = connection.execute(
            "SELECT source_row_number,top1_score,current_status FROM match_items WHERE task_id='t-score' AND source_row_id='1'"
        ).fetchone()
        stored_candidate = connection.execute(
            "SELECT target_row_number,score FROM match_candidates WHERE task_id='t-score' AND source_row_id='1' AND rank=1"
        ).fetchone()
    assert int(stored["source_row_number"]) == 5
    assert float(stored["top1_score"]) == candidate_score.display_score
    assert str(stored["current_status"]) == "UNMATCHED"
    assert int(stored_candidate["target_row_number"]) == 7
    assert float(stored_candidate["score"]) == candidate_score.display_score

    finalized = service.finalize("t-score")
    result_record = files.get(str(finalized["result_file_id"]))
    exported = load_workbook(Path(str(result_record["stored_path"])), data_only=True)
    try:
        summary_values = [exported["匹配摘要"].cell(row=row, column=1).value for row in range(1, exported["匹配摘要"].max_row + 1)]
        assert "配置SHA256" not in summary_values

        result = exported["匹配结果"]
        assert result.cell(2, 1).value == 5
        assert result.cell(2, 2).value == "S-1"
        assert result.cell(2, 7).value == 7
        assert float(result.cell(2, 8).value) == candidate_score.display_score

        topn = exported["TopN候选"]
        assert topn.cell(2, 1).value == 5
        assert topn.cell(2, 2).value == "S-1"
        assert topn.cell(2, 4).value == 7
        assert topn.cell(2, 5).value == "G-1"
        assert float(topn.cell(2, 6).value) == candidate_score.display_score
    finally:
        exported.close()

    with meta.connect() as connection:
        task = connection.execute("SELECT config_sha256 FROM tasks WHERE task_id='t-score'").fetchone()
    assert str(task["config_sha256"]) == "audit-sha-kept-in-db"


def _insert_redecide_row(
    meta: MetadataRepository,
    *,
    source_row_id: str,
    status: str,
    score: float,
    second_score: float,
    critical_conflict: bool = False,
    second_group: str | None = None,
) -> None:
    now = "2026-09-16T00:00:00+08:00"
    with meta.connect() as connection:
        connection.execute(
            """INSERT INTO match_items(
                task_id,source_row_id,source_row_number,source_id,source_payload,original_status,current_status,
                top1_group_code,top1_score,second_score,score_gap,critical_conflict,final_group_code,created_at,updated_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                "t-red",
                source_row_id,
                10 + int(source_row_id[1:]),
                source_row_id,
                "{}",
                status,
                status,
                "G1",
                score,
                second_score,
                score - second_score,
                1 if critical_conflict else 0,
                "G1" if status in {"MATCHED", "CONFIRMED"} else None,
                now,
                now,
            ),
        )
        connection.execute(
            """INSERT INTO match_candidates(
                task_id,source_row_id,rank,target_row_number,target_group_code,target_payload,score,field_scores,critical_conflict
            ) VALUES(?,?,?,?,?,?,?,?,?)""",
            ("t-red", source_row_id, 1, 20, "G1", "{}", score, "[]", 1 if critical_conflict else 0),
        )
        if second_group is not None:
            connection.execute(
                """INSERT INTO match_candidates(
                    task_id,source_row_id,rank,target_row_number,target_group_code,target_payload,score,field_scores,critical_conflict
                ) VALUES(?,?,?,?,?,?,?,?,?)""",
                ("t-red", source_row_id, 2, 21, second_group, "{}", second_score, "[]", 0),
            )


def test_redecide_preview_is_side_effect_free_and_apply_keeps_protections(tmp_path: Path) -> None:
    service, meta, _ = _service(tmp_path)
    _seed_task(meta, _config(), task_id="t-red", rows=6)
    with meta.connect() as connection:
        connection.execute("UPDATE tasks SET stage='REVIEW' WHERE task_id='t-red'")

    _insert_redecide_row(meta, source_row_id="r1", status="REVIEW", score=90.0, second_score=80.0)
    _insert_redecide_row(meta, source_row_id="r2", status="REVIEW", score=90.0, second_score=80.0, critical_conflict=True)
    _insert_redecide_row(meta, source_row_id="r3", status="REVIEW", score=90.0, second_score=90.0, second_group="G2")
    _insert_redecide_row(meta, source_row_id="r4", status="UNMATCHED", score=40.0, second_score=20.0)
    _insert_redecide_row(meta, source_row_id="r5", status="UNMATCHED", score=95.0, second_score=80.0)
    _insert_redecide_row(meta, source_row_id="r6", status="CONFIRMED", score=95.0, second_score=80.0)

    now = "2026-09-16T00:00:00+08:00"
    with meta.connect() as connection:
        connection.execute(
            "INSERT INTO reviews VALUES(?,?,?,?,?,?,?,?,?)",
            ("review-r5", "t-red", "r5", "REVIEW", None, "REJECT_ALL", "reviewer", "人工未匹配", now),
        )
        connection.execute(
            "INSERT INTO reviews VALUES(?,?,?,?,?,?,?,?,?)",
            ("review-r6", "t-red", "r6", "REVIEW", "G1", "CONFIRM_CANDIDATE", "reviewer", "人工确认", now),
        )

    preview = service.re_decide("t-red", 70, 55, mode="preview")
    assert preview["before"] == {"matched": 0, "review": 3, "unmatched": 1}
    assert preview["after"] == {"matched": 1, "review": 2, "unmatched": 1}
    assert preview["matched_delta"] == 1
    assert preview["review_reduction"] == 1
    assert preview["unmatched_delta"] == 0
    assert preview["affected_rows"] == 1
    assert preview["critical_conflict_protected"] == 1
    assert preview["ambiguity_protected"] == 1

    with meta.connect() as connection:
        before_apply = {
            str(row["source_row_id"]): str(row["current_status"])
            for row in connection.execute("SELECT source_row_id,current_status FROM match_items WHERE task_id='t-red'").fetchall()
        }
    assert before_apply == {"r1": "REVIEW", "r2": "REVIEW", "r3": "REVIEW", "r4": "UNMATCHED", "r5": "UNMATCHED", "r6": "CONFIRMED"}

    applied = service.re_decide("t-red", 70, 55, mode="apply")
    assert applied["after"] == preview["after"]
    with meta.connect() as connection:
        after_apply = {
            str(row["source_row_id"]): str(row["current_status"])
            for row in connection.execute("SELECT source_row_id,current_status FROM match_items WHERE task_id='t-red'").fetchall()
        }
        audit = connection.execute(
            "SELECT payload FROM audit_events WHERE entity_type='task' AND entity_id='t-red' AND action='REDECIDE_THRESHOLDS'"
        ).fetchone()
    assert after_apply == {"r1": "MATCHED", "r2": "REVIEW", "r3": "REVIEW", "r4": "UNMATCHED", "r5": "UNMATCHED", "r6": "CONFIRMED"}
    assert json.loads(str(audit["payload"]))["success_threshold"] == 70
