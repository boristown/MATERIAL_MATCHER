from __future__ import annotations

import io
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _upload(client: TestClient, role: str, name: str, content: str):
    return client.post("/api/files/upload", data={"role": role}, files={"file": (name, content.encode("utf-8"), "text/csv")}).json()


def _prepare_minimal_draft(authed: TestClient, source_content: str, tag: str) -> str:
    source = _upload(authed, "source", f"{tag}-src.csv", source_content)["file"]["file_id"]
    target = _upload(authed, "target", f"{tag}-tgt.csv", "集团码,物料名称,型号\nG001,测试电机,X1\n")["file"]["file_id"]
    catalog = authed.post("/api/catalogs", json={"name": f"{tag}-目录", "source_file_id": target, "group_code_column": "集团码"}).json()
    created = authed.post("/api/task-drafts", json={}).json()
    draft_id = created["draft_id"]
    config = {
        "source_id_column": "物料编码",
        "scope_mode": "GLOBAL",
        "rules": [{
            "id": "nm",
            "source": {"fields": ["物料名称"], "combine": "concat", "separator": " ", "pipeline": []},
            "target": {"fields": ["物料名称"], "combine": "concat", "separator": " ", "pipeline": []},
            "matcher": "fuzzy", "weight": 100, "critical": False, "matcher_options": {},
        }],
        "decision": {"success_threshold": 60, "review_enabled": True, "review_threshold": 40, "top_n": 5},
        "retrieval": {"mode": "scan", "retrieval_top_k": 20, "oversample": 2},
        "advanced": {"workspace_target": {"file_id": target, "group_code_column": "集团码"}},
    }
    assert authed.put(f"/api/task-drafts/{draft_id}/data", json={"source_file_id": source, "catalog_version_id": catalog["version_id"]}).status_code == 200
    assert authed.put(f"/api/task-drafts/{draft_id}/rules", json=config).status_code == 200
    return draft_id


def _start_minimal_task(authed: TestClient, source_content: str, tag: str) -> str:
    draft_id = _prepare_minimal_draft(authed, source_content, tag)
    started = authed.post(f"/api/task-drafts/{draft_id}/start")
    assert started.status_code in (200, 202), started.text
    return started.json()["task_id"]


def _wait(authed: TestClient, task_id: str):
    from time import sleep
    for _ in range(120):
        progress = authed.get(f"/api/tasks/{task_id}/progress").json()
        if progress.get("status") in {"COMPLETED", "FAILED"}:
            return progress
        sleep(0.25)
    raise AssertionError("task did not finish")


SOURCE = "物料编码,物料名称,型号\nM001,测试电机,X1\n"


def test_draft_and_start_no_longer_take_a_business_name(authed: TestClient) -> None:
    created = authed.post("/api/task-drafts", json={})
    assert created.status_code == 200
    draft = created.json()
    assert str(draft["name"]).startswith("run-")

    legacy = authed.post("/api/task-drafts", json={"name": "客户甲九月任务"})
    assert legacy.status_code == 200
    assert legacy.json()["name"] == draft["name"] or not str(legacy.json()["name"]).endswith("九月任务")
    assert str(legacy.json()["name"]).startswith("run-")

    draft_id = legacy.json()["draft_id"]
    patched = authed.patch(f"/api/task-drafts/{draft_id}", json={"name": "改名试试"})
    assert patched.status_code == 200
    assert patched.json()["name"] == legacy.json()["name"]

    task_id = _start_minimal_task(authed, SOURCE, "noname")
    assert _wait(authed, task_id)["status"] == "COMPLETED"
    task = authed.get(f"/api/tasks/{task_id}").json()
    assert str(task["name"]).startswith("run-")
    assert task["scheme_name"] == "未命名方案"
    assert task["run_number"] == 1


def test_internal_name_smoke_never_leaks_to_business_surfaces(authed: TestClient, tmp_path: Path) -> None:
    task_id = _start_minimal_task(authed, SOURCE, "leak")
    assert _wait(authed, task_id)["status"] == "COMPLETED"
    finalized = authed.post(f"/api/tasks/{task_id}/finalize", json={"allow_unresolved_review": True})
    assert finalized.status_code == 200, finalized.text
    task = authed.get(f"/api/tasks/{task_id}").json()
    files = authed.get("/api/files").json()
    result_record = next(item for item in files if item["file_id"] == task["result_file_id"])
    assert str(result_record["original_name"]).startswith("物料集团码匹配结果_未命名方案_")

    workbook = load_workbook(io.BytesIO(authed.get(f"/api/tasks/{task_id}/result").content))
    summary = workbook["匹配摘要"]
    assert summary.cell(3, 1).value == "方案名称"
    all_text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "任务名称" not in all_text
    assert "SMOKE-1.1.5" not in all_text

    manual = authed.get(f"/api/tasks/{task_id}/manual-review.xlsx")
    assert manual.status_code == 200
    assert manual.headers["content-disposition"].startswith("attachment; filename*=UTF-8''%E4%BA%BA%E5%B7%A5%E5%8C%B9%E9%85%8D_")
    manual_book = load_workbook(io.BytesIO(manual.content))
    guide = manual_book["填写说明"]
    assert guide.cell(2, 1).value == "方案名称"
    manual_text = "\n".join(str(cell.value) for sheet in manual_book.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "任务名称" not in manual_text


def test_history_displays_frozen_scheme_even_after_internal_name_rewrite(authed: TestClient) -> None:
    profile = authed.post("/api/profiles", json={"name": "电子元器件集团码匹配方案"}).json()
    profile_id = profile["profile_id"]
    document = {
        "source_id_column": "物料编码",
        "scope_mode": "GLOBAL",
        "rules": [{
            "id": "nm",
            "source": {"fields": ["物料名称"], "combine": "concat", "separator": " ", "pipeline": []},
            "target": {"fields": ["物料名称"], "combine": "concat", "separator": " ", "pipeline": []},
            "matcher": "fuzzy", "weight": 100, "critical": False, "matcher_options": {},
        }],
        "decision": {"success_threshold": 60, "review_enabled": True, "review_threshold": 40, "top_n": 5},
        "retrieval": {"mode": "scan", "retrieval_top_k": 20, "oversample": 2},
    }
    assert authed.put(f"/api/profiles/{profile_id}/draft", json=document).status_code == 200
    assert authed.post(f"/api/profiles/{profile_id}/validate").status_code == 200
    assert authed.post(f"/api/profiles/{profile_id}/publish").status_code == 200

    source = _upload(authed, "source", "tpl-src.csv", SOURCE)["file"]["file_id"]
    target = _upload(authed, "target", "tpl-tgt.csv", "集团码,物料名称,型号\nG001,测试电机,X1\n")["file"]["file_id"]
    catalog = authed.post("/api/catalogs", json={"name": "模板目录", "source_file_id": target, "group_code_column": "集团码"}).json()
    draft = authed.post("/api/task-drafts", json={"name": "ignored"}).json()
    assert authed.patch(f"/api/task-drafts/{draft['draft_id']}", json={
        "source_file_id": source, "catalog_version_id": catalog["version_id"],
        "template_profile_id": profile_id, "template_profile_version": 1, "config_document": document,
    }).status_code == 200
    task_id = authed.post(f"/api/task-drafts/{draft['draft_id']}/start").json()["task_id"]
    assert _wait(authed, task_id)["status"] == "COMPLETED"

    # An operator (or legacy tooling) rewrites the internal column with a SMOKE name.
    with authed.app.state.meta.connect() as connection:
        connection.execute("UPDATE tasks SET name='SMOKE-1.1.5' WHERE task_id=?", (task_id,))

    listed = next(item for item in authed.get("/api/tasks").json() if item["task_id"] == task_id)
    assert listed["name"] == "SMOKE-1.1.5"
    assert listed["scheme_name"] == "电子元器件集团码匹配方案"
    assert str(listed["run_number"]) == "1"

    authed.patch(f"/api/profiles/{profile_id}", json={"name": "后来改过的方案名"})
    after_rename = next(item for item in authed.get("/api/tasks").json() if item["task_id"] == task_id)
    assert after_rename["scheme_name"] == "电子元器件集团码匹配方案"

    finalized = authed.post(f"/api/tasks/{task_id}/finalize", json={"allow_unresolved_review": True})
    assert finalized.status_code == 200
    workbook = load_workbook(io.BytesIO(authed.get(f"/api/tasks/{task_id}/result").content))
    assert workbook["匹配摘要"].cell(3, 1).value == "方案名称"
    assert workbook["匹配摘要"].cell(3, 2).value == "电子元器件集团码匹配方案"
    all_text = "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert "SMOKE-1.1.5" not in all_text and "后来改过的方案名" not in all_text


def test_same_config_runs_get_increasing_run_numbers(authed: TestClient) -> None:
    draft_id = _prepare_minimal_draft(authed, SOURCE, "rn")
    first = authed.post(f"/api/task-drafts/{draft_id}/start").json()["task_id"]
    assert _wait(authed, first)["status"] == "COMPLETED"
    second = authed.post(f"/api/task-drafts/{draft_id}/start").json()["task_id"]
    assert _wait(authed, second)["status"] == "COMPLETED"
    numbers = {item["task_id"]: item["run_number"] for item in authed.get("/api/tasks").json() if item["task_id"] in (first, second)}
    assert sorted(numbers.values()) == [1, 2]


def test_safe_business_filename_sanitiser() -> None:
    from material_matcher.services.task_service import safe_business_filename

    assert safe_business_filename('A/B:C*D?E"F<G>H|I') == "ABCDEFGHI"
    assert safe_business_filename("   ") == "未命名方案"
    assert safe_business_filename(None) == "未命名方案"
    assert len(safe_business_filename("很长的方案名称" * 20)) <= 60
    assert safe_business_filename("..hidden..") == "hidden"
