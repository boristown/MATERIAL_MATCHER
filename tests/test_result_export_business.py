from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from material_matcher.matching.engine import CandidateResult, RowResult
from material_matcher.services.match_service import MatchService
from material_matcher.settings import Settings
from material_matcher.storage.files import FileRepository
from material_matcher.storage.metadata import MetadataRepository


EXACT = "00E2F0D9"
PARTIAL = "00FFF2CC"
DIFFERENT = "00FCE4D6"
MISSING = "00E7E6E6"


def _service(tmp_path: Path) -> tuple[MatchService, MetadataRepository, FileRepository]:
    settings = Settings(
        data_dir=tmp_path / "data",
        config_dir=tmp_path / "etc",
        log_dir=tmp_path / "log",
        admin_password="x",
    )
    settings.ensure_dirs()
    meta = MetadataRepository(settings.data_dir / "meta" / "material_matcher.db")
    files = FileRepository(settings.data_dir, meta)
    return MatchService(meta, files, settings), meta, files


def _config() -> dict[str, object]:
    return {
        "source_id_column": "物料号",
        "rules": [
            {
                "id": "name",
                "source": {"fields": ["名称"]},
                "target": {"fields": ["名称"]},
                "matcher": "exact",
                "weight": 40,
            },
            {
                "id": "model",
                "source": {"fields": ["型号"]},
                "target": {"fields": ["型号"]},
                "matcher": "fuzzy",
                "weight": 40,
            },
            {
                "id": "category",
                "source": {"fields": ["类别"]},
                "target": {"fields": ["类别"]},
                "matcher": "exact",
                "weight": 20,
            },
        ],
        "decision": {"success_threshold": 78, "review_enabled": True, "review_threshold": 50, "top_n": 5},
    }


def _seed_task(meta: MetadataRepository) -> None:
    now = "2026-09-16T09:00:00+08:00"
    with meta.connect() as connection:
        connection.execute(
            """INSERT INTO tasks(
                task_id,name,source_file_id,catalog_version_id,config_snapshot,config_sha256,
                stage,status,progress,processed_rows,total_rows,created_at,started_at,finished_at
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                "business-result",
                "正式物料匹配",
                "source-file",
                "catalog-v1",
                json.dumps(_config(), ensure_ascii=False),
                "kept-only-in-database",
                "RESULT",
                "COMPLETED",
                100.0,
                3,
                3,
                now,
                "2026-09-16T09:01:00+08:00",
                "2026-09-16T09:05:00+08:00",
            ),
        )
        connection.execute(
            "INSERT INTO audit_events VALUES(?,?,?,?,?,?)",
            ("task-created", "task", "business-result", "TASK_CREATED", json.dumps({"operator": "creator01"}), now),
        )
        connection.execute(
            "INSERT INTO audit_events VALUES(?,?,?,?,?,?)",
            ("task-started", "task", "business-result", "TASK_STARTED", json.dumps({"operator": "starter02"}), "2026-09-16T09:01:00+08:00"),
        )


def _candidate(rank: int, code: str, row: int, score: float) -> CandidateResult:
    target = {"集团码": code, "名称": "六角螺栓", "型号": "M8X22", "类别": "B", "备注": None}
    field_scores = [
        {"rule_id": "name", "score": 1.0, "weight": 40, "source_value": "六角螺栓", "target_value": "六角螺栓"},
        {"rule_id": "model", "score": 0.8, "weight": 40, "source_value": "M8X20", "target_value": "M8X22"},
        {"rule_id": "category", "score": 0.0, "weight": 20, "source_value": "A", "target_value": "B"},
    ]
    return CandidateResult(
        rank=rank,
        group_code=code,
        score=score,
        raw_score=score / 100.0,
        critical_conflict=False,
        target_payload=target,
        field_scores=field_scores,
        target_row_number=row,
    )


def _row(source_row_id: str, source_row_number: int, status: str, code: str, target_row: int, score: float) -> RowResult:
    candidate = _candidate(1, code, target_row, score)
    return RowResult(
        source_row_id=source_row_id,
        source_id=f"00012{source_row_id}",
        source_payload={"物料号": f"00012{source_row_id}", "名称": "六角螺栓", "型号": "M8X20", "类别": "A", "备注": None},
        status=status,
        final_group_code=code if status == "MATCHED" else None,
        first_score=score,
        second_score=0.0,
        score_gap=score,
        critical_conflict=False,
        candidates=[candidate],
        source_row_number=source_row_number,
    )


def _section_range(sheet, title: str) -> tuple[int, int]:
    for merged in sheet.merged_cells.ranges:
        if sheet.cell(merged.min_row, merged.min_col).value == title:
            return merged.min_col, merged.max_col
    raise AssertionError(f"section {title!r} not found")


def _header_col(sheet, title: str, section: str) -> int:
    start, end = _section_range(sheet, section)
    for column in range(start, end + 1):
        if sheet.cell(2, column).value == title:
            return column
    raise AssertionError(f"header {title!r} not found in {section!r}")


def test_final_excel_is_business_complete_colored_and_auditable(tmp_path: Path) -> None:
    service, meta, files = _service(tmp_path)
    _seed_task(meta)
    service._persist_rows(
        "business-result",
        [
            _row("1", 8, "MATCHED", "0000456", 34567, 88.0),
            _row("2", 9, "REVIEW", "0000789", 34568, 82.0),
            _row("3", 10, "REVIEW", "0000999", 34569, 61.0),
        ],
    )
    service.confirm("business-result", "2", "0000789", "确认集团码", operator="alice")
    service.reject("business-result", "3", "无合适候选", operator="bob")

    business_summary = service.summary("business-result")
    assert business_summary["created_by"] == "creator01"
    assert business_summary["started_by"] == "starter02"
    preview = {str(row["source_row_id"]): row for row in business_summary["preview_rows"]}
    assert preview["1"]["source_row_number"] == 8
    assert preview["1"]["target_row_number"] == 34567
    assert preview["1"]["match_method"] == "自动匹配"
    assert preview["2"]["last_operator"] == "alice"
    assert preview["2"]["match_method"] == "人工匹配"
    assert preview["3"]["last_operator"] == "bob"
    assert preview["3"]["match_method"] == "人工标记未匹配"

    finalized = service.finalize("business-result")
    result_record = files.get(str(finalized["result_file_id"]))
    workbook = load_workbook(Path(str(result_record["stored_path"]),), data_only=True)
    try:
        assert workbook.sheetnames == ["匹配摘要", "最终匹配结果", "Top5候选", "人工操作记录", "未匹配清单"]

        banned = ("sha256", "config hash", "internal id", "vector index", "embedding", "bbq", "rerank", "provider", "critical conflict", "关键字段冲突")
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        lowered = cell.value.lower()
                        assert not any(term in lowered for term in banned)

        result = workbook["最终匹配结果"]
        assert result.freeze_panes == "A3"
        assert result.auto_filter.ref is not None
        assert result.auto_filter.ref.startswith("A2:")
        assert int(result.auto_filter.ref.split(":")[1][1:]) >= result.max_row

        source_code_col = _header_col(result, "物料号", "源数据区")
        source_name_col = _header_col(result, "名称", "源数据区")
        source_model_col = _header_col(result, "型号", "源数据区")
        source_category_col = _header_col(result, "类别", "源数据区")
        source_note_col = _header_col(result, "备注", "源数据区")
        status_col = _header_col(result, "匹配状态", "匹配结果区")
        code_col = _header_col(result, "最终集团码", "匹配结果区")
        method_col = _header_col(result, "匹配方式", "匹配结果区")
        operator_col = _header_col(result, "操作账号", "匹配结果区")
        target_row_col = _header_col(result, "目标表原始行号", "目标数据区")
        target_code_col = _header_col(result, "集团码", "目标数据区")
        target_name_col = _header_col(result, "名称", "目标数据区")
        target_model_col = _header_col(result, "型号", "目标数据区")
        target_category_col = _header_col(result, "类别", "目标数据区")
        target_note_col = _header_col(result, "备注", "目标数据区")

        assert result.cell(3, 1).value == 8
        assert result.cell(3, source_code_col).value == "000121"
        assert result.cell(3, source_code_col).number_format == "@"
        assert result.cell(3, code_col).value == "0000456"
        assert result.cell(3, code_col).number_format == "@"
        assert result.cell(3, target_row_col).value == 34567
        assert result.cell(3, target_code_col).value == "0000456"
        assert result.cell(3, target_code_col).number_format == "@"
        assert result.cell(3, status_col).value == "自动匹配"
        assert result.cell(3, method_col).value == "自动匹配"

        assert result.cell(3, source_name_col).fill.fgColor.rgb == EXACT
        assert result.cell(3, target_name_col).fill.fgColor.rgb == EXACT
        assert result.cell(3, source_model_col).fill.fgColor.rgb == PARTIAL
        assert result.cell(3, target_model_col).fill.fgColor.rgb == PARTIAL
        assert result.cell(3, source_category_col).fill.fgColor.rgb == DIFFERENT
        assert result.cell(3, target_category_col).fill.fgColor.rgb == DIFFERENT
        assert result.cell(3, source_note_col).fill.fgColor.rgb == MISSING
        assert result.cell(3, target_note_col).fill.fgColor.rgb == MISSING

        assert result.cell(4, status_col).value == "人工匹配"
        assert result.cell(4, method_col).value == "人工匹配"
        assert result.cell(4, operator_col).value == "alice"
        assert result.cell(5, status_col).value == "未匹配"
        assert result.cell(5, method_col).value == "人工标记未匹配"
        assert result.cell(5, operator_col).value == "bob"
        assert result.cell(5, code_col).value == ""
        assert result.cell(5, target_row_col).value is None

        top5 = workbook["Top5候选"]
        assert top5.freeze_panes == "A3"
        assert _header_col(top5, "候选排名", "候选结果区") > 0
        assert _header_col(top5, "目标表原始行号", "候选结果区") > 0

        audit = workbook["人工操作记录"]
        action_col = _header_col(audit, "操作", "人工操作")
        audit_target_col = _header_col(audit, "目标表原始行号", "人工操作")
        audit_code_col = _header_col(audit, "集团码", "人工操作")
        audit_operator_col = _header_col(audit, "操作账号", "人工操作")
        audit_time_col = _header_col(audit, "操作时间", "人工操作")
        assert audit.cell(3, action_col).value == "匹配"
        assert audit.cell(3, audit_target_col).value == 34568
        assert audit.cell(3, audit_code_col).value == "0000789"
        assert audit.cell(3, audit_operator_col).value == "alice"
        assert audit.cell(3, audit_time_col).value is not None
        assert audit.cell(4, action_col).value == "标记未匹配"
        assert audit.cell(4, audit_operator_col).value == "bob"

        unmatched = workbook["未匹配清单"]
        assert unmatched.max_row == 3
        assert unmatched.cell(3, 1).value == 10
    finally:
        workbook.close()
