from pathlib import Path

from openpyxl import Workbook, load_workbook

from material_matcher.matching import run_matching, write_result_xlsx


def save_book(path: Path, headers: list[str], rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_matching_respects_threshold_and_topn(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    target = tmp_path / "target.xlsx"
    save_book(source, ["物料号", "名称", "型号"], [["0001", "环氧玻璃布板", "3240"], ["0002", "完全不同", "XYZ"]])
    save_book(target, ["集团码", "名称", "牌号"], [["G100", "环氧玻璃布板", "3240"], ["G200", "酚醛板", "3021"]])

    config = {
        "source_sheet": "数据",
        "source_header_row": 1,
        "target_sheet": "数据",
        "target_header_row": 1,
        "source_id_column": "物料号",
        "group_code_column": "集团码",
        "mappings": [
            {"source_header": "名称", "target_header": "名称", "weight": 0.4, "method": "hybrid", "name": "名称"},
            {"source_header": "型号", "target_header": "牌号", "weight": 0.6, "method": "hybrid", "name": "型号", "critical": True},
        ],
        "threshold": 0.85,
        "review_threshold": 0.70,
        "top_n": 2,
        "candidate_limit": 20,
    }
    result = run_matching(source, target, config)
    assert result["summary"]["source_rows"] == 2
    assert result["rows"][0]["status"] == "MATCHED"
    assert result["rows"][0]["matched_group_code"] == "G100"
    assert result["rows"][1]["matched_group_code"] == ""
    assert len(result["rows"][0]["candidates"]) == 2

    output = tmp_path / "result.xlsx"
    write_result_xlsx(result, output)
    wb = load_workbook(output, read_only=True)
    assert "匹配结果" in wb.sheetnames
    assert "TopN" in wb.sheetnames
