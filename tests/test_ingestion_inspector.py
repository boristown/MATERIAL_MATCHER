from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from material_matcher.ingestion.inspector import inspect_tabular_file
from material_matcher.ingestion.reader import detect_layout, iter_tabular_rows
from material_matcher.matching.engine import load_target_rows_with_position
from material_matcher.services.match_service import MatchService


def _write_csv(path: Path, rows: int) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        stream.write("物料编码,物料描述\n")
        for index in range(rows):
            stream.write(f"{index},六角螺栓 M{index}x30\n")


def _write_inflated_xlsx(
    path: Path,
    *,
    rows: int = 100,
    header_row: int = 1,
    sheet_name: str = "Data",
    phantom_row: int = 100_000,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    if header_row > 1:
        worksheet.cell(1, 1, "导入说明")
    worksheet.cell(header_row, 1, "物料编码")
    worksheet.cell(header_row, 2, "物料描述")
    for index in range(rows):
        worksheet.cell(header_row + index + 1, 1, f"M{index + 1:04d}")
        worksheet.cell(header_row + index + 1, 2, f"测试物料 {index + 1}")
    # A style-only remote cell persists a large worksheet dimension without
    # containing business data, reproducing the Excel used-range residue seen
    # in production files.
    worksheet.cell(phantom_row, 1).number_format = "@"
    workbook.save(path)


def _assert_inflated_dimension(path: Path, minimum: int = 100_000) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.active.max_row >= minimum
    finally:
        workbook.close()


def test_csv_row_count_estimate_is_not_capped_by_sample_window(tmp_path: Path) -> None:
    path = tmp_path / "big.csv"
    _write_csv(path, 5200)
    inspected = inspect_tabular_file(path)
    sheet = inspected["sheets"][0]
    assert sheet["row_count_estimate"] == 5200
    assert detect_layout(path).row_count_estimate == 5200


def test_csv_row_count_estimate_handles_small_and_nonterminated_files(tmp_path: Path) -> None:
    full = tmp_path / "full.csv"
    _write_csv(full, 10)
    assert inspect_tabular_file(full)["sheets"][0]["row_count_estimate"] == 10

    unterminated = tmp_path / "tail.csv"
    unterminated.write_bytes("物料编码,物料描述\n1,A\n2,B".encode("utf-8"))
    assert inspect_tabular_file(unterminated)["sheets"][0]["row_count_estimate"] == 2
    assert len(list(iter_tabular_rows(unterminated))) == 2


def test_csv_effective_rows_ignore_empty_and_whitespace_only_records(tmp_path: Path) -> None:
    path = tmp_path / "blank-records.csv"
    path.write_text(
        "物料编码,物料描述\n"
        "1,A\n"
        ",\n"
        "   ,   \n"
        "2,B\n",
        encoding="utf-8",
    )
    assert inspect_tabular_file(path)["sheets"][0]["row_count_estimate"] == 2
    assert len(list(iter_tabular_rows(path))) == 2


def test_xlsx_inflated_dimension_counts_only_effective_business_rows(tmp_path: Path) -> None:
    path = tmp_path / "inflated.xlsx"
    _write_inflated_xlsx(path, rows=100)
    _assert_inflated_dimension(path)

    inspected = inspect_tabular_file(path)
    sheet = inspected["sheets"][0]

    assert sheet["recommended_header_row"] == 1
    assert sheet["row_count_estimate"] == 100
    assert detect_layout(path).row_count_estimate == 100

    rows = list(iter_tabular_rows(path))
    assert len(rows) == 100
    assert rows[0]["物料编码"] == "M0001"
    assert rows[-1]["物料编码"] == "M0100"


def test_xlsx_inflated_dimension_does_not_trip_baseline_limit(tmp_path: Path) -> None:
    path = tmp_path / "target.xlsx"
    _write_inflated_xlsx(path, rows=100)
    _assert_inflated_dimension(path)

    loaded = load_target_rows_with_position(path, max_target_rows=500)
    assert len(loaded) == 100

    service = MatchService.__new__(MatchService)
    service.settings = SimpleNamespace(baseline_max_target_rows=500)
    config = SimpleNamespace(
        retrieval=SimpleNamespace(mode="auto"),
        rules=[],
    )
    assert service._use_vector(config, path) is False


def test_upload_inspection_api_returns_effective_xlsx_row_count(
    tmp_path: Path,
    authed: TestClient,
) -> None:
    path = tmp_path / "upload-inflated.xlsx"
    _write_inflated_xlsx(path, rows=100)
    _assert_inflated_dimension(path)

    response = authed.post(
        "/api/files/upload",
        data={"role": "target"},
        files={
            "file": (
                path.name,
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    sheet = payload["inspection"]["sheets"][0]
    assert sheet["row_count_estimate"] == 100
    assert payload["inspection"]["column_profile"]["row_count_estimate"] == 100

    reloaded = authed.get(f"/api/files/{payload['file']['file_id']}/inspection")
    assert reloaded.status_code == 200, reloaded.text
    assert reloaded.json()["inspection"]["sheets"][0]["row_count_estimate"] == 100


def test_xlsx_header_not_on_first_row_and_whitespace_rows_are_not_counted(tmp_path: Path) -> None:
    path = tmp_path / "offset-header.xlsx"
    _write_inflated_xlsx(path, rows=2, header_row=3)
    workbook = load_workbook(path)
    try:
        worksheet = workbook.active
        worksheet.cell(6, 1, "   ")
        worksheet.cell(6, 2, "\t")
        workbook.save(path)
    finally:
        workbook.close()

    layout = detect_layout(path)
    assert layout.header_row == 3
    assert layout.row_count_estimate == 2
    assert len(list(iter_tabular_rows(path))) == 2


def test_xlsx_empty_and_header_only_sheets_report_zero_data_rows(tmp_path: Path) -> None:
    empty = tmp_path / "empty.xlsx"
    Workbook().save(empty)
    assert inspect_tabular_file(empty)["sheets"][0]["row_count_estimate"] == 0

    header_only = tmp_path / "header-only.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["物料编码", "物料描述"])
    worksheet.cell(100_000, 1).number_format = "@"
    workbook.save(header_only)
    _assert_inflated_dimension(header_only)
    assert inspect_tabular_file(header_only)["sheets"][0]["row_count_estimate"] == 0


def test_xlsx_multi_sheet_recommendation_uses_effective_row_counts(tmp_path: Path) -> None:
    path = tmp_path / "multi.xlsx"
    workbook = Workbook()
    empty = workbook.active
    empty.title = "Empty"
    empty.append(["物料编码", "物料描述"])
    empty.cell(100_000, 1).number_format = "@"

    data = workbook.create_sheet("Data")
    data.append(["物料编码", "物料描述"])
    data.append(["M001", "电阻"])
    data.append(["M002", "电容"])
    workbook.save(path)

    inspected = inspect_tabular_file(path)
    counts = {sheet["sheet_name"]: sheet["row_count_estimate"] for sheet in inspected["sheets"]}
    assert counts == {"Empty": 0, "Data": 2}
    assert inspected["recommended_sheet"] == "Data"
