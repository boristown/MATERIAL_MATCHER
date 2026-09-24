from __future__ import annotations

from pathlib import Path

from material_matcher.domain.models import MatchingConfig
from material_matcher.matching.engine import match_rows, source_filter_allows


def _config(**over) -> MatchingConfig:
    doc = {
        "source_id_column": "编码",
        "scope_mode": "GLOBAL",
        "rules": [{"id": "name", "source": {"fields": ["名称"]}, "target": {"fields": ["名称"]}, "matcher": "exact", "weight": 100}],
        "decision": {"success_threshold": 60, "review_enabled": True, "review_threshold": 30, "top_n": 5},
        "retrieval": {"mode": "scan"},
    }
    doc.update(over)
    return MatchingConfig.model_validate(doc)


def test_source_filter_include_exclude(tmp_path: Path) -> None:
    cfg = _config()
    row = {"编码": "1", "名称": "电阻", "物料类型": "Z001"}
    assert source_filter_allows(row, cfg) is True
    cfg_inc = _config(source_filter={"field": "物料类型", "values": ["Z001", "A001"], "mode": "include"})
    cfg_exc = _config(source_filter={"field": "物料类型", "values": ["Z001"], "mode": "exclude"})
    cfg_contains = _config(source_filter={"field": "名称", "values": ["电"], "mode": "include", "match": "contains"})
    assert source_filter_allows(row, cfg_inc) is True
    assert source_filter_allows(row, cfg_exc) is False
    assert source_filter_allows(row, cfg_contains) is True
    assert source_filter_allows({"编码": "2", "名称": "电容", "物料类型": "Z002"}, cfg_inc) is False


def test_match_rows_skips_filtered_rows(tmp_path: Path) -> None:
    target = tmp_path / "t.csv"
    target.write_text("编码,集团码,名称\nT1,G1,电阻\nT2,G2,电容\n", encoding="utf-8")
    source = tmp_path / "s.csv"
    source.write_text("编码,名称,物料类型\n0001,电阻,Z001\n0002,电容,Z002\n", encoding="utf-8")
    cfg = _config(source_filter={"field": "物料类型", "values": ["Z001"], "mode": "include"})
    rows = match_rows(source, target, config=cfg, group_code_column="集团码", max_target_rows=1000)
    assert [row.source_id for row in rows] == ["0001"]
    assert rows[0].candidates[0].group_code == "G1"


def test_redecide_and_workbench_search_api(tmp_path: Path, authed) -> None:
    client = authed

    def upload(name: str, content: str, role: str) -> str:
        path = tmp_path / name
        path.write_text(content, encoding="utf-8")
        with path.open("rb") as stream:
            response = client.post("/api/files/upload", data={"role": role}, files={"file": (name, stream, "text/csv")})
        response.raise_for_status()
        return response.json()["file"]["file_id"]

    target_file = upload("目录.csv", "集团码,名称\nG1,电阻\nG2,电容\nG3,电感\n", "target")
    catalog = client.post("/api/catalogs", json={"name": "阈值测试", "source_file_id": target_file, "group_code_column": "集团码"}).json()
    source_file = upload("物料.csv", "编码,名称\n0001,电阻\n0002,电容器\n0003,完全不存在\n", "source")
    draft = client.post("/api/task-drafts", json={"name": "阈值"}).json()["draft_id"]
    client.put(f"/api/task-drafts/{draft}/data", json={"source_file_id": source_file, "catalog_version_id": catalog["version_id"]}).raise_for_status()
    client.put(f"/api/task-drafts/{draft}/rules", json={
        "source_id_column": "编码", "scope_mode": "GLOBAL",
        "rules": [{"id": "name", "source": {"fields": ["名称"]}, "target": {"fields": ["名称"]}, "matcher": "fuzzy", "weight": 100}],
        "decision": {"success_threshold": 90, "review_enabled": True, "review_threshold": 40, "top_n": 5},
        "retrieval": {"mode": "scan"},
    }).raise_for_status()
    task = client.post(f"/api/task-drafts/{draft}/start").json()["task_id"]
    import time as _time
    for _ in range(100):
        progress = client.get(f"/api/tasks/{task}/progress").json()
        if progress["status"] in {"COMPLETED", "FAILED"}:
            break
        _time.sleep(0.05)
    assert progress["status"] == "COMPLETED"

    # 模糊搜索待确认行
    searched = client.get(f"/api/tasks/{task}/workbench/items", params={"q": "电容"}).json()
    assert all("电容" in str(item["source_payload"]) for item in searched["items"])

    # 降低自动阈值 → 待确认行变为自动匹配
    before = client.get(f"/api/tasks/{task}/workbench/summary").json()
    result = client.post(f"/api/tasks/{task}/re-decide", json={"success_threshold": 55})
    assert result.status_code == 200, result.text
    after = result.json()["summary"]
    assert after["pending_review"] <= before["pending_review"]
    assert "review_threshold" not in result.json()
    # 旧客户端多传 review_threshold 仍兼容，但该字段不再参与判定或出现在响应中。
    legacy = client.post(f"/api/tasks/{task}/re-decide", json={"success_threshold": 40, "review_threshold": 99})
    assert legacy.status_code == 200
    assert "review_threshold" not in legacy.json()
    # 自动匹配阈值本身仍需合法。
    assert client.post(f"/api/tasks/{task}/re-decide", json={"success_threshold": 0}).status_code == 422

    # 正式结果生成后仍允许全局调参，但必须保留旧结果并形成新版本。
    from conftest import finalize_wait
    fin_v1 = finalize_wait(client, task)
    first_file = fin_v1["result_file_id"]
    first_versions = client.get(f"/api/tasks/{task}/result-revisions").json()
    assert len(first_versions) == 1

    redecide_v2 = client.post(f"/api/tasks/{task}/re-decide", json={"success_threshold": 70})
    assert redecide_v2.status_code == 200
    assert redecide_v2.json()["revision_no"] >= 2
    fin_v2 = finalize_wait(client, task)
    assert fin_v2["result_file_id"] != first_file
    versions = client.get(f"/api/tasks/{task}/result-revisions").json()
    assert len(versions) == 2
    assert {row["file_id"] for row in versions} == {first_file, fin_v2["result_file_id"]}
    assert client.get(f"/api/tasks/{task}/result-revisions/1").status_code == 200

    export = client.get(f"/api/tasks/{task}/result")
    assert export.status_code == 200
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(export.content))
    assert wb.sheetnames[0] == "匹配摘要"
    assert "最终匹配结果" in wb.sheetnames
