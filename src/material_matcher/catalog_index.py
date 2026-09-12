from __future__ import annotations

import json
import sqlite3
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from .bbq import EmbeddedBBQFlatIndex
from .embedding import EmbeddingProvider, create_embedding_provider
from .matching import MatchConfig, normalize_text


class CatalogIndexError(ValueError):
    pass


def _json_default(value: Any) -> str:
    return str(value)


def _target_text(row: dict[str, Any], config: MatchConfig) -> str:
    parts: list[str] = []
    for rule in config.mappings:
        for header in [part.strip() for part in rule.target_header.split(" + ")]:
            value = normalize_text(row.get(header))
            if value and value not in parts:
                parts.append(value)
    return " | ".join(parts)


def build_catalog_sqlite(
    source_path: Path,
    sqlite_path: Path,
    raw_config: dict[str, Any],
    *,
    batch_size: int = 2000,
) -> dict[str, Any]:
    """Convert an arbitrary target Excel to a stable dynamic SQLite catalog.

    The original row is stored as JSON so customer-specific fields remain
    dynamic. Frequently used retrieval columns are materialized separately.
    """
    config = MatchConfig.from_dict(raw_config)
    sqlite_path.parent.mkdir(parents=True, exist_ok=True)
    if sqlite_path.exists():
        sqlite_path.unlink()

    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if config.target_sheet and config.target_sheet in workbook.sheetnames:
        ws = workbook[config.target_sheet]
    else:
        ws = workbook[workbook.sheetnames[0]]
    headers = ["" if cell.value is None else str(cell.value).strip() for cell in ws[config.target_header_row]]
    if config.group_code_column not in headers:
        workbook.close()
        raise CatalogIndexError(f"集团码列不存在：{config.group_code_column}")

    connection = sqlite3.connect(sqlite_path)
    try:
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA synchronous=NORMAL")
        connection.execute(
            """
            CREATE TABLE records (
                row_id INTEGER PRIMARY KEY,
                group_code TEXT,
                group_value TEXT,
                search_text TEXT NOT NULL,
                payload_json TEXT NOT NULL
            )
            """
        )
        connection.execute("CREATE INDEX idx_records_group_value ON records(group_value)")
        connection.execute("CREATE INDEX idx_records_group_code ON records(group_code)")
        rows: list[tuple[str, str, str, str]] = []
        count = 0
        for values in ws.iter_rows(min_row=config.target_header_row + 1, values_only=True):
            row = {headers[idx]: value for idx, value in enumerate(values) if idx < len(headers) and headers[idx]}
            if not any(value is not None and str(value).strip() for value in row.values()):
                continue
            group_code = "" if row.get(config.group_code_column) is None else str(row.get(config.group_code_column))
            group_value = ""
            if config.target_group_column:
                group_value = normalize_text(row.get(config.target_group_column))
            search_text = _target_text(row, config)
            payload_json = json.dumps(row, ensure_ascii=False, default=_json_default)
            rows.append((group_code, group_value, search_text, payload_json))
            count += 1
            if len(rows) >= batch_size:
                connection.executemany(
                    "INSERT INTO records(group_code, group_value, search_text, payload_json) VALUES(?,?,?,?)",
                    rows,
                )
                connection.commit()
                rows.clear()
        if rows:
            connection.executemany(
                "INSERT INTO records(group_code, group_value, search_text, payload_json) VALUES(?,?,?,?)",
                rows,
            )
            connection.commit()
        connection.execute("CREATE TABLE metadata(key TEXT PRIMARY KEY, value TEXT NOT NULL)")
        metadata = {
            "count": count,
            "group_code_column": config.group_code_column,
            "target_group_column": config.target_group_column or "",
            "target_sheet": config.target_sheet or ws.title,
            "target_header_row": config.target_header_row,
        }
        connection.executemany(
            "INSERT INTO metadata(key,value) VALUES(?,?)",
            [(key, json.dumps(value, ensure_ascii=False)) for key, value in metadata.items()],
        )
        connection.commit()
        return metadata
    finally:
        connection.close()
        workbook.close()


def catalog_count(sqlite_path: Path) -> int:
    with sqlite3.connect(sqlite_path) as connection:
        row = connection.execute("SELECT COUNT(*) FROM records").fetchone()
    return int(row[0]) if row else 0


def iter_catalog_embedding_batches(
    sqlite_path: Path,
    provider: EmbeddingProvider,
    *,
    batch_size: int = 2048,
) -> Iterator[tuple[list[str], np.ndarray]]:
    connection = sqlite3.connect(sqlite_path)
    try:
        cursor = connection.execute("SELECT row_id, search_text FROM records ORDER BY row_id")
        while True:
            batch = cursor.fetchmany(batch_size)
            if not batch:
                break
            ids = [str(row[0]) for row in batch]
            vectors = provider.encode(row[1] or "" for row in batch)
            yield ids, vectors
    finally:
        connection.close()


def build_catalog_bbq_index(
    sqlite_path: Path,
    index_dir: Path,
    embedding_config: dict[str, Any] | None = None,
    *,
    batch_size: int = 2048,
) -> dict[str, Any]:
    provider = create_embedding_provider(embedding_config)
    count = catalog_count(sqlite_path)
    if count == 0:
        raise CatalogIndexError("集团码 Catalog 为空，无法建立向量索引")

    def batches():
        return iter_catalog_embedding_batches(sqlite_path, provider, batch_size=batch_size)

    index = EmbeddedBBQFlatIndex.build_streaming_to_directory(
        batches,
        count=count,
        dimensions=provider.dimensions,
        directory=index_dir,
        store_float16_rerank=False,
    )
    return {
        "index_type": "embedded_bbq_flat",
        "embedding_provider": provider.provider_id,
        "dimensions": provider.dimensions,
        "count": index.count,
        "packed_bytes_per_vector": index.packed_bytes_per_vector,
        "target_bits_per_dimension": 1,
        "query_bits_per_dimension": 4,
        "index_dir": str(index_dir),
    }


class CatalogStore:
    def __init__(self, sqlite_path: str | Path) -> None:
        self.sqlite_path = Path(sqlite_path)

    def fetch_rows(self, row_ids: list[int]) -> dict[int, dict[str, Any]]:
        if not row_ids:
            return {}
        result: dict[int, dict[str, Any]] = {}
        connection = sqlite3.connect(self.sqlite_path)
        try:
            for start in range(0, len(row_ids), 500):
                chunk = row_ids[start : start + 500]
                placeholders = ",".join("?" for _ in chunk)
                cursor = connection.execute(
                    f"SELECT row_id, payload_json FROM records WHERE row_id IN ({placeholders})",
                    chunk,
                )
                for row_id, payload_json in cursor:
                    result[int(row_id)] = json.loads(payload_json)
        finally:
            connection.close()
        return result

    def group_candidate_indices(self, groups: list[str]) -> np.ndarray:
        normalized = [normalize_text(group) for group in groups if normalize_text(group)]
        if not normalized:
            return np.array([], dtype=np.int64)
        connection = sqlite3.connect(self.sqlite_path)
        try:
            placeholders = ",".join("?" for _ in normalized)
            rows = connection.execute(
                f"SELECT row_id FROM records WHERE group_value IN ({placeholders}) ORDER BY row_id",
                normalized,
            ).fetchall()
        finally:
            connection.close()
        # BBQ positions are zero-based while SQLite row_id begins at one.
        return np.asarray([int(row[0]) - 1 for row in rows], dtype=np.int64)
