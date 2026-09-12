from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass
from typing import Any, BinaryIO

from openpyxl import load_workbook


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("matnr", "物料", "物料号", "物料编码", "物资编码", "编码"),
    "group_code": ("集团码", "集团编码", "集团物资编码", "zjtm"),
    "material_name": ("物料名称", "物料描述", "名称", "品名", "zwlmc", "maktx"),
    "model": ("型号", "牌号", "系列", "zxh", "zx h"),
    "specification": ("规格", "型号规格", "详细型号规格", "尺寸", "zxhgg", "zgg"),
    "manufacturer": ("生产厂家", "制造商", "厂家", "品牌", "zsccj", "herst"),
    "material_group": ("物料组", "物料类型", "分类", "matkl", "zwlyx"),
    "standard": ("标准", "采用标准", "技术标准", "规范", "标准号", "zcgbz", "zjsbz", "zbzh"),
    "unit": ("计量单位", "单位", "meins"),
}


@dataclass
class ColumnInspection:
    index: int
    header: str
    logical_hint: str | None
    confidence: float
    samples: list[str]
    leading_zero_risk: bool
    null_tokens_seen: list[str]


@dataclass
class SheetInspection:
    name: str
    header_row: int
    header_confidence: float
    row_count: int
    column_count: int
    columns: list[ColumnInspection]
    warnings: list[str]


@dataclass
class WorkbookInspection:
    sheets: list[SheetInspection]
    recommended_sheet: str | None

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    return re.sub(r"[\s_\-—–/（）()【】\[\]:：]+", "", text)


def _guess_field(header: str) -> tuple[str | None, float]:
    normalized = _normalize_header(header)
    if not normalized:
        return None, 0.0

    best_field: str | None = None
    best_score = 0.0
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            candidate = _normalize_header(alias)
            if normalized == candidate:
                score = 0.99
            elif candidate and (candidate in normalized or normalized in candidate):
                score = 0.82
            else:
                continue
            if score > best_score:
                best_field, best_score = field, score
    return best_field, best_score


def _header_row_score(values: list[Any]) -> float:
    non_empty = [v for v in values if v is not None and str(v).strip()]
    if not non_empty:
        return 0.0
    strings = [v for v in non_empty if isinstance(v, str)]
    unique = len({_normalize_header(v) for v in non_empty})
    hint_count = sum(1 for v in non_empty if _guess_field(str(v))[0])
    # Header rows are normally text-heavy, unique and often contain known business labels.
    return len(non_empty) + len(strings) * 1.7 + unique * 0.35 + hint_count * 2.5


def _detect_header_row(ws, max_scan_rows: int = 12) -> tuple[int, float]:
    candidates: list[tuple[int, float]] = []
    max_row = min(ws.max_row or 1, max_scan_rows)
    max_col = min(ws.max_column or 1, 200)
    for row_idx in range(1, max_row + 1):
        values = [ws.cell(row=row_idx, column=col).value for col in range(1, max_col + 1)]
        candidates.append((row_idx, _header_row_score(values)))
    candidates.sort(key=lambda item: item[1], reverse=True)
    if not candidates:
        return 1, 0.0
    best_row, best_score = candidates[0]
    second = candidates[1][1] if len(candidates) > 1 else 0.0
    confidence = 1.0 if best_score <= 0 else min(1.0, 0.55 + max(0.0, best_score - second) / max(best_score, 1.0))
    return best_row, round(confidence, 3)


def _sample_values(ws, column: int, start_row: int, limit: int = 5) -> tuple[list[str], list[Any]]:
    rendered: list[str] = []
    raw_values: list[Any] = []
    for row in range(start_row, min(ws.max_row or start_row, start_row + 200) + 1):
        value = ws.cell(row=row, column=column).value
        if value is None or str(value).strip() == "":
            continue
        raw_values.append(value)
        rendered.append(str(value).strip())
        if len(rendered) >= limit:
            break
    return rendered, raw_values


def inspect_excel(file_obj: BinaryIO | bytes, max_header_scan_rows: int = 12, sample_limit: int = 5) -> WorkbookInspection:
    stream: BinaryIO
    if isinstance(file_obj, bytes):
        stream = io.BytesIO(file_obj)
    else:
        stream = file_obj

    workbook = load_workbook(stream, read_only=True, data_only=True)
    sheets: list[SheetInspection] = []

    for ws in workbook.worksheets:
        header_row, header_confidence = _detect_header_row(ws, max_header_scan_rows)
        warnings: list[str] = []
        columns: list[ColumnInspection] = []
        non_empty_headers = 0

        for column in range(1, min(ws.max_column or 1, 300) + 1):
            raw_header = ws.cell(row=header_row, column=column).value
            header = "" if raw_header is None else str(raw_header).strip()
            if not header:
                continue
            non_empty_headers += 1
            logical_hint, confidence = _guess_field(header)
            samples, raw_samples = _sample_values(ws, column, header_row + 1, sample_limit)
            leading_zero_risk = logical_hint in {"source_id", "group_code", "material_group"} and any(
                isinstance(value, (int, float)) and not isinstance(value, bool) for value in raw_samples
            )
            null_tokens = sorted({value for value in samples if value in {"88", "-", "/", "<NULL>", "NULL", "N/A"}})
            columns.append(
                ColumnInspection(
                    index=column,
                    header=header,
                    logical_hint=logical_hint,
                    confidence=confidence,
                    samples=samples,
                    leading_zero_risk=leading_zero_risk,
                    null_tokens_seen=null_tokens,
                )
            )
            if leading_zero_risk:
                warnings.append(f"列“{header}”疑似编码字段，但 Excel 中包含数值单元格，前导零可能已丢失。")

        if header_confidence < 0.65:
            warnings.append("表头识别置信度较低，建议人工确认表头所在行。")
        if non_empty_headers == 0:
            warnings.append("未识别到有效表头。")

        sheets.append(
            SheetInspection(
                name=ws.title,
                header_row=header_row,
                header_confidence=header_confidence,
                row_count=ws.max_row or 0,
                column_count=non_empty_headers,
                columns=columns,
                warnings=warnings,
            )
        )

    recommended_sheet = None
    if sheets:
        recommended_sheet = max(sheets, key=lambda item: (item.header_confidence, item.row_count, item.column_count)).name
    return WorkbookInspection(sheets=sheets, recommended_sheet=recommended_sheet)
