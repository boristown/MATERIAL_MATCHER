from __future__ import annotations

from collections import Counter
from datetime import date, datetime
import csv
from pathlib import Path
import re

from openpyxl import load_workbook

DEFAULT_COLUMN_PROFILE_SCAN_LIMIT = 5_000
MAX_COLUMN_PROFILE_SCAN_LIMIT = 50_000
DEFAULT_TOP_VALUES_LIMIT = 20
MAX_ENUM_UNIQUE_VALUES = 20
SAMPLE_VALUES_LIMIT = 5

_INTEGER_RE = re.compile(r"^[+-]?\d+$")
_NUMBER_RE = re.compile(r"^[+-]?(?:\d+\.\d*|\d*\.\d+)(?:[eE][+-]?\d+)?$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(
    r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?(?:Z|[+-]\d{2}:?\d{2})?$"
)


def _render_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def _value_kind(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, datetime):
        return "DATETIME"
    if isinstance(value, date):
        return "DATE"
    if isinstance(value, int):
        return "INTEGER"
    if isinstance(value, float):
        return "INTEGER" if value.is_integer() else "NUMBER"

    text = str(value).strip()
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"true", "false"}:
        return "BOOLEAN"
    if _INTEGER_RE.fullmatch(text):
        unsigned = text.lstrip("+-")
        if len(unsigned) > 1 and unsigned.startswith("0"):
            return "TEXT"
        return "INTEGER"
    if _NUMBER_RE.fullmatch(text):
        return "NUMBER"
    if _DATE_RE.fullmatch(text):
        try:
            date.fromisoformat(text)
            return "DATE"
        except ValueError:
            pass
    if _DATETIME_RE.fullmatch(text):
        try:
            datetime.fromisoformat(text.replace("Z", "+00:00"))
            return "DATETIME"
        except ValueError:
            pass
    return "TEXT"


def _infer_datatype(kinds: set[str]) -> str:
    if not kinds:
        return "EMPTY"
    if kinds <= {"INTEGER"}:
        return "INTEGER"
    if kinds <= {"INTEGER", "NUMBER"}:
        return "NUMBER" if "NUMBER" in kinds else "INTEGER"
    if kinds <= {"DATE"}:
        return "DATE"
    if kinds <= {"DATE", "DATETIME"}:
        return "DATETIME" if "DATETIME" in kinds else "DATE"
    if kinds <= {"BOOLEAN"}:
        return "BOOLEAN"
    if kinds <= {"TEXT"}:
        return "TEXT"
    return "MIXED"


def _is_enum_candidate(*, unique_count: int, non_empty_count: int, datatype: str) -> bool:
    if unique_count == 0 or unique_count > MAX_ENUM_UNIQUE_VALUES:
        return False
    if unique_count == 1:
        return True

    # Small templates may legitimately contain one example per permitted value.
    # This remains a candidate hint only; callers also receive cardinality and
    # truncation metadata and should keep the final mapping user-confirmed.
    if non_empty_count <= 3 and unique_count == non_empty_count:
        return datatype in {"BOOLEAN", "INTEGER", "NUMBER", "DATE", "TEXT"}

    repetition_ratio = unique_count / max(non_empty_count, 1)
    if unique_count <= 5:
        return repetition_ratio <= 0.80
    return non_empty_count >= unique_count * 2 and repetition_ratio <= 0.50


def _new_accumulator(index: int, header: object) -> dict[str, object]:
    return {
        "column_index": index,
        "field_name": "" if header is None else str(header).strip(),
        "counter": Counter(),
        "sample_values": [],
        "sample_seen": set(),
        "kinds": set(),
        "non_empty_count": 0,
    }


def _add_value(accumulator: dict[str, object], raw_value: object) -> None:
    rendered = _render_value(raw_value)
    if rendered is None or rendered.strip() == "":
        return

    accumulator["non_empty_count"] = int(accumulator["non_empty_count"]) + 1
    counter = accumulator["counter"]
    assert isinstance(counter, Counter)
    counter[rendered] += 1

    kinds = accumulator["kinds"]
    assert isinstance(kinds, set)
    kind = _value_kind(raw_value)
    if kind:
        kinds.add(kind)

    seen = accumulator["sample_seen"]
    samples = accumulator["sample_values"]
    assert isinstance(seen, set) and isinstance(samples, list)
    if rendered not in seen and len(samples) < SAMPLE_VALUES_LIMIT:
        seen.add(rendered)
        samples.append(rendered)


def _finalize_columns(
    accumulators: list[dict[str, object]],
    *,
    top_values_limit: int,
) -> list[dict[str, object]]:
    columns: list[dict[str, object]] = []
    for accumulator in accumulators:
        field_name = str(accumulator["field_name"])
        if not field_name:
            continue

        counter = accumulator["counter"]
        kinds = accumulator["kinds"]
        assert isinstance(counter, Counter) and isinstance(kinds, set)
        datatype = _infer_datatype(kinds)
        unique_count = len(counter)
        non_empty_count = int(accumulator["non_empty_count"])
        top_values = [
            {"value": value, "count": count}
            for value, count in counter.most_common(top_values_limit)
        ]
        columns.append(
            {
                "field_name": field_name,
                "column_index": int(accumulator["column_index"]),
                "datatype": datatype,
                "unique_count": unique_count,
                "non_empty_count": non_empty_count,
                "sample_values": list(accumulator["sample_values"]),
                "top_values": top_values,
                "enum_candidate": _is_enum_candidate(
                    unique_count=unique_count,
                    non_empty_count=non_empty_count,
                    datatype=datatype,
                ),
            }
        )
    return columns


def _csv_encoding(path: Path) -> str:
    with path.open("rb") as stream:
        raw = stream.read(256 * 1024)
    try:
        raw.decode("utf-8-sig")
        return "utf-8-sig"
    except UnicodeDecodeError:
        raw.decode("gb18030")
        return "gb18030"


def _resolve_layout(
    path: Path,
    *,
    sheet_name: str | None,
    header_row: int | None,
) -> tuple[str, int, int]:
    from material_matcher.ingestion.inspector import inspect_tabular_file

    inspection = inspect_tabular_file(path)
    sheets = list(inspection.get("sheets") or [])
    if not sheets:
        raise ValueError("未识别到可读取的数据表")

    selected_sheet = sheet_name or inspection.get("recommended_sheet") or sheets[0].get("sheet_name")
    selected = next((item for item in sheets if item.get("sheet_name") == selected_sheet), None)
    if selected is None:
        raise ValueError(f"工作表不存在: {selected_sheet}")

    selected_header = int(header_row or selected.get("recommended_header_row") or 1)
    return str(selected_sheet), selected_header, int(selected.get("row_count_estimate") or 0)


def profile_tabular_columns(
    path: Path,
    *,
    sheet_name: str | None = None,
    header_row: int | None = None,
    scan_limit: int = DEFAULT_COLUMN_PROFILE_SCAN_LIMIT,
    top_values_limit: int = DEFAULT_TOP_VALUES_LIMIT,
) -> dict[str, object]:
    if scan_limit < 1 or scan_limit > MAX_COLUMN_PROFILE_SCAN_LIMIT:
        raise ValueError(f"scan_limit 必须在 1..{MAX_COLUMN_PROFILE_SCAN_LIMIT} 之间")
    if top_values_limit < 1:
        raise ValueError("top_values_limit 必须大于 0")

    selected_sheet, selected_header, row_count_estimate = _resolve_layout(
        path,
        sheet_name=sheet_name,
        header_row=header_row,
    )
    suffix = path.suffix.lower()
    scanned_rows = 0

    if suffix == ".csv":
        if selected_sheet != "CSV":
            raise ValueError(f"CSV 文件不存在工作表: {selected_sheet}")
        with path.open("r", encoding=_csv_encoding(path), newline="") as stream:
            reader = csv.reader(stream)
            for _ in range(selected_header - 1):
                next(reader, None)
            headers = next(reader, []) or []
            accumulators = [_new_accumulator(index + 1, header) for index, header in enumerate(headers)]
            for row in reader:
                if scanned_rows >= scan_limit:
                    break
                if not any(value != "" for value in row):
                    continue
                scanned_rows += 1
                for index, accumulator in enumerate(accumulators):
                    _add_value(accumulator, row[index] if index < len(row) else None)

    elif suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if selected_sheet not in workbook.sheetnames:
                raise ValueError(f"工作表不存在: {selected_sheet}")
            worksheet = workbook[selected_sheet]
            headers = next(
                worksheet.iter_rows(
                    min_row=selected_header,
                    max_row=selected_header,
                    values_only=True,
                ),
                (),
            )
            accumulators = [_new_accumulator(index + 1, header) for index, header in enumerate(headers)]
            for row in worksheet.iter_rows(min_row=selected_header + 1, values_only=True):
                if scanned_rows >= scan_limit:
                    break
                if not any(value is not None and str(value).strip() != "" for value in row):
                    continue
                scanned_rows += 1
                for index, accumulator in enumerate(accumulators):
                    _add_value(accumulator, row[index] if index < len(row) else None)
        finally:
            workbook.close()
    else:
        raise ValueError("unsupported tabular file")

    return {
        "sheet_name": selected_sheet,
        "header_row": selected_header,
        "scanned_rows": scanned_rows,
        "row_count_estimate": row_count_estimate,
        "truncated": row_count_estimate > scanned_rows and scanned_rows >= scan_limit,
        "columns": _finalize_columns(accumulators, top_values_limit=top_values_limit),
    }


def inspect_tabular_file_with_profiles(
    path: Path,
    *,
    scan_limit: int = DEFAULT_COLUMN_PROFILE_SCAN_LIMIT,
) -> dict[str, object]:
    from material_matcher.ingestion.inspector import inspect_tabular_file

    inspection = inspect_tabular_file(path)
    recommended_sheet = inspection.get("recommended_sheet")
    sheets = list(inspection.get("sheets") or [])
    selected = next((item for item in sheets if item.get("sheet_name") == recommended_sheet), None)
    if selected is None:
        return inspection

    profile = profile_tabular_columns(
        path,
        sheet_name=str(selected["sheet_name"]),
        header_row=int(selected.get("recommended_header_row") or 1),
        scan_limit=scan_limit,
    )
    by_index = {int(item["column_index"]): item for item in profile["columns"]}
    for column in list(selected.get("columns") or []):
        item = by_index.get(int(column.get("index") or 0))
        if item is None:
            continue
        column.update(
            {
                "field_name": item["field_name"],
                "datatype": item["datatype"],
                "unique_count": item["unique_count"],
                "sample_values": item["sample_values"],
                "top_values": item["top_values"],
                "enum_candidate": item["enum_candidate"],
            }
        )

    inspection["column_profile"] = {
        key: profile[key]
        for key in ("sheet_name", "header_row", "scanned_rows", "row_count_estimate", "truncated")
    }
    return inspection
