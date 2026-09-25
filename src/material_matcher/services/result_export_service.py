from __future__ import annotations

from collections import OrderedDict
from datetime import datetime
from io import BytesIO
import json
from typing import Any, Iterable, Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from material_matcher.settings import Settings
from material_matcher.services.export_profile import load_export_profile
from material_matcher.storage.files import FileRepository
from material_matcher.services.task_input_asset_service import TaskInputAssetService
from material_matcher.services.task_service import (
    format_duration_ms,
    resolve_task_scheme_name,
    safe_business_filename,
    task_time_fields,
)
from material_matcher.storage.metadata import MetadataRepository



def _now() -> str:
    return datetime.now().astimezone().isoformat()


def _json_object(value: object) -> dict[str, object]:
    if isinstance(value, dict):
        return dict(value)
    if not value:
        return {}
    try:
        decoded = json.loads(str(value))
    except (TypeError, json.JSONDecodeError):
        return {}
    return dict(decoded) if isinstance(decoded, dict) else {}


def _json_list(value: object) -> list[dict[str, object]]:
    if isinstance(value, list):
        return [dict(item) for item in value if isinstance(item, dict)]
    if not value:
        return []
    try:
        decoded = json.loads(str(value))
    except (TypeError, json.JSONDecodeError):
        return []
    return [dict(item) for item in decoded if isinstance(item, dict)] if isinstance(decoded, list) else []


def _display_time(value: object) -> datetime | None:
    if value is None or value == "":
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.replace(tzinfo=None)
    except ValueError:
        return None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value)


def _is_missing(value: object) -> bool:
    return value is None or str(value).strip() == ""


def _ordered_fields(payloads: Iterable[Mapping[str, object]]) -> list[str]:
    fields: OrderedDict[str, None] = OrderedDict()
    for payload in payloads:
        for key in payload:
            text = str(key).strip()
            if text and not text.startswith("__material_matcher_"):
                fields.setdefault(text, None)
    return list(fields)


class ResultExportService:
    """Build the customer-facing final workbook without exposing internal diagnostics."""

    def __init__(self, metadata: MetadataRepository, files: FileRepository, settings: Settings) -> None:
        self.meta = metadata
        self.files = files
        self.settings = settings
        self.export_profile = load_export_profile(settings.config_dir)
        self.task_input_assets = TaskInputAssetService(metadata, files)

    def _task_actor(self, task_id: str, task: Mapping[str, object], kind: str) -> str:
        direct_keys = (
            ("created_by", "creator", "created_user", "created_username")
            if kind == "created"
            else ("started_by", "starter", "started_user", "started_username")
        )
        for key in direct_keys:
            value = task.get(key)
            if value:
                return str(value)

        action_tokens = ("CREATE", "CREATED") if kind == "created" else ("START", "STARTED")
        with self.meta.connect() as connection:
            rows = connection.execute(
                "SELECT action,payload FROM audit_events WHERE entity_type='task' AND entity_id=? ORDER BY created_at",
                (task_id,),
            ).fetchall()
        for row in rows:
            action = str(row["action"] or "").upper()
            if not any(token in action for token in action_tokens):
                continue
            payload = _json_object(row["payload"])
            for key in ("operator", "username", "user", "actor", "created_by", "started_by"):
                value = payload.get(key)
                if value:
                    return str(value)
        return ""

    @staticmethod
    def _latest_reviews(reviews: Iterable[Mapping[str, object]]) -> dict[str, Mapping[str, object]]:
        result: dict[str, Mapping[str, object]] = {}
        for review in reviews:
            source_row_id = str(review.get("source_row_id") or "")
            if not source_row_id:
                continue
            previous = result.get(source_row_id)
            if previous is None or str(review.get("created_at") or "") >= str(previous.get("created_at") or ""):
                result[source_row_id] = review
        return result

    @staticmethod
    def _candidate_map(candidates: Iterable[Mapping[str, object]]) -> dict[str, list[Mapping[str, object]]]:
        result: dict[str, list[Mapping[str, object]]] = {}
        for candidate in candidates:
            result.setdefault(str(candidate.get("source_row_id") or ""), []).append(candidate)
        for rows in result.values():
            rows.sort(key=lambda row: int(row.get("rank") or 999))
        return result

    @staticmethod
    def _selected_candidate(item: Mapping[str, object], candidates: list[Mapping[str, object]]) -> Mapping[str, object] | None:
        final_code = item.get("final_group_code")
        if final_code is None or str(final_code) == "":
            return None
        return next(
            (candidate for candidate in candidates if str(candidate.get("target_group_code")) == str(final_code)),
            None,
        )

    @staticmethod
    def _review_target_row(review: Mapping[str, object] | None, candidates: list[Mapping[str, object]]) -> object:
        if not review:
            return None
        for key in ("target_row_number", "selected_target_row", "target_excel_row", "target_row"):
            if review.get(key) not in {None, ""}:
                return review.get(key)
        selected = review.get("selected_group_code")
        if selected in {None, ""}:
            return None
        candidate = next(
            (row for row in candidates if str(row.get("target_group_code")) == str(selected)),
            None,
        )
        return candidate.get("target_row_number") if candidate else None

    @staticmethod
    def _rules(task: Mapping[str, object]) -> dict[str, dict[str, list[str]]]:
        snapshot = task.get("config_snapshot")
        if isinstance(snapshot, str):
            snapshot = _json_object(snapshot)
        if not isinstance(snapshot, dict):
            return {}
        result: dict[str, dict[str, list[str]]] = {}
        for raw_rule in snapshot.get("rules") or []:
            if not isinstance(raw_rule, dict):
                continue
            rule_id = str(raw_rule.get("id") or "")
            if not rule_id:
                continue
            source = raw_rule.get("source") if isinstance(raw_rule.get("source"), dict) else {}
            target = raw_rule.get("target") if isinstance(raw_rule.get("target"), dict) else {}
            result[rule_id] = {
                "source": [str(value) for value in source.get("fields") or [] if str(value)],
                "target": [str(value) for value in target.get("fields") or [] if str(value)],
            }
        return result

    @staticmethod
    def _field_states(
        source_payload: Mapping[str, object],
        target_payload: Mapping[str, object],
        candidate: Mapping[str, object] | None,
        rules: Mapping[str, Mapping[str, list[str]]],
    ) -> tuple[dict[str, str], dict[str, str]]:
        source_states: dict[str, str] = {}
        target_states: dict[str, str] = {}
        priority = {"exact": 1, "partial": 2, "different": 3, "missing": 4}

        def set_state(states: dict[str, str], field: str, state: str, payload: Mapping[str, object]) -> None:
            if _is_missing(payload.get(field)):
                state = "missing"
            current = states.get(field)
            if current is None or priority[state] > priority[current]:
                states[field] = state

        field_scores = _json_list(candidate.get("field_scores") if candidate else None)
        for field_score in field_scores:
            rule_id = str(field_score.get("rule_id") or "")
            rule = rules.get(rule_id)
            if not rule:
                continue
            try:
                score = float(field_score.get("score") or 0.0)
            except (TypeError, ValueError):
                score = 0.0
            state = "exact" if score >= 0.999999 else "partial" if score > 0 else "different"
            for field in rule.get("source", []):
                set_state(source_states, field, state, source_payload)
            for field in rule.get("target", []):
                set_state(target_states, field, state, target_payload)

        for field in set(source_payload).intersection(target_payload):
            if field in source_states or field in target_states:
                continue
            source_value = source_payload.get(field)
            target_value = target_payload.get(field)
            if _is_missing(source_value) or _is_missing(target_value):
                state = "missing"
            else:
                state = "exact" if str(source_value).strip() == str(target_value).strip() else "different"
            source_states[field] = state
            target_states[field] = state

        for field, value in source_payload.items():
            if _is_missing(value):
                source_states[field] = "missing"
        for field, value in target_payload.items():
            if _is_missing(value):
                target_states[field] = "missing"
        return source_states, target_states

    def _solid_fill(self, key: str) -> PatternFill:
        return PatternFill("solid", fgColor=self.export_profile.color(key))

    def _thin_border(self) -> Border:
        color = self.export_profile.color("border_color")
        return Border(
            left=Side(style="thin", color=color),
            right=Side(style="thin", color=color),
            top=Side(style="thin", color=color),
            bottom=Side(style="thin", color=color),
        )

    def _fill_for_state(self, state: str | None) -> PatternFill | None:
        key = {
            "exact": "exact_fill",
            "partial": "partial_fill",
            "different": "different_fill",
            "missing": "missing_fill",
        }.get(state or "")
        return self._solid_fill(key) if key else None

    def _append_grouped_headers(
        self,
        sheet: Any,
        groups: list[tuple[str, list[str], PatternFill]],
    ) -> dict[str, tuple[int, int]]:
        ranges: dict[str, tuple[int, int]] = {}
        column = 1
        border = self._thin_border()
        header_fill = self._solid_fill("header_fill")
        header_font = Font(color=self.export_profile.color("header_font_color"), bold=True, size=10)
        for group_name, headers, fill in groups:
            if not headers:
                continue
            start = column
            end = column + len(headers) - 1
            sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            top = sheet.cell(1, start, group_name)
            top.fill = fill
            top.font = Font(bold=True, color=self.export_profile.color("group_font_color"))
            top.alignment = Alignment(horizontal="center", vertical="center")
            for current in range(start, end + 1):
                cell = sheet.cell(1, current)
                cell.fill = fill
                cell.border = border
            for offset, header in enumerate(headers):
                cell = sheet.cell(2, start + offset, header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            ranges[group_name] = (start, end)
            column = end + 1
        sheet.row_dimensions[1].height = self.export_profile.group_header_height
        sheet.row_dimensions[2].height = self.export_profile.column_header_height
        sheet.freeze_panes = self.export_profile.freeze_panes
        return ranges

    def _style_data_sheet(self, sheet: Any, *, max_width: int | None = None) -> None:
        sheet.sheet_view.showGridLines = False
        border = self._thin_border()
        resolved_max_width = max_width or self.export_profile.max_column_width
        for row in sheet.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
        for index in range(1, sheet.max_column + 1):
            letter = get_column_letter(index)
            width = 10
            for cell in list(sheet[letter])[:250]:
                if cell.value is not None:
                    width = max(width, min(resolved_max_width, len(str(cell.value)) + 2))
            sheet.column_dimensions[letter].width = min(resolved_max_width, width)
        if sheet.max_column:
            sheet.auto_filter.ref = f"A2:{get_column_letter(sheet.max_column)}{max(2, sheet.max_row)}"

    @staticmethod
    def _write_text(cell: Any, value: object) -> None:
        cell.value = _text(value)
        cell.number_format = "@"

    def _write_time(self, cell: Any, value: object) -> None:
        parsed = _display_time(value)
        if parsed is None:
            cell.value = ""
            cell.number_format = "@"
        else:
            cell.value = parsed
            cell.number_format = self.export_profile.date_format


    class _Formats:
        _STATE_FILL_KEYS = {"exact": "exact_fill", "partial": "partial_fill", "different": "different_fill", "missing": "missing_fill"}

        def __init__(self, workbook, profile) -> None:
            self.workbook = workbook
            self.profile = profile
            self._cache: dict = {}

        @staticmethod
        def _hex(color: object) -> str:
            text = str(color or "").strip()
            if not text:
                return ""
            return text if text.startswith("#") else "#" + text

        def text(self, state: str | None = None, *, bold: bool = False) -> object:
            key = ("text", state, bold)
            if key not in self._cache:
                options = {"font_name": "微软雅黑", "font_size": 10, "valign": "vcenter", "num_format": "@"}
                fill = self._state_hex(state)
                if fill:
                    options["bg_color"] = fill
                if bold:
                    options["bold"] = True
                self._cache[key] = self.workbook.add_format(options)
            return self._cache[key]

        def _state_hex(self, state: str | None) -> str:
            if not state:
                return ""
            color_key = self._STATE_FILL_KEYS.get(state)
            return self._hex(self.profile.color(color_key)) if color_key else ""

        def number(self) -> object:
            return self._cache.setdefault("number", self.workbook.add_format({"num_format": "0.0", "font_name": "微软雅黑", "font_size": 10, "valign": "vcenter"}))

        def datetime(self) -> object:
            return self._cache.setdefault("datetime", self.workbook.add_format({"num_format": "yyyy-mm-dd hh:mm", "font_name": "微软雅黑", "font_size": 10, "valign": "vcenter"}))

        def group(self, fill_key: str) -> object:
            key = ("group", fill_key)
            if key not in self._cache:
                self._cache[key] = self.workbook.add_format({
                    "bold": True, "font_color": self._hex(self.profile.color("group_font_color")),
                    "bg_color": self._hex(self.profile.color(fill_key)),
                    "align": "center", "valign": "vcenter", "border": 1, "border_color": self._hex(self.profile.color("header_font_color")),
                })
            return self._cache[key]

        def column_header(self) -> object:
            return self._cache.setdefault("column_header", self.workbook.add_format({
                "bold": True, "font_size": 10,
                "font_color": self._hex(self.profile.color("header_font_color")),
                "bg_color": self._hex(self.profile.color("header_fill")),
                "align": "center", "valign": "vcenter", "text_wrap": True,
            }))

        def summary_title(self) -> object:
            return self._cache.setdefault("summary_title", self.workbook.add_format({"bold": True, "font_size": 16, "font_color": self._hex(self.profile.color("summary_title_color"))}))

        def summary_key(self) -> object:
            return self._cache.setdefault("summary_key", self.workbook.add_format({
                "bold": True, "font_color": self._hex(self.profile.color("summary_key_color")),
                "bg_color": self._hex(self.profile.color("summary_key_fill")), "valign": "vcenter",
            }))

        def legend(self, color_key: str) -> object:
            key = ("legend", color_key)
            if key not in self._cache:
                self._cache[key] = self.workbook.add_format({"bold": True, "align": "center", "valign": "vcenter", "bg_color": self._hex(self.profile.color(color_key))})
            return self._cache[key]

        def status(self, status_color: str | None) -> object:
            key = ("status", status_color or "")
            if key not in self._cache:
                options = {"font_name": "微软雅黑", "font_size": 10, "valign": "vcenter"}
                if status_color:
                    options["bg_color"] = self._hex(status_color)
                self._cache[key] = self.workbook.add_format(options)
            return self._cache[key]

    def _cell_text(self, value: object) -> str:
        return _text(value)

    def _xl_write(self, sheet, row: int, col: int, value: object, *, is_time: bool, formats: "_Formats") -> None:
        if is_time:
            parsed = _display_time(value)
            if parsed is None:
                sheet.write_blank(row, col, None)
            else:
                sheet.write_datetime(row, col, parsed, formats.datetime())
            return
        if value is None or value == "":
            sheet.write_blank(row, col, None, formats.text())
            return
        sheet.write_string(row, col, str(value), formats.text())

    def _ordered_pairs(self, task, source_fields, target_fields) -> list:
        snapshot = task.get("config_snapshot")
        if isinstance(snapshot, str):
            snapshot = _json_object(snapshot)
        rules = snapshot.get("rules") if isinstance(snapshot, dict) else None
        ordered = [rule for rule in (rules or []) if isinstance(rule, dict)]
        ordered.sort(key=lambda rule: (0 if rule.get("critical") else 1, -float(rule.get("weight") or 0)))
        available_source = set(map(str, source_fields))
        available_target = set(map(str, target_fields))
        pairs: list[dict] = []
        for rule in ordered:
            source_side = rule.get("source") if isinstance(rule.get("source"), dict) else {}
            target_side = rule.get("target") if isinstance(rule.get("target"), dict) else {}
            source_cols = [field for field in (str(item) for item in source_side.get("fields") or []) if field in available_source]
            target_cols = [field for field in (str(item) for item in target_side.get("fields") or []) if field in available_target]
            if not source_cols and not target_cols:
                continue
            if not (source_cols and target_cols):
                continue
            available_source.difference_update(source_cols)
            available_target.difference_update(target_cols)
            source_label = " / ".join(source_cols)
            target_label = " / ".join(target_cols)
            label = source_label if source_label == target_label else f"{source_label} ↔ {target_label}"
            pairs.append({"source": source_cols, "target": target_cols, "label": label})
        return pairs

    def _xl_group_headers(self, sheet, formats: "_Formats", groups: list) -> dict:
        # constant_memory flushes a row once the next row is written: place ALL
        # row-0 merges before ANY row-1 header cells.
        column = 0
        ranges: dict = {}
        layout: list[tuple[str, int, int, str, list]] = []
        for group_name, headers, fill_key in groups:
            if not headers:
                continue
            start, end = column, column + len(headers) - 1
            layout.append((group_name, start, end, fill_key, headers))
            ranges[group_name] = (start, end)
            column = end + 1
        for group_name, start, end, fill_key, _headers in layout:
            if start == end:
                sheet.write(0, start, group_name, formats.group(fill_key))
            else:
                sheet.merge_range(0, start, 0, end, group_name, formats.group(fill_key))
        for _group_name, start, _end, _fill_key, headers in layout:
            for offset, header in enumerate(headers):
                sheet.write(1, start + offset, str(header), formats.column_header())
        sheet.set_row(0, 20)
        sheet.set_row(1, 20)
        sheet.freeze_panes(2, 0)
        return ranges

    def _autofit(self, sheet, widths: dict, max_column: int, last_row: int = 1) -> None:
        for index, width in widths.items():
            sheet.set_column(index, index, width)
        if max_column:
            sheet.autofilter(1, 0, max(1, last_row), max_column - 1)

    def export_task(self, task_id: str, task: Mapping[str, object], *, unresolved_review: int) -> dict[str, object]:
        profile = self.export_profile
        with self.meta.connect() as connection:
            items = [dict(row) for row in connection.execute(
                "SELECT * FROM match_items WHERE task_id=? ORDER BY COALESCE(source_row_number, CAST(source_row_id AS INTEGER)), CAST(source_row_id AS INTEGER)",
                (task_id,),
            ).fetchall()]
            candidates = [dict(row) for row in connection.execute(
                "SELECT * FROM match_candidates WHERE task_id=? ORDER BY CAST(source_row_id AS INTEGER), rank",
                (task_id,),
            ).fetchall()]
            reviews = [dict(row) for row in connection.execute(
                "SELECT * FROM reviews WHERE task_id=? ORDER BY created_at",
                (task_id,),
            ).fetchall()]
        for item in items:
            item["source_payload"] = _json_object(item.get("source_payload"))
        for candidate in candidates:
            candidate["target_payload"] = _json_object(candidate.get("target_payload"))
            candidate["field_scores"] = _json_list(candidate.get("field_scores"))
        item_by_source = {str(item.get("source_row_id")): item for item in items}
        candidate_by_source = self._candidate_map(candidates)
        latest_review = self._latest_reviews(reviews)
        source_fields = _ordered_fields(item["source_payload"] for item in items)
        target_fields = _ordered_fields(candidate["target_payload"] for candidate in candidates)
        rules = self._rules(task)
        generated_at = _now()
        asset_summary = self.task_input_assets.describe(task_id)
        source_asset = asset_summary.get("source") if isinstance(asset_summary.get("source"), dict) else {}
        target_asset = asset_summary.get("target") if isinstance(asset_summary.get("target"), dict) else {}
        source_file_name = str(source_asset.get("original_name") or "")
        target_file_name = str(target_asset.get("original_name") or "")
        counts = {"MATCHED": 0, "CONFIRMED": 0, "REVIEW": 0, "UNMATCHED": 0}
        for item in items:
            status = str(item.get("current_status") or "")
            if status in counts:
                counts[status] += 1

        import io as _io
        import xlsxwriter as _XL
        output = _io.BytesIO()
        workbook = _XL.Workbook(output, {"constant_memory": True, "use_zip64": True})
        formats = self._Formats(workbook, profile)

        summary = workbook.add_worksheet(profile.sheet("summary"))
        summary.hide_gridlines(2)
        summary.merge_range(0, 0, 0, 3, profile.label("summary_title"), formats.summary_title())
        started_by = self._task_actor(task_id, task, "started")
        scheme_name = resolve_task_scheme_name(self.meta, dict(task))
        business_digits = "".join(ch for ch in str(generated_at) if ch.isdigit())[:14]
        business_stamp = f"{business_digits[:8]}_{business_digits[8:14]}"
        timing = task_time_fields(self.meta, dict(task))
        summary_rows = [
            (profile.summary_label("scheme_name"), scheme_name, False, profile.summary_label("source_total"), len(items), False),
            (profile.summary_label("automatic_matches"), counts["MATCHED"], False, profile.summary_label("manual_matches"), counts["CONFIRMED"], False),
            (profile.summary_label("unmatched_count"), counts["UNMATCHED"], False, profile.summary_label("pending_count"), counts["REVIEW"], False),
            (profile.summary_label("started_by"), started_by, False, profile.summary_label("started_at"), timing.get("started_at"), True),
            (profile.summary_label("compute_duration"), format_duration_ms(timing.get("compute_duration_ms")), False, profile.summary_label("compute_completed_at"), timing.get("compute_completed_at"), True),
            (profile.summary_label("generated_at"), generated_at, True, profile.summary_label("source_file"), source_file_name, False),
            (profile.summary_label("target_file"), target_file_name, False, profile.summary_label("pending_records"), unresolved_review, False),
        ]
        for row_index, values in enumerate(summary_rows, start=2):
            summary.write(row_index, 0, values[0], formats.summary_key())
            summary.write(row_index, 2, values[3], formats.summary_key())
            self._xl_write(summary, row_index, 1, values[1], is_time=bool(values[2]), formats=formats)
            self._xl_write(summary, row_index, 3, values[4], is_time=bool(values[5]), formats=formats)
        summary.write(11, 0, profile.label("legend_title"), formats.summary_key())
        for offset, (label, color_key) in enumerate([
            (profile.legend_label("exact"), "exact_fill"),
            (profile.legend_label("partial"), "partial_fill"),
            (profile.legend_label("different"), "different_fill"),
            (profile.legend_label("missing"), "missing_fill"),
        ]):
            summary.write(12, offset, label, formats.legend(color_key))
        for column, width in profile.summary_column_widths.items():
            index = ord(str(column).upper()[0]) - ord("A")
            summary.set_column(index, index, width)
        summary.set_row(0, 22)

        result_sheet = workbook.add_worksheet(profile.sheet("final_result"))
        result_sheet.hide_gridlines(2)
        pairs = self._ordered_pairs(task, source_fields, target_fields)
        assigned_source = {field for pair in pairs for field in pair["source"]}
        assigned_target = {field for pair in pairs for field in pair["target"]}
        leftover_source = [field for field in source_fields if field not in assigned_source]
        leftover_target = [field for field in target_fields if field not in assigned_target]

        pair_headers: list[str] = []
        pair_cells: list[tuple[str, str]] = []
        for pair in pairs:
            for field in pair["source"]:
                pair_headers.append(f"源·{field}")
                pair_cells.append(("source", field))
            for field in pair["target"]:
                pair_headers.append(f"候选·{field}")
                pair_cells.append(("target", field))
        overview_name = profile.group("result")
        groups: list[tuple[str, list[str], str]] = [
            (overview_name, [profile.header("source_row_number"), profile.header("status"), profile.header("final_group_code"), profile.header("similarity"), profile.header("target_row_number"), profile.header("match_method"), profile.header("operator"), profile.header("match_time")], "group_result_fill"),
        ]
        if pair_headers:
            groups.append(("映射依据（按权重成对）", pair_headers, "group_source_fill"))
        if leftover_source:
            groups.append((profile.group("source"), leftover_source, "group_target_fill"))
        if leftover_target:
            groups.append(("完整候选数据", leftover_target, "group_target_fill"))
        ranges = self._xl_group_headers(result_sheet, formats, groups)
        ov = ranges[overview_name][0]
        pair_start = ranges.get("映射依据（按权重成对）", (None, None))[0]
        src_zone = ranges.get(profile.group("source"), (None, None))[0]
        tgt_zone = ranges.get("完整候选数据", (None, None))[0]
        for excel_row, item in enumerate(items, start=2):
            source_payload = item["source_payload"]
            source_key = str(item.get("source_row_id") or "")
            row_candidates = candidate_by_source.get(source_key, [])
            selected = self._selected_candidate(item, row_candidates)
            target_payload = selected.get("target_payload", {}) if selected else {}
            review = latest_review.get(source_key)
            status = str(item.get("current_status") or "")
            match_time = review.get("created_at") if review else item.get("created_at") if status in {"MATCHED", "UNMATCHED"} else None
            operator = review.get("operator") if review else None
            similarity = selected.get("score") if selected else item.get("top1_score")
            result_sheet.write_number(excel_row, ov, float(item.get("source_row_number") or 0) if str(item.get("source_row_number") or "").isdigit() else 0, formats.text())
            result_sheet.write_string(excel_row, ov + 1, profile.status_label(status), formats.status(profile.status_color(status)))
            result_sheet.write_string(excel_row, ov + 2, self._cell_text(item.get("final_group_code")), formats.text())
            if isinstance(similarity, (int, float)):
                result_sheet.write_number(excel_row, ov + 3, float(similarity), formats.number())
            else:
                result_sheet.write_blank(excel_row, ov + 3, None, formats.text())
            target_row_number = selected.get("target_row_number") if selected else None
            if isinstance(target_row_number, (int, float)):
                result_sheet.write_number(excel_row, ov + 4, float(target_row_number), formats.text())
            else:
                result_sheet.write_blank(excel_row, ov + 4, None, formats.text())
            result_sheet.write_string(excel_row, ov + 5, profile.method_label(status, reviewed=bool(review)), formats.text())
            result_sheet.write_string(excel_row, ov + 6, self._cell_text(operator), formats.text())
            self._xl_write(result_sheet, excel_row, ov + 7, match_time, is_time=True, formats=formats)
            if pair_start is not None:
                source_states, target_states = self._field_states(source_payload, target_payload, selected, rules)
                column = pair_start
                for side, field in pair_cells:
                    state = (source_states if side == "source" else target_states).get(field)
                    result_sheet.write_string(excel_row, column, self._cell_text((source_payload if side == "source" else target_payload).get(field)), formats.text(state))
                    column += 1
            if src_zone is not None:
                for offset, field in enumerate(leftover_source):
                    result_sheet.write_string(excel_row, src_zone + offset, self._cell_text(source_payload.get(field)), formats.text())
            if tgt_zone is not None:
                for offset, field in enumerate(leftover_target):
                    result_sheet.write_string(excel_row, tgt_zone + offset, self._cell_text(target_payload.get(field)), formats.text())
        def _result_width(header: str) -> float:
            name = str(header)
            if name.endswith("行号"):
                return 9
            if "集团码" in name or "编码" in name:
                return 15
            if "相似度" in name:
                return 8
            if "状态" in name or "方式" in name:
                return 11
            if "账号" in name:
                return 10
            if "时间" in name or "日期" in name:
                return 16
            if "名称" in name or "描述" in name:
                return 26
            if "规格" in name or "型号" in name:
                return 20
            return 13
        widths: dict[int, float] = {}
        column = 0
        for _group, headers, _fill in groups:
            for header in headers:
                widths[column] = _result_width(header)
                column += 1
        self._autofit(result_sheet, widths, column, last_row=2 + max(0, len(items) - 1))

        top5 = workbook.add_worksheet(profile.sheet("top_candidates"))
        top5.hide_gridlines(2)
        top_groups = [
            (profile.group("source"), [profile.header("source_row_number"), profile.header("source_material_code"), *source_fields], "group_source_fill"),
            (profile.group("candidate_result"), [profile.header("candidate_rank"), profile.header("target_row_number"), profile.header("group_code"), profile.header("similarity")], "group_result_fill"),
            (profile.group("target"), target_fields, "group_target_fill"),
        ]
        top_ranges = self._xl_group_headers(top5, formats, top_groups)
        ts, tc, tt = top_ranges[profile.group("source")][0], top_ranges[profile.group("candidate_result")][0], top_ranges[profile.group("target")][0]
        candidate_excel_row = 2
        for candidate in candidates:
            if int(candidate.get("rank") or 999) > profile.candidate_top_n:
                continue
            item = item_by_source.get(str(candidate.get("source_row_id") or ""))
            if item is None:
                continue
            source_payload = item["source_payload"]
            target_payload = candidate.get("target_payload") or {}
            top5.write_number(candidate_excel_row, ts, float(item.get("source_row_number") or 0), formats.text())
            top5.write_string(candidate_excel_row, ts + 1, self._cell_text(item.get("source_id")), formats.text())
            source_states, target_states = self._field_states(source_payload, target_payload, candidate, rules)
            for offset, field in enumerate(source_fields):
                top5.write_string(candidate_excel_row, ts + 2 + offset, self._cell_text(source_payload.get(field)), formats.text(source_states.get(field)))
            top5.write_number(candidate_excel_row, tc, float(candidate.get("rank") or 0), formats.text())
            top5.write_number(candidate_excel_row, tc + 1, float(candidate.get("target_row_number") or 0), formats.text())
            top5.write_string(candidate_excel_row, tc + 2, self._cell_text(candidate.get("target_group_code")), formats.text())
            if isinstance(candidate.get("score"), (int, float)):
                top5.write_number(candidate_excel_row, tc + 3, float(candidate["score"]), formats.number())
            else:
                top5.write_blank(candidate_excel_row, tc + 3, None, formats.text())
            for offset, field in enumerate(target_fields):
                top5.write_string(candidate_excel_row, tt + offset, self._cell_text(target_payload.get(field)), formats.text(target_states.get(field)))
            candidate_excel_row += 1
        self._autofit(top5, {}, ts + 2 + len(source_fields) + 4 + len(target_fields) + 1, last_row=candidate_excel_row - 1)

        audit = workbook.add_worksheet(profile.sheet("audit"))
        audit.hide_gridlines(2)
        audit_groups = [
            (profile.group("audit_source"), [profile.header("source_row_number"), profile.header("source_material_code")], "group_source_fill"),
            (profile.group("audit_action"), [profile.header("action"), profile.header("target_row_number"), profile.header("group_code"), profile.header("operator"), profile.header("comment"), profile.header("action_time")], "group_audit_fill"),
        ]
        audit_ranges = self._xl_group_headers(audit, formats, audit_groups)
        asc, aac = audit_ranges[profile.group("audit_source")][0], audit_ranges[profile.group("audit_action")][0]
        for excel_row, review in enumerate(reviews, start=2):
            source_key = str(review.get("source_row_id") or "")
            item = item_by_source.get(source_key)
            row_candidates = candidate_by_source.get(source_key, [])
            audit.write_number(excel_row, asc, float(item.get("source_row_number") or 0) if item else 0, formats.text())
            audit.write_string(excel_row, asc + 1, self._cell_text(item.get("source_id")) if item else "", formats.text())
            audit.write_string(excel_row, aac, profile.action_label(review.get("action")), formats.text())
            target_row = self._review_target_row(review, row_candidates)
            if isinstance(target_row, (int, float)):
                audit.write_number(excel_row, aac + 1, float(target_row), formats.text())
            else:
                audit.write_blank(excel_row, aac + 1, None, formats.text())
            audit.write_string(excel_row, aac + 2, self._cell_text(review.get("selected_group_code")), formats.text())
            audit.write_string(excel_row, aac + 3, self._cell_text(review.get("operator")), formats.text())
            audit.write_string(excel_row, aac + 4, self._cell_text(review.get("comment")), formats.text())
            self._xl_write(audit, excel_row, aac + 5, review.get("created_at"), is_time=True, formats=formats)
        self._autofit(audit, {}, aac + 6)

        unmatched = workbook.add_worksheet(profile.sheet("unmatched"))
        unmatched.hide_gridlines(2)
        unmatched_groups = [
            (profile.group("source"), [profile.header("source_row_number"), *source_fields], "group_source_fill"),
            (profile.group("unmatched_reference"), [profile.header("unmatched_status"), profile.header("first_target_row"), profile.header("first_group_code"), profile.header("first_similarity"), profile.header("last_operator"), profile.header("last_action_time")], "group_result_fill"),
        ]
        unmatched_ranges = self._xl_group_headers(unmatched, formats, unmatched_groups)
        usc, urc = unmatched_ranges[profile.group("source")][0], unmatched_ranges[profile.group("unmatched_reference")][0]
        unmatched_excel_row = 2
        for item in items:
            if str(item.get("current_status") or "") != "UNMATCHED":
                continue
            source_key = str(item.get("source_row_id") or "")
            row_candidates = candidate_by_source.get(source_key, [])
            top1 = row_candidates[0] if row_candidates else None
            review = latest_review.get(source_key)
            unmatched.write_number(unmatched_excel_row, usc, float(item.get("source_row_number") or 0), formats.text())
            for offset, field in enumerate(source_fields):
                unmatched.write_string(unmatched_excel_row, usc + 1 + offset, self._cell_text(item["source_payload"].get(field)), formats.text())
            unmatched.write_string(unmatched_excel_row, urc, profile.status_label("UNMATCHED"), formats.status(profile.status_color("UNMATCHED")))
            if isinstance((top1 or {}).get("target_row_number"), (int, float)):
                unmatched.write_number(unmatched_excel_row, urc + 1, float(top1["target_row_number"]), formats.text())
            else:
                unmatched.write_blank(unmatched_excel_row, urc + 1, None, formats.text())
            unmatched.write_string(unmatched_excel_row, urc + 2, self._cell_text((top1 or {}).get("target_group_code")), formats.text())
            if isinstance((top1 or {}).get("score"), (int, float)):
                unmatched.write_number(unmatched_excel_row, urc + 3, float(top1["score"]), formats.number())
            else:
                unmatched.write_blank(unmatched_excel_row, urc + 3, None, formats.text())
            unmatched.write_string(unmatched_excel_row, urc + 4, self._cell_text((review or {}).get("operator")), formats.text())
            self._xl_write(unmatched, unmatched_excel_row, urc + 5, (review or {}).get("created_at"), is_time=True, formats=formats)
            unmatched_excel_row += 1
        self._autofit(unmatched, {}, usc + 1 + len(source_fields) + 6)

        workbook.close()
        output.seek(0)
        return self.files.save_stream(
            f"{safe_business_filename(profile.filename_prefix)}_{safe_business_filename(scheme_name)}_{business_stamp}.xlsx",
            "result",
            output,
            self.settings.max_total_upload_bytes,
        )
