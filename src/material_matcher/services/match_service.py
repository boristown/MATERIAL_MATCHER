from __future__ import annotations

from dataclasses import asdict
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
from typing import Any
import uuid

from openpyxl import Workbook

from material_matcher.domain.errors import DomainError
from material_matcher.domain.models import MatchingConfig
from material_matcher.matching.engine import RowResult, match_rows, summarize
from material_matcher.settings import Settings
from material_matcher.storage.files import FileRepository
from material_matcher.storage.metadata import MetadataRepository


def _now() -> str:
    return datetime.now().astimezone().isoformat()


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


class MatchService:
    def __init__(self, metadata: MetadataRepository, files: FileRepository, settings: Settings) -> None:
        self.meta=metadata; self.files=files; self.settings=settings

    def _catalog(self, version_id: str) -> dict[str, object]:
        with self.meta.connect() as connection:
            row=connection.execute("SELECT c.name, v.* FROM catalog_versions v JOIN catalogs c ON c.catalog_id=v.catalog_id WHERE v.version_id=?",(version_id,)).fetchone()
        if row is None: raise DomainError("CATALOG_NOT_FOUND","集团码目录版本不存在",status_code=404)
        return dict(row)

    def _draft_context(self, draft_id: str) -> tuple[dict[str, object], MatchingConfig, dict[str, object], dict[str, object]]:
        with self.meta.connect() as connection:
            row=connection.execute("SELECT * FROM task_drafts WHERE draft_id=?",(draft_id,)).fetchone()
        if row is None: raise DomainError("TASK_DRAFT_NOT_FOUND","任务草稿不存在",status_code=404)
        draft=dict(row); config=MatchingConfig.model_validate(json.loads(str(draft["config_document"] or "{}")))
        if not draft.get("source_file_id") or not draft.get("catalog_version_id") or not config.rules:
            raise DomainError("TASK_DRAFT_INCOMPLETE","请先完成数据选择和匹配规则配置",status_code=422)
        source=self.files.get(str(draft["source_file_id"])); catalog=self._catalog(str(draft["catalog_version_id"])); target=self.files.get(str(catalog["source_file_id"]))
        return draft, config, source, {**catalog,"file":target}

    def dry_run(self, draft_id: str, sample_rows: int = 100) -> dict[str, object]:
        _, config, source, catalog=self._draft_context(draft_id)
        rows=match_rows(Path(str(source["stored_path"])),Path(str(catalog["file"]["stored_path"])),config=config,group_code_column=str(catalog["group_code_column"]),max_target_rows=self.settings.baseline_max_target_rows,max_source_rows=sample_rows)
        return {"summary":summarize(rows),"rows":[self._row_response(row) for row in rows]}

    @staticmethod
    def _row_response(row: RowResult) -> dict[str, object]:
        return {
            "source_row_id":row.source_row_id,"source_id":row.source_id,"status":row.status,"final_group_code":row.final_group_code,
            "first_score":row.first_score,"second_score":row.second_score,"score_gap":row.score_gap,"critical_conflict":row.critical_conflict,
            "candidates":[asdict(candidate) for candidate in row.candidates],
        }

    def claim_next_task(self) -> str | None:
        with self.meta.connect() as connection:
            row=connection.execute("SELECT task_id FROM tasks WHERE status IN ('PENDING','RECOVERING') ORDER BY created_at LIMIT 1").fetchone()
            if row is None: return None
            task_id=str(row["task_id"])
            updated=connection.execute("UPDATE tasks SET status='PREPARING', started_at=COALESCE(started_at,?), error_code=NULL, error_message=NULL WHERE task_id=? AND status IN ('PENDING','RECOVERING')",(_now(),task_id)).rowcount
        return task_id if updated else None

    def execute_task(self, task_id: str) -> None:
        try:
            with self.meta.connect() as connection:
                row=connection.execute("SELECT * FROM tasks WHERE task_id=?",(task_id,)).fetchone()
            if row is None: raise DomainError("TASK_NOT_FOUND","任务不存在",status_code=404)
            task=dict(row); config=MatchingConfig.model_validate(json.loads(str(task["config_snapshot"]))); source=self.files.get(str(task["source_file_id"])); catalog=self._catalog(str(task["catalog_version_id"])); target=self.files.get(str(catalog["source_file_id"]))
            with self.meta.connect() as connection:
                connection.execute("DELETE FROM match_candidates WHERE task_id=?",(task_id,)); connection.execute("DELETE FROM match_items WHERE task_id=?",(task_id,)); connection.execute("UPDATE tasks SET status='RUNNING', progress=2 WHERE task_id=?",(task_id,))
            def progress(done:int,total:int)->None:
                pct=5.0 + 90.0*(done/max(total,1))
                with self.meta.connect() as connection:
                    connection.execute("UPDATE tasks SET processed_rows=?, total_rows=?, progress=? WHERE task_id=?",(done,total,min(95.0,pct),task_id))
            rows=match_rows(Path(str(source["stored_path"])),Path(str(target["stored_path"])),config=config,group_code_column=str(catalog["group_code_column"]),max_target_rows=self.settings.baseline_max_target_rows,on_progress=progress)
            now=_now()
            with self.meta.connect() as connection:
                for item in rows:
                    connection.execute("INSERT INTO match_items VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(task_id,item.source_row_id,item.source_id,_json(item.source_payload),item.status,item.status,item.candidates[0].group_code if item.candidates else None,item.first_score,item.second_score,item.score_gap,1 if item.critical_conflict else 0,item.final_group_code,now,now))
                    for candidate in item.candidates:
                        connection.execute("INSERT INTO match_candidates VALUES(?,?,?,?,?,?,?,?)",(task_id,item.source_row_id,candidate.rank,candidate.group_code,_json(candidate.target_payload),candidate.score,_json(candidate.field_scores),1 if candidate.critical_conflict else 0))
                review_count=connection.execute("SELECT COUNT(*) FROM match_items WHERE task_id=? AND current_status='REVIEW'",(task_id,)).fetchone()[0]
                stage="REVIEW" if review_count else "RESULT"
                connection.execute("UPDATE tasks SET status='COMPLETED', stage=?, progress=100, processed_rows=?, total_rows=?, finished_at=? WHERE task_id=?",(stage,len(rows),len(rows),now,task_id))
        except Exception as exc:
            code=exc.code if isinstance(exc,DomainError) else "INTERNAL_ERROR"; message=exc.message if isinstance(exc,DomainError) else str(exc)
            with self.meta.connect() as connection:
                connection.execute("UPDATE tasks SET status='FAILED', error_code=?, error_message=?, finished_at=? WHERE task_id=?",(code,message,_now(),task_id))
            if isinstance(exc,DomainError): return
            raise

    def summary(self, task_id: str) -> dict[str, int]:
        with self.meta.connect() as connection:
            rows=connection.execute("SELECT current_status, COUNT(*) count FROM match_items WHERE task_id=? GROUP BY current_status",(task_id,)).fetchall()
        counts={str(row["current_status"]):int(row["count"]) for row in rows}
        return {"pending_review":counts.get("REVIEW",0),"confirmed":counts.get("CONFIRMED",0),"unmatched":counts.get("UNMATCHED",0),"automatic_matched":counts.get("MATCHED",0)}

    def workbench_items(self, task_id: str, *, first_score_min: float|None=None, first_score_max:float|None=None, second_score_min:float|None=None, second_score_max:float|None=None, gap_min:float|None=None, gap_max:float|None=None, critical_conflict:bool|None=None, page:int=1, page_size:int=50) -> dict[str, object]:
        conditions=["task_id=?","current_status='REVIEW'"]; params:list[object]=[task_id]
        for column,low,high in (("top1_score",first_score_min,first_score_max),("second_score",second_score_min,second_score_max),("score_gap",gap_min,gap_max)):
            if low is not None: conditions.append(f"{column}>=?"); params.append(low)
            if high is not None: conditions.append(f"{column}<=?"); params.append(high)
        if critical_conflict is not None: conditions.append("critical_conflict=?"); params.append(1 if critical_conflict else 0)
        where=" AND ".join(conditions)
        with self.meta.connect() as connection:
            total=int(connection.execute(f"SELECT COUNT(*) FROM match_items WHERE {where}",params).fetchone()[0])
            rows=connection.execute(f"SELECT * FROM match_items WHERE {where} ORDER BY top1_score DESC, source_row_id LIMIT ? OFFSET ?",(*params,page_size,(page-1)*page_size)).fetchall()
        return {"total":total,"page":page,"page_size":page_size,"items":[self._decode_item(row) for row in rows]}

    @staticmethod
    def _decode_item(row: Any) -> dict[str, object]:
        result=dict(row); result["source_payload"]=json.loads(str(result["source_payload"])); result["critical_conflict"]=bool(result["critical_conflict"]); return result

    def candidates(self, task_id: str, source_row_id: str) -> list[dict[str, object]]:
        with self.meta.connect() as connection:
            rows=connection.execute("SELECT * FROM match_candidates WHERE task_id=? AND source_row_id=? ORDER BY rank",(task_id,source_row_id)).fetchall()
        result=[]
        for row in rows:
            item=dict(row); item["target_payload"]=json.loads(str(item["target_payload"])); item["field_scores"]=json.loads(str(item["field_scores"])); item["critical_conflict"]=bool(item["critical_conflict"]); result.append(item)
        return result

    def confirm(self, task_id: str, source_row_id: str, target_group_code: str, comment: str="") -> dict[str, object]:
        candidates=self.candidates(task_id,source_row_id)
        if target_group_code not in {str(item["target_group_code"]) for item in candidates}: raise DomainError("CANDIDATE_NOT_FOUND","所选候选不存在",status_code=404)
        return self._review_action(task_id,source_row_id,"CONFIRM_CANDIDATE",target_group_code,comment)

    def reject(self, task_id: str, source_row_id: str, comment: str="") -> dict[str, object]:
        return self._review_action(task_id,source_row_id,"REJECT_ALL",None,comment)

    def _review_action(self,task_id:str,source_row_id:str,action:str,group_code:str|None,comment:str)->dict[str,object]:
        with self.meta.connect() as connection:
            row=connection.execute("SELECT * FROM match_items WHERE task_id=? AND source_row_id=?",(task_id,source_row_id)).fetchone()
            if row is None: raise DomainError("MATCH_ITEM_NOT_FOUND","待处理记录不存在",status_code=404)
            if row["current_status"] != "REVIEW": raise DomainError("TASK_STATE_CONFLICT","该记录已经处理",status_code=409)
            new_status="CONFIRMED" if action=="CONFIRM_CANDIDATE" else "UNMATCHED"; now=_now()
            connection.execute("UPDATE match_items SET current_status=?, final_group_code=?, updated_at=? WHERE task_id=? AND source_row_id=?",(new_status,group_code,now,task_id,source_row_id))
            connection.execute("INSERT INTO reviews VALUES(?,?,?,?,?,?,?,?,?)",(uuid.uuid4().hex,task_id,source_row_id,str(row["original_status"]),group_code,action,"admin",comment,now))
            connection.execute("INSERT INTO audit_events VALUES(?,?,?,?,?,?)",(uuid.uuid4().hex,"match_item",f"{task_id}:{source_row_id}",action,_json({"selected_group_code":group_code,"comment":comment}),now))
        return {"source_row_id":source_row_id,"status":new_status,"final_group_code":group_code}

    def batch_confirm_top1(self,task_id:str,source_row_ids:list[str])->dict[str,object]:
        return self._batch(task_id,source_row_ids,confirm=True)
    def batch_reject(self,task_id:str,source_row_ids:list[str])->dict[str,object]:
        return self._batch(task_id,source_row_ids,confirm=False)
    def _batch(self,task_id:str,ids:list[str],*,confirm:bool)->dict[str,object]:
        success=[]; failed=[]
        for source_row_id in ids:
            try:
                if confirm:
                    candidates=self.candidates(task_id,source_row_id)
                    if not candidates: raise DomainError("NO_CANDIDATE","没有可确认候选",status_code=409)
                    self.confirm(task_id,source_row_id,str(candidates[0]["target_group_code"])); success.append(source_row_id)
                else:
                    self.reject(task_id,source_row_id); success.append(source_row_id)
            except DomainError as exc: failed.append({"source_row_id":source_row_id,"code":exc.code,"message":exc.message})
        return {"success":success,"failed":failed}

    def finalize(self,task_id:str,allow_unresolved_review:bool=False)->dict[str,object]:
        with self.meta.connect() as connection:
            task=connection.execute("SELECT * FROM tasks WHERE task_id=?",(task_id,)).fetchone()
            if task is None: raise DomainError("TASK_NOT_FOUND","任务不存在",status_code=404)
            if task["status"] != "COMPLETED" or task["stage"] not in {"REVIEW", "RESULT"}:
                raise DomainError("TASK_STATE_CONFLICT","比对计算尚未完成，暂不能生成最终结果",status_code=409)
            if task["result_file_id"]:
                return {"task_id":task_id,"result_file_id":task["result_file_id"],"unresolved_review":0,"summary":self.summary(task_id)}
            unresolved=int(connection.execute("SELECT COUNT(*) FROM match_items WHERE task_id=? AND current_status='REVIEW'",(task_id,)).fetchone()[0])
        if unresolved and not allow_unresolved_review:
            raise DomainError("TASK_STATE_CONFLICT",f"仍有 {unresolved} 条待人工处理记录，请确认是否保留为空后继续生成",status_code=409,details={"unresolved_review":unresolved})
        workbook=Workbook(); result_sheet=workbook.active; result_sheet.title="匹配结果"; result_sheet.append(["客户物料编码","状态","最终集团码","第一候选集团码","第一候选分","第二候选分","分差","关键字段冲突","配置SHA256"])
        with self.meta.connect() as connection:
            items=connection.execute("SELECT * FROM match_items WHERE task_id=? ORDER BY CAST(source_row_id AS INTEGER)",(task_id,)).fetchall(); candidates=connection.execute("SELECT * FROM match_candidates WHERE task_id=? ORDER BY CAST(source_row_id AS INTEGER), rank",(task_id,)).fetchall(); reviews=connection.execute("SELECT * FROM reviews WHERE task_id=? ORDER BY created_at",(task_id,)).fetchall()
        for item in items:
            final_code=item["final_group_code"] if item["current_status"] != "REVIEW" else None
            result_sheet.append([str(item["source_id"]),str(item["current_status"]),final_code,item["top1_group_code"],item["top1_score"],item["second_score"],item["score_gap"],"是" if item["critical_conflict"] else "否",str(task["config_sha256"])])
        cand_sheet=workbook.create_sheet("TopN候选"); cand_sheet.append(["客户行","排名","集团码","得分"])
        for candidate in candidates: cand_sheet.append([str(candidate["source_row_id"]),candidate["rank"],str(candidate["target_group_code"]),candidate["score"]])
        review_sheet=workbook.create_sheet("人工确认记录"); review_sheet.append(["客户行","动作","集团码","操作人","备注","时间"])
        for review in reviews: review_sheet.append([str(review["source_row_id"]),review["action"],review["selected_group_code"],review["operator"],review["comment"],review["created_at"]])
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value,str): cell.number_format="@"
        output=BytesIO(); workbook.save(output); output.seek(0); file_record=self.files.save_stream(f"material_matcher_{task_id}.xlsx","result",output,self.settings.max_total_upload_bytes)
        with self.meta.connect() as connection:
            connection.execute("UPDATE tasks SET result_file_id=?, stage='RESULT', status='COMPLETED' WHERE task_id=?",(file_record["file_id"],task_id))
        return {"task_id":task_id,"result_file_id":file_record["file_id"],"unresolved_review":unresolved,"summary":self.summary(task_id)}
