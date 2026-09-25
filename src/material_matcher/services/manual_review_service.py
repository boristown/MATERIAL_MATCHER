from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json
from typing import Any, BinaryIO
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from material_matcher.domain.errors import DomainError
from material_matcher.services.task_service import resolve_task_scheme_name, safe_business_filename
from material_matcher.storage.metadata import MetadataRepository


def _now() -> str:
    return datetime.now().astimezone().isoformat()


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


_DECISION_ACTIONS = {"MATCH", "REMATCH", "IMPORT_MATCH", "MARK_UNMATCHED"}
_CANCEL_ACTIONS = {"CANCEL_MATCH", "CANCEL_UNMATCHED"}


class ManualReviewService:
    """Own STEP3 manual decisions, offline collaboration and structured audit history.

    Legacy ``reviews`` rows are still appended for compatibility. The new
    ``match_operation_logs`` table is the authoritative business audit trail.
    """

    def __init__(self, repository: MetadataRepository) -> None:
        self.repo = repository

    @staticmethod
    def _decode_payload(value: object) -> dict[str, object]:
        if not value:
            return {}
        try:
            decoded = json.loads(str(value))
        except (TypeError, json.JSONDecodeError):
            return {}
        return decoded if isinstance(decoded, dict) else {}

    def _task(self, task_id: str) -> dict[str, object]:
        with self.repo.connect() as connection:
            row = connection.execute("SELECT * FROM tasks WHERE task_id=?", (task_id,)).fetchone()
        if row is None:
            raise DomainError("TASK_NOT_FOUND", "任务不存在", status_code=404)
        return dict(row)

    def _item(self, task_id: str, source_row_id: str) -> dict[str, object]:
        with self.repo.connect() as connection:
            row = connection.execute(
                "SELECT * FROM match_items WHERE task_id=? AND source_row_id=?",
                (task_id, source_row_id),
            ).fetchone()
        if row is None:
            raise DomainError("MATCH_ITEM_NOT_FOUND", "匹配记录不存在", status_code=404)
        result = dict(row)
        result["source_payload"] = self._decode_payload(result.get("source_payload"))
        return result

    def candidates(self, task_id: str, source_row_id: str) -> list[dict[str, object]]:
        with self.repo.connect() as connection:
            rows = connection.execute(
                "SELECT * FROM match_candidates WHERE task_id=? AND source_row_id=? ORDER BY rank",
                (task_id, source_row_id),
            ).fetchall()
        result: list[dict[str, object]] = []
        for row in rows:
            item = dict(row)
            item["target_payload"] = self._decode_payload(item.get("target_payload"))
            item["field_scores"] = self._decode_payload(item.get("field_scores"))
            item["critical_conflict"] = bool(item.get("critical_conflict"))
            result.append(item)
        return result

    def _candidate(self, task_id: str, source_row_id: str, *, group_code: str | None = None, rank: int | None = None) -> dict[str, object]:
        rows = self.candidates(task_id, source_row_id)
        for row in rows:
            if int(row.get("rank") or 0) > 5:
                continue
            if rank is not None and int(row.get("rank") or 0) == rank:
                return row
            if group_code is not None and str(row.get("target_group_code")) == str(group_code):
                return row
        raise DomainError("CANDIDATE_NOT_FOUND", "所选 Top5 候选不存在", status_code=404)

    @staticmethod
    def _baseline(item: dict[str, object]) -> tuple[str, str | None]:
        status = str(item.get("original_status") or "REVIEW")
        group_code = str(item.get("top1_group_code")) if status == "MATCHED" and item.get("top1_group_code") is not None else None
        return status, group_code

    def _latest_active_manual_decision(self, task_id: str, source_row_id: str, item: dict[str, object] | None = None) -> tuple[str, str | None] | None:
        with self.repo.connect() as connection:
            row = connection.execute(
                "SELECT operation_type, selected_group_code FROM match_operation_logs WHERE task_id=? AND source_row_id=? ORDER BY operated_at DESC, rowid DESC LIMIT 1",
                (task_id, source_row_id),
            ).fetchone()
            if row is not None:
                action = str(row["operation_type"])
                if action in _CANCEL_ACTIONS:
                    return None
                if action in {"MATCH", "REMATCH", "IMPORT_MATCH"}:
                    return ("MATCH", str(row["selected_group_code"]) if row["selected_group_code"] is not None else None)
                if action == "MARK_UNMATCHED":
                    return ("UNMATCHED", None)
            legacy = connection.execute(
                "SELECT action, selected_group_code FROM reviews WHERE task_id=? AND source_row_id=? ORDER BY created_at DESC, rowid DESC LIMIT 1",
                (task_id, source_row_id),
            ).fetchone()
        if legacy is not None:
            action = str(legacy["action"])
            if action in {"CONFIRM_CANDIDATE", "MATCH", "REMATCH", "IMPORT_MATCH"}:
                return ("MATCH", str(legacy["selected_group_code"]) if legacy["selected_group_code"] is not None else None)
            if action in {"REJECT_ALL", "MARK_UNMATCHED"}:
                return ("UNMATCHED", None)
            if action in _CANCEL_ACTIONS:
                return None
        current = item or self._item(task_id, source_row_id)
        if str(current.get("current_status")) == "CONFIRMED":
            return ("MATCH", str(current.get("final_group_code")) if current.get("final_group_code") is not None else None)
        return None

    def _append_history(
        self,
        connection: Any,
        *,
        item: dict[str, object],
        operator: str,
        operation_type: str,
        before_status: str,
        after_status: str,
        previous_group_code: str | None,
        selected_group_code: str | None,
        previous_target_row_number: int | None,
        target_row_number: int | None,
        source: str,
        comment: str,
    ) -> str:
        now = _now()
        operation_id = uuid.uuid4().hex
        connection.execute(
            """INSERT INTO match_operation_logs(
                operation_id,task_id,source_row_id,source_row_number,operator,operated_at,
                operation_type,before_status,after_status,previous_group_code,selected_group_code,
                previous_target_row_number,target_row_number,source,comment
            ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                operation_id,
                item["task_id"],
                item["source_row_id"],
                item.get("source_row_number"),
                operator or "system",
                now,
                operation_type,
                before_status,
                after_status,
                previous_group_code,
                selected_group_code,
                previous_target_row_number,
                target_row_number,
                source,
                comment,
            ),
        )
        connection.execute(
            "INSERT INTO reviews(review_id,task_id,source_row_id,original_status,selected_group_code,action,operator,comment,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
            (
                uuid.uuid4().hex,
                item["task_id"],
                item["source_row_id"],
                before_status,
                selected_group_code,
                operation_type,
                operator or "system",
                comment,
                now,
            ),
        )
        connection.execute(
            "INSERT INTO audit_events VALUES(?,?,?,?,?,?)",
            (
                uuid.uuid4().hex,
                "match_item",
                f"{item['task_id']}:{item['source_row_id']}",
                operation_type,
                _json(
                    {
                        "operator": operator or "system",
                        "source": source,
                        "before_status": before_status,
                        "after_status": after_status,
                        "previous_group_code": previous_group_code,
                        "selected_group_code": selected_group_code,
                        "previous_target_row_number": previous_target_row_number,
                        "target_row_number": target_row_number,
                        "comment": comment,
                    }
                ),
                now,
            ),
        )
        return operation_id

    def _target_row_for_group(self, task_id: str, source_row_id: str, group_code: str | None) -> int | None:
        if not group_code:
            return None
        with self.repo.connect() as connection:
            row = connection.execute(
                "SELECT target_row_number FROM match_candidates WHERE task_id=? AND source_row_id=? AND target_group_code=? ORDER BY rank LIMIT 1",
                (task_id, source_row_id, group_code),
            ).fetchone()
        return int(row["target_row_number"]) if row is not None and row["target_row_number"] is not None else None

    def match(
        self,
        task_id: str,
        source_row_id: str,
        target_group_code: str,
        *,
        operator: str,
        source: str = "WEB",
        comment: str = "",
        operation_type: str | None = None,
    ) -> dict[str, object]:
        self._task(task_id)
        item = self._item(task_id, source_row_id)
        candidate = self._candidate(task_id, source_row_id, group_code=target_group_code)
        active = self._latest_active_manual_decision(task_id, source_row_id, item)
        selected = str(candidate["target_group_code"])
        if active == ("MATCH", selected):
            return {"source_row_id": source_row_id, "status": str(item["current_status"]), "final_group_code": selected, "idempotent": True}
        before_status = str(item["current_status"])
        previous_group = str(item["final_group_code"]) if item.get("final_group_code") is not None else None
        previous_target = self._target_row_for_group(task_id, source_row_id, previous_group)
        action = operation_type or ("REMATCH" if before_status == "CONFIRMED" and previous_group and previous_group != selected else ("IMPORT_MATCH" if source == "EXCEL" else "MATCH"))
        now = _now()
        with self.repo.connect() as connection:
            connection.execute(
                "UPDATE match_items SET current_status='CONFIRMED', final_group_code=?, updated_at=? WHERE task_id=? AND source_row_id=?",
                (selected, now, task_id, source_row_id),
            )
            operation_id = self._append_history(
                connection,
                item=item,
                operator=operator,
                operation_type=action,
                before_status=before_status,
                after_status="CONFIRMED",
                previous_group_code=previous_group,
                selected_group_code=selected,
                previous_target_row_number=previous_target,
                target_row_number=int(candidate["target_row_number"]) if candidate.get("target_row_number") is not None else None,
                source=source,
                comment=comment,
            )
        return {"source_row_id": source_row_id, "status": "CONFIRMED", "final_group_code": selected, "operation_id": operation_id, "idempotent": False}

    def mark_unmatched(self, task_id: str, source_row_id: str, *, operator: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        self._task(task_id)
        item = self._item(task_id, source_row_id)
        active = self._latest_active_manual_decision(task_id, source_row_id, item)
        if active == ("UNMATCHED", None):
            return {"source_row_id": source_row_id, "status": "UNMATCHED", "final_group_code": None, "idempotent": True}
        before_status = str(item["current_status"])
        previous_group = str(item["final_group_code"]) if item.get("final_group_code") is not None else None
        previous_target = self._target_row_for_group(task_id, source_row_id, previous_group)
        now = _now()
        with self.repo.connect() as connection:
            connection.execute(
                "UPDATE match_items SET current_status='UNMATCHED', final_group_code=NULL, updated_at=? WHERE task_id=? AND source_row_id=?",
                (now, task_id, source_row_id),
            )
            operation_id = self._append_history(
                connection,
                item=item,
                operator=operator,
                operation_type="MARK_UNMATCHED",
                before_status=before_status,
                after_status="UNMATCHED",
                previous_group_code=previous_group,
                selected_group_code=None,
                previous_target_row_number=previous_target,
                target_row_number=None,
                source=source,
                comment=comment,
            )
        return {"source_row_id": source_row_id, "status": "UNMATCHED", "final_group_code": None, "operation_id": operation_id, "idempotent": False}

    def _cancel(self, task_id: str, source_row_id: str, *, operator: str, kind: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        self._task(task_id)
        item = self._item(task_id, source_row_id)
        active = self._latest_active_manual_decision(task_id, source_row_id, item)
        expected = "MATCH" if kind == "match" else "UNMATCHED"
        if active is None or active[0] != expected:
            raise DomainError("TASK_STATE_CONFLICT", "当前记录没有可撤销的对应人工判断", status_code=409)
        before_status = str(item["current_status"])
        previous_group = str(item["final_group_code"]) if item.get("final_group_code") is not None else None
        previous_target = self._target_row_for_group(task_id, source_row_id, previous_group)
        after_status, after_group = self._baseline(item)
        action = "CANCEL_MATCH" if kind == "match" else "CANCEL_UNMATCHED"
        now = _now()
        with self.repo.connect() as connection:
            connection.execute(
                "UPDATE match_items SET current_status=?, final_group_code=?, updated_at=? WHERE task_id=? AND source_row_id=?",
                (after_status, after_group, now, task_id, source_row_id),
            )
            operation_id = self._append_history(
                connection,
                item=item,
                operator=operator,
                operation_type=action,
                before_status=before_status,
                after_status=after_status,
                previous_group_code=previous_group,
                selected_group_code=after_group,
                previous_target_row_number=previous_target,
                target_row_number=self._target_row_for_group(task_id, source_row_id, after_group),
                source=source,
                comment=comment,
            )
        return {"source_row_id": source_row_id, "status": after_status, "final_group_code": after_group, "operation_id": operation_id}

    def cancel_match(self, task_id: str, source_row_id: str, *, operator: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        return self._cancel(task_id, source_row_id, operator=operator, kind="match", source=source, comment=comment)

    def cancel_unmatched(self, task_id: str, source_row_id: str, *, operator: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        return self._cancel(task_id, source_row_id, operator=operator, kind="unmatched", source=source, comment=comment)

    def undo(self, task_id: str, source_row_id: str, *, operator: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        item = self._item(task_id, source_row_id)
        active = self._latest_active_manual_decision(task_id, source_row_id, item)
        if active is None:
            raise DomainError("TASK_STATE_CONFLICT", "当前记录没有可撤销的人工判断", status_code=409)
        if active[0] == "MATCH":
            return self.cancel_match(task_id, source_row_id, operator=operator, source=source, comment=comment)
        return self.cancel_unmatched(task_id, source_row_id, operator=operator, source=source, comment=comment)

    def rematch(self, task_id: str, source_row_id: str, target_group_code: str, *, operator: str, source: str = "WEB", comment: str = "") -> dict[str, object]:
        item = self._item(task_id, source_row_id)
        active = self._latest_active_manual_decision(task_id, source_row_id, item)
        if active is None or active[0] != "MATCH":
            raise DomainError("TASK_STATE_CONFLICT", "当前记录尚未人工匹配，不能执行重新匹配", status_code=409)
        return self.match(task_id, source_row_id, target_group_code, operator=operator, source=source, comment=comment, operation_type="REMATCH")

    def batch_confirm_top1(self, task_id: str, source_row_ids: list[str], *, operator: str) -> dict[str, object]:
        success: list[str] = []
        failed: list[dict[str, object]] = []
        for source_row_id in source_row_ids:
            try:
                candidate = self._candidate(task_id, source_row_id, rank=1)
                self.match(task_id, source_row_id, str(candidate["target_group_code"]), operator=operator)
                success.append(source_row_id)
            except DomainError as exc:
                failed.append({"source_row_id": source_row_id, "code": exc.code, "message": exc.message})
        return {"success": success, "failed": failed}

    def batch_mark_unmatched(self, task_id: str, source_row_ids: list[str], *, operator: str) -> dict[str, object]:
        success: list[str] = []
        failed: list[dict[str, object]] = []
        for source_row_id in source_row_ids:
            try:
                self.mark_unmatched(task_id, source_row_id, operator=operator)
                success.append(source_row_id)
            except DomainError as exc:
                failed.append({"source_row_id": source_row_id, "code": exc.code, "message": exc.message})
        return {"success": success, "failed": failed}

    def logs(self, task_id: str, *, source_row_id: str | None = None, limit: int = 200, offset: int = 0) -> dict[str, object]:
        self._task(task_id)
        conditions = ["task_id=?"]
        params: list[object] = [task_id]
        if source_row_id is not None:
            conditions.append("source_row_id=?")
            params.append(source_row_id)
        where = " AND ".join(conditions)
        with self.repo.connect() as connection:
            total = int(connection.execute(f"SELECT COUNT(*) FROM match_operation_logs WHERE {where}", params).fetchone()[0])
            rows = connection.execute(
                f"SELECT * FROM match_operation_logs WHERE {where} ORDER BY operated_at DESC, rowid DESC LIMIT ? OFFSET ?",
                (*params, limit, offset),
            ).fetchall()
        return {"total": total, "items": [dict(row) for row in rows]}

    @staticmethod
    def _business_keys(rows: list[dict[str, object]], payload_key: str) -> list[str]:
        keys: set[str] = set()
        for row in rows:
            payload = row.get(payload_key)
            if isinstance(payload, dict):
                keys.update(str(key) for key in payload if not str(key).startswith("__"))
        return sorted(keys)

    @staticmethod
    def _ordered_business_keys(
        rows: list[dict[str, object]],
        payload_key: str,
        document: dict[str, object],
        side: str,
    ) -> list[str]:
        keys = ManualReviewService._business_keys(rows, payload_key)
        weights: dict[str, float] = {}
        for rule in document.get("rules") or []:
            if not isinstance(rule, dict):
                continue
            rank = float(rule.get("weight") or 0) + (1_000_000 if rule.get("critical") else 0)
            for field in ((rule.get(side) or {}).get("fields") if isinstance(rule.get(side), dict) else []) or []:
                name = str(field)
                weights[name] = max(weights.get(name, -1.0), rank)
        return sorted(keys, key=lambda key: (-weights.get(key, -1.0), keys.index(key)))

    @staticmethod
    def _ordered_from_keys(base_keys: list[str], document: dict[str, object], side: str) -> list[str]:
        weights: dict[str, float] = {}
        for rule in document.get("rules") or []:
            if not isinstance(rule, dict):
                continue
            rank = float(rule.get("weight") or 0) + (1_000_000 if rule.get("critical") else 0)
            side_obj = rule.get(side)
            for field in (side_obj.get("fields") if isinstance(side_obj, dict) else []) or []:
                name = str(field)
                weights[name] = max(weights.get(name, -1.0), rank)
        return sorted(base_keys, key=lambda key: (-weights.get(key, -1.0), base_keys.index(key)))

    def export_workbook(self, task_id: str) -> BytesIO:
        import xlsxwriter

        task = self._task(task_id)
        try:
            document = json.loads(str(task.get("config_snapshot") or "{}"))
        except Exception:  # noqa: BLE001
            document = {}

        source_key_set: set[str] = set()
        target_key_set: set[str] = set()
        with self.repo.connect() as connection:
            for (raw,) in connection.execute("SELECT source_payload FROM match_items WHERE task_id=?", (task_id,)):
                for key in self._decode_payload(raw):
                    if not key.startswith("__"):
                        source_key_set.add(key)
            for (raw,) in connection.execute("SELECT target_payload FROM match_candidates WHERE task_id=? AND rank<=5", (task_id,)):
                for key in self._decode_payload(raw):
                    if not key.startswith("__"):
                        target_key_set.add(key)
        source_keys = self._ordered_from_keys(sorted(source_key_set), document, "source")
        target_keys = self._ordered_from_keys(sorted(target_key_set), document, "target")

        headers = ["源表原始行号"] + [f"源.{key}" for key in source_keys]
        headers.append("人工选择")
        for rank in range(1, 6):
            headers.extend([f"候选{rank}.目标表原始行号", f"候选{rank}.集团码", f"候选{rank}.相似度"])
            headers.extend(f"候选{rank}.{key}" for key in target_keys)
        headers.extend(["__task_id", "__source_row_id"])
        selection_col = headers.index("人工选择")
        task_col = headers.index("__task_id")
        source_row_id_col = headers.index("__source_row_id")

        output = BytesIO()
        workbook = xlsxwriter.Workbook(output, {"constant_memory": True, "use_zip64": True})
        header_format = workbook.add_format({"bold": True, "font_color": "FFFFFF", "bg_color": "1F4E79", "align": "center", "valign": "vcenter"})
        source_header_format = workbook.add_format({"bold": True, "font_color": "FFFFFF", "bg_color": "2E75B6", "align": "center", "valign": "vcenter"})
        candidate_header_format = workbook.add_format({"bold": True, "font_color": "FFFFFF", "bg_color": "548235", "align": "center", "valign": "vcenter"})
        selection_header_format = workbook.add_format({"bold": True, "font_color": "7F5F00", "bg_color": "FFD966", "align": "center", "valign": "vcenter"})
        text_format = workbook.add_format({"num_format": "@"})
        score_format = workbook.add_format({"num_format": "0.0"})
        selection_body_format = workbook.add_format({"bg_color": "FFF2CC", "border": 1, "border_color": "BF8F00"})
        sheet = workbook.add_worksheet("人工匹配")
        for index, header in enumerate(headers):
            name = str(header)
            if name == "人工选择":
                fmt = selection_header_format
            elif name.startswith("候选"):
                fmt = candidate_header_format
            else:
                fmt = source_header_format
            sheet.write(0, index, name, fmt)
        sheet.freeze_panes(1, 2)
        def _width_for(header: str) -> int:
            name = str(header)
            if name == "人工选择":
                return 16
            if "相似度" in name:
                return 9
            if name.endswith("行号"):
                return 9
            if "集团码" in name or "编码" in name:
                return 15
            if "名称" in name or "描述" in name:
                return 26
            if "规格" in name or "型号" in name:
                return 20
            return 12
        for index, header in enumerate(headers):
            if index in (task_col, source_row_id_col):
                sheet.set_column(index, index, 10, None, {"hidden": True})
            else:
                sheet.set_column(index, index, _width_for(header))
        with self.repo.connect() as count_connection:
            total_items = int(count_connection.execute(
                "SELECT COUNT(*) FROM match_items WHERE task_id=?", (task_id,),
            ).fetchone()[0])
        if total_items:
            sheet.data_validation(1, selection_col, total_items, selection_col, {
                "validate": "list",
                "source": '"候选1,候选2,候选3,候选4,候选5,均不匹配"',
                "allow_blank": True,
            })

        def write_cell(row: int, column: int, value: object) -> None:
            if value is None or value == "":
                sheet.write_blank(row, column, None, text_format)
            elif isinstance(value, bool):
                sheet.write_string(row, column, str(value), text_format)
            elif isinstance(value, (int, float)):
                sheet.write_number(row, column, float(value))
            else:
                sheet.write_string(row, column, str(value), text_format)

        with self.repo.connect() as connection:
            cursor = connection.execute(
                "SELECT * FROM match_items WHERE task_id=? ORDER BY COALESCE(source_row_number, 2147483647), CAST(source_row_id AS INTEGER), source_row_id",
                (task_id,),
            )
            row_index = 1
            while True:
                batch = cursor.fetchmany(500)
                if not batch:
                    break
                batch_ids = [str(row["source_row_id"]) for row in batch]
                placeholders = ",".join("?" * len(batch_ids))
                candidates_by_source: dict[str, dict[int, dict[str, object]]] = {}
                for candidate_row in connection.execute(
                    f"SELECT * FROM match_candidates WHERE task_id=? AND rank<=5 AND source_row_id IN ({placeholders}) ORDER BY source_row_id, rank",
                    (task_id, *batch_ids),
                ).fetchall():
                    candidate = dict(candidate_row)
                    candidate["target_payload"] = self._decode_payload(candidate.get("target_payload"))
                    candidates_by_source.setdefault(str(candidate["source_row_id"]), {})[int(candidate["rank"])] = candidate
                for item_row in batch:
                    item = dict(item_row)
                    payload = self._decode_payload(item.get("source_payload"))
                    column = 0
                    write_cell(row_index, column, item.get("source_row_number")); column += 1
                    for key in source_keys:
                        write_cell(row_index, column, payload.get(key)); column += 1
                    per_row = candidates_by_source.get(str(item["source_row_id"]), {})
                    for rank in range(1, 6):
                        candidate = per_row.get(rank)
                        if candidate is None:
                            column += 3 + len(target_keys)
                        else:
                            write_cell(row_index, column, candidate.get("target_row_number")); column += 1
                            write_cell(row_index, column, candidate.get("target_group_code")); column += 1
                            score_value = candidate.get("score")
                            if isinstance(score_value, (int, float)):
                                sheet.write_number(row_index, column, float(score_value), score_format)
                            else:
                                write_cell(row_index, column, score_value)
                            column += 1
                            target_payload = candidate.get("target_payload") if isinstance(candidate.get("target_payload"), dict) else {}
                            for key in target_keys:
                                write_cell(row_index, column, target_payload.get(key)); column += 1
                    sheet.write_blank(row_index, selection_col, None, selection_body_format)
                    write_cell(row_index, task_col, task_id)
                    write_cell(row_index, source_row_id_col, str(item["source_row_id"]))
                    row_index += 1

        guide = workbook.add_worksheet("填写说明")
        guide.write(0, 0, "人工匹配 Excel")
        guide.write(1, 0, "方案名称")
        guide.write(1, 1, resolve_task_scheme_name(self.repo, dict(task)))
        guide.write(2, 0, "填写方式")
        guide.write(2, 1, "仅填写“人工选择”列；可选择候选1～候选5或“均不匹配”。空白行会被忽略。")
        guide.write(3, 0, "多人协作")
        guide.write(3, 1, "可复制或拆分本文件给多人处理，再分别上传；系统会幂等合并，相同结果安全跳过，不同结果返回冲突且不会覆盖。")
        guide.set_column(0, 0, 18)
        guide.set_column(1, 1, 90)
        workbook.close()
        output.seek(0)
        return output

    @staticmethod
    def _parse_selection(value: object) -> tuple[str, int | None] | None:
        if value is None or str(value).strip() == "":
            return None
        text = str(value).strip().replace(" ", "")
        if text in {"均不匹配", "不匹配", "无匹配", "NONE", "none"}:
            return ("UNMATCHED", None)
        if text.isdigit() and 1 <= int(text) <= 5:
            return ("MATCH", int(text))
        if text.startswith("候选") and text[2:].isdigit() and 1 <= int(text[2:]) <= 5:
            return ("MATCH", int(text[2:]))
        raise DomainError("INVALID_MANUAL_SELECTION", f"人工选择值不正确: {value}", status_code=422)

    def import_workbook(self, task_id: str, stream: BinaryIO, *, operator: str, filename: str = "") -> dict[str, object]:
        self._task(task_id)
        try:
            workbook = load_workbook(stream, data_only=True, read_only=True)
        except Exception as exc:
            raise DomainError("MANUAL_EXCEL_PARSE_FAILED", "人工匹配 Excel 无法解析", status_code=400) from exc
        if "人工匹配" not in workbook.sheetnames:
            workbook.close()
            raise DomainError("MANUAL_EXCEL_INVALID", "未找到“人工匹配”工作表", status_code=422)
        sheet = workbook["人工匹配"]
        row_iter = sheet.iter_rows(values_only=True)
        header_row = next(row_iter, None) or ()
        headers = {str(value).strip(): index for index, value in enumerate(header_row) if value is not None}
        required = {"人工选择", "__task_id", "__source_row_id"}
        if not required.issubset(headers):
            workbook.close()
            raise DomainError("MANUAL_EXCEL_INVALID", "人工匹配 Excel 缺少内部识别列，请使用系统下载的模板", status_code=422)

        def cell_value(cells: tuple, name: str) -> object:
            position = headers.get(name)
            if position is None or position >= len(cells):
                return None
            return cells[position]

        success = skipped = conflicts = errors = 0
        details: list[dict[str, object]] = []
        seen: dict[str, tuple[str, int | None]] = {}
        for excel_row, cells in enumerate(row_iter, start=2):
            selection_value = cell_value(cells, "人工选择")
            if selection_value is None or str(selection_value).strip() == "":
                skipped += 1
                continue
            try:
                parsed = self._parse_selection(selection_value)
                assert parsed is not None
                embedded_task_id = str(cell_value(cells, "__task_id") or "").strip()
                source_row_id = str(cell_value(cells, "__source_row_id") or "").strip()
                if embedded_task_id != task_id:
                    raise DomainError("MANUAL_EXCEL_TASK_MISMATCH", "Excel 中的任务与当前任务不一致", status_code=409)
                if not source_row_id:
                    raise DomainError("MANUAL_EXCEL_INVALID", "Excel 行缺少源记录标识", status_code=422)
                if source_row_id in seen and seen[source_row_id] != parsed:
                    conflicts += 1
                    details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "conflict", "code": "DUPLICATE_ROW_CONFLICT", "message": "同一文件对同一源记录给出了不同结果"})
                    continue
                seen[source_row_id] = parsed
                item = self._item(task_id, source_row_id)
                active = self._latest_active_manual_decision(task_id, source_row_id, item)
                desired: tuple[str, str | None]
                if parsed[0] == "UNMATCHED":
                    desired = ("UNMATCHED", None)
                else:
                    candidate = self._candidate(task_id, source_row_id, rank=int(parsed[1] or 0))
                    desired = ("MATCH", str(candidate["target_group_code"]))
                if active is not None:
                    if active == desired:
                        skipped += 1
                        details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "skipped", "code": "IDEMPOTENT", "message": "该人工结果已存在"})
                    else:
                        conflicts += 1
                        details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "conflict", "code": "MANUAL_RESULT_CONFLICT", "message": "该源记录已存在不同的人工结果，未覆盖", "existing": {"decision": active[0], "group_code": active[1]}, "incoming": {"decision": desired[0], "group_code": desired[1]}})
                    continue
                if desired[0] == "UNMATCHED":
                    self.mark_unmatched(task_id, source_row_id, operator=operator, source="EXCEL", comment=f"导入文件: {filename}" if filename else "Excel 上传")
                else:
                    self.match(task_id, source_row_id, str(desired[1]), operator=operator, source="EXCEL", comment=f"导入文件: {filename}" if filename else "Excel 上传")
                success += 1
            except DomainError as exc:
                errors += 1
                details.append({"excel_row": excel_row, "source_row_id": str(cell_value(cells, "__source_row_id") or ""), "type": "error", "code": exc.code, "message": exc.message})
            except Exception as exc:
                errors += 1
                details.append({"excel_row": excel_row, "source_row_id": str(cell_value(cells, "__source_row_id") or ""), "type": "error", "code": "MANUAL_IMPORT_ERROR", "message": str(exc)})
        workbook.close()
        return {
            "task_id": task_id,
            "success_count": success,
            "skipped_count": skipped,
            "conflict_count": conflicts,
            "error_count": errors,
            "details": details,
        }
