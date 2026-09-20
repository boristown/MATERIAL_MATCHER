from __future__ import annotations

from datetime import datetime
import json
import sqlite3
from typing import BinaryIO, Iterable
import uuid

from openpyxl import load_workbook

from material_matcher.domain.errors import DomainError
from material_matcher.storage.metadata import MetadataRepository


_ALLOWED_STATUSES = {"ALL", "MATCHED", "REVIEW", "CONFIRMED", "UNMATCHED"}
_ALLOWED_ACTIONS = {
    "CONFIRM_TOP1",
    "MARK_UNMATCHED",
    "CANCEL_MATCH",
    "CANCEL_UNMATCHED",
    "RESTORE_ALGORITHM",
}
_MATCH_ACTIONS = {"MATCH", "REMATCH", "IMPORT_MATCH", "CONFIRM_TOP1"}
_UNMATCH_ACTIONS = {"MARK_UNMATCHED"}
_CLEAR_ACTIONS = {"CANCEL_MATCH", "CANCEL_UNMATCHED", "RESTORE_ALGORITHM"}
_CHUNK_SIZE = 500


def _now() -> str:
    return datetime.now().astimezone().isoformat()


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _chunks(values: list[str], size: int = _CHUNK_SIZE) -> Iterable[list[str]]:
    for index in range(0, len(values), size):
        yield values[index : index + size]


class ReviewWorkbenchService:
    """Large-volume STEP3 query, selection, mutation, concurrency and Excel merge service.

    ``updated_at`` is the optimistic-concurrency version token. Every mutation
    performed here updates it in the same transaction as the business state and
    audit rows, so a stale client cannot silently overwrite another reviewer.
    """

    def __init__(self, repository: MetadataRepository) -> None:
        self.repo = repository

    @staticmethod
    def _decode(value: object) -> dict[str, object]:
        if not value:
            return {}
        try:
            decoded = json.loads(str(value))
        except (TypeError, json.JSONDecodeError):
            return {}
        return decoded if isinstance(decoded, dict) else {}

    def _ensure_task(self, connection: sqlite3.Connection, task_id: str) -> None:
        if connection.execute("SELECT 1 FROM tasks WHERE task_id=?", (task_id,)).fetchone() is None:
            raise DomainError("TASK_NOT_FOUND", "任务不存在", status_code=404)

    @staticmethod
    def _filter_sql(filters: dict[str, object]) -> tuple[str, list[object]]:
        status = str(filters.get("status") or "REVIEW").upper()
        if status not in _ALLOWED_STATUSES:
            raise DomainError("INVALID_REVIEW_STATUS", "人工工作台状态筛选不正确", status_code=422)
        conditions = ["m.task_id=?"]
        params: list[object] = [filters["task_id"]]
        if status != "ALL":
            conditions.append("m.current_status=?")
            params.append(status)
        q = str(filters.get("q") or "").strip()
        if q:
            like = f"%{q}%"
            conditions.append(
                "(m.source_id LIKE ? OR m.source_payload LIKE ? OR m.top1_group_code LIKE ? "
                "OR m.final_group_code LIKE ? OR EXISTS(SELECT 1 FROM match_candidates qmc "
                "WHERE qmc.task_id=m.task_id AND qmc.source_row_id=m.source_row_id AND qmc.rank<=5 "
                "AND (qmc.target_group_code LIKE ? OR qmc.target_payload LIKE ?)))"
            )
            params.extend([like, like, like, like, like, like])
        ranges = (
            ("m.top1_score", filters.get("first_score_min"), filters.get("first_score_max")),
            ("m.second_score", filters.get("second_score_min"), filters.get("second_score_max")),
            ("m.score_gap", filters.get("gap_min"), filters.get("gap_max")),
        )
        for column, low, high in ranges:
            if low is not None:
                conditions.append(f"{column}>=?")
                params.append(float(low))
            if high is not None:
                conditions.append(f"{column}<=?")
                params.append(float(high))
        if filters.get("critical_conflict") is not None:
            conditions.append("m.critical_conflict=?")
            params.append(1 if bool(filters["critical_conflict"]) else 0)
        return " AND ".join(conditions), params

    def list_items(
        self,
        task_id: str,
        *,
        status: str = "REVIEW",
        q: str | None = None,
        page: int = 1,
        page_size: int = 50,
        include_candidates: int = 0,
        first_score_min: float | None = None,
        first_score_max: float | None = None,
        second_score_min: float | None = None,
        second_score_max: float | None = None,
        gap_min: float | None = None,
        gap_max: float | None = None,
        critical_conflict: bool | None = None,
    ) -> dict[str, object]:
        if page < 1 or page_size < 1 or page_size > 200:
            raise DomainError("INVALID_PAGINATION", "分页参数不正确", status_code=422)
        if include_candidates < 0 or include_candidates > 5:
            raise DomainError("INVALID_CANDIDATE_LIMIT", "候选数量仅支持 0～5", status_code=422)
        filters: dict[str, object] = {
            "task_id": task_id,
            "status": status,
            "q": q,
            "first_score_min": first_score_min,
            "first_score_max": first_score_max,
            "second_score_min": second_score_min,
            "second_score_max": second_score_max,
            "gap_min": gap_min,
            "gap_max": gap_max,
            "critical_conflict": critical_conflict,
        }
        where, params = self._filter_sql(filters)
        offset = (page - 1) * page_size
        with self.repo.connect() as connection:
            self._ensure_task(connection, task_id)
            total = int(connection.execute(f"SELECT COUNT(*) FROM match_items m WHERE {where}", params).fetchone()[0])
            rows = connection.execute(
                f"""SELECT m.task_id,m.source_row_id,m.source_row_number,m.source_id,m.source_payload,
                           m.original_status,m.current_status,m.top1_group_code,m.top1_score,m.second_score,
                           m.score_gap,m.critical_conflict,m.final_group_code,m.created_at,m.updated_at
                    FROM match_items m WHERE {where}
                    ORDER BY m.top1_score DESC, COALESCE(m.source_row_number,2147483647), m.source_row_id
                    LIMIT ? OFFSET ?""",
                (*params, page_size, offset),
            ).fetchall()
            items = [dict(row) for row in rows]
            candidates = self._candidates_for(connection, task_id, [str(item["source_row_id"]) for item in items], include_candidates)
        for item in items:
            item["source_payload"] = self._decode(item.get("source_payload"))
            item["critical_conflict"] = bool(item.get("critical_conflict"))
            item["version"] = str(item["updated_at"])
            if include_candidates:
                item["candidates"] = candidates.get(str(item["source_row_id"]), [])
        return {
            "total": total,
            "page": page,
            "page_size": page_size,
            "status": str(status).upper(),
            "include_candidates": include_candidates,
            "items": items,
        }

    def get_item(self, task_id: str, source_row_id: str, *, include_candidates: int = 0) -> dict[str, object]:
        if include_candidates < 0 or include_candidates > 5:
            raise DomainError("INVALID_CANDIDATE_LIMIT", "候选数量仅支持 0～5", status_code=422)
        with self.repo.connect() as connection:
            self._ensure_task(connection, task_id)
            row = connection.execute(
                """SELECT m.task_id,m.source_row_id,m.source_row_number,m.source_id,m.source_payload,
                          m.original_status,m.current_status,m.top1_group_code,m.top1_score,m.second_score,
                          m.score_gap,m.critical_conflict,m.final_group_code,m.created_at,m.updated_at
                   FROM match_items m WHERE m.task_id=? AND m.source_row_id=?""",
                (task_id, source_row_id),
            ).fetchone()
            if row is None:
                raise DomainError("MATCH_ITEM_NOT_FOUND", "匹配记录不存在", status_code=404)
            item = dict(row)
            candidates = self._candidates_for(connection, task_id, [source_row_id], include_candidates)
        item["source_payload"] = self._decode(item.get("source_payload"))
        item["critical_conflict"] = bool(item.get("critical_conflict"))
        item["version"] = str(item["updated_at"])
        if include_candidates:
            item["candidates"] = candidates.get(source_row_id, [])
        return item

    def _candidates_for(
        self,
        connection: sqlite3.Connection,
        task_id: str,
        source_row_ids: list[str],
        limit: int,
    ) -> dict[str, list[dict[str, object]]]:
        result: dict[str, list[dict[str, object]]] = {}
        if not source_row_ids or limit <= 0:
            return result
        for chunk in _chunks(source_row_ids):
            marks = ",".join("?" for _ in chunk)
            rows = connection.execute(
                f"""SELECT task_id,source_row_id,rank,target_row_number,target_group_code,target_payload,
                           score,field_scores,critical_conflict,child_profile_id,child_profile_version,
                           child_profile_name,target_file_id,target_file_name
                    FROM match_candidates
                    WHERE task_id=? AND source_row_id IN ({marks}) AND rank<=?
                    ORDER BY source_row_id,rank""",
                (task_id, *chunk, limit),
            ).fetchall()
            for row in rows:
                item = dict(row)
                item["target_payload"] = self._decode(item.get("target_payload"))
                item["field_scores"] = self._decode(item.get("field_scores"))
                item["critical_conflict"] = bool(item.get("critical_conflict"))
                result.setdefault(str(item["source_row_id"]), []).append(item)
        return result

    def selection_count(self, task_id: str, selection: dict[str, object]) -> dict[str, object]:
        mode = str(selection.get("mode") or "explicit").lower()
        with self.repo.connect() as connection:
            self._ensure_task(connection, task_id)
            if mode == "filter":
                raw_filter = selection.get("filter")
                if not isinstance(raw_filter, dict):
                    raise DomainError("INVALID_SELECTION", "筛选选择缺少 filter", status_code=422)
                where, params = self._filter_sql({**raw_filter, "task_id": task_id})
                count = int(connection.execute(f"SELECT COUNT(*) FROM match_items m WHERE {where}", params).fetchone()[0])
                return {"mode": "filter", "count": count}
            if mode == "explicit":
                ids, _ = self._normalize_explicit(selection)
                found = 0
                for chunk in _chunks(ids):
                    marks = ",".join("?" for _ in chunk)
                    found += int(connection.execute(
                        f"SELECT COUNT(*) FROM match_items WHERE task_id=? AND source_row_id IN ({marks})",
                        (task_id, *chunk),
                    ).fetchone()[0])
                return {"mode": "explicit", "count": found, "requested_count": len(ids)}
        raise DomainError("INVALID_SELECTION", "selection.mode 仅支持 explicit 或 filter", status_code=422)

    @staticmethod
    def _normalize_explicit(selection: dict[str, object]) -> tuple[list[str], dict[str, str | None]]:
        raw = selection.get("items")
        if raw is None:
            raw = selection.get("ids")
        if not isinstance(raw, list) or not raw:
            raise DomainError("INVALID_SELECTION", "显式选择至少需要一条记录", status_code=422)
        ids: list[str] = []
        versions: dict[str, str | None] = {}
        for value in raw:
            if isinstance(value, str):
                source_row_id, version = value, None
            elif isinstance(value, dict):
                source_row_id = str(value.get("source_row_id") or "").strip()
                version_raw = value.get("expected_version")
                version = str(version_raw) if version_raw is not None else None
            else:
                raise DomainError("INVALID_SELECTION", "显式选择格式不正确", status_code=422)
            if not source_row_id:
                raise DomainError("INVALID_SELECTION", "显式选择缺少源记录标识", status_code=422)
            if source_row_id not in versions:
                ids.append(source_row_id)
            versions[source_row_id] = version
        return ids, versions

    def _selection_rows(
        self,
        connection: sqlite3.Connection,
        task_id: str,
        selection: dict[str, object],
    ) -> tuple[list[dict[str, object]], dict[str, str | None], str]:
        mode = str(selection.get("mode") or "explicit").lower()
        if mode == "explicit":
            ids, expected = self._normalize_explicit(selection)
            found: dict[str, dict[str, object]] = {}
            for chunk in _chunks(ids):
                marks = ",".join("?" for _ in chunk)
                rows = connection.execute(
                    f"SELECT * FROM match_items WHERE task_id=? AND source_row_id IN ({marks})",
                    (task_id, *chunk),
                ).fetchall()
                found.update({str(row["source_row_id"]): dict(row) for row in rows})
            rows = [found[source_row_id] for source_row_id in ids if source_row_id in found]
            return rows, expected, mode
        if mode == "filter":
            raw_filter = selection.get("filter")
            if not isinstance(raw_filter, dict):
                raise DomainError("INVALID_SELECTION", "筛选选择缺少 filter", status_code=422)
            filters = {**raw_filter, "task_id": task_id}
            where, params = self._filter_sql(filters)
            rows = connection.execute(
                f"SELECT m.* FROM match_items m WHERE {where} ORDER BY m.source_row_id",
                params,
            ).fetchall()
            return [dict(row) for row in rows], {}, mode
        raise DomainError("INVALID_SELECTION", "selection.mode 仅支持 explicit 或 filter", status_code=422)

    @staticmethod
    def _latest_manual_states(
        connection: sqlite3.Connection,
        task_id: str,
        source_row_ids: list[str],
    ) -> dict[str, tuple[str, str | None] | None]:
        result: dict[str, tuple[str, str | None] | None] = {source_row_id: None for source_row_id in source_row_ids}
        if not source_row_ids:
            return result
        for chunk in _chunks(source_row_ids):
            marks = ",".join("?" for _ in chunk)
            log_rows = connection.execute(
                f"""SELECT source_row_id,operation_type,selected_group_code FROM (
                        SELECT source_row_id,operation_type,selected_group_code,
                               ROW_NUMBER() OVER(PARTITION BY source_row_id ORDER BY operated_at DESC,rowid DESC) AS rn
                        FROM match_operation_logs WHERE task_id=? AND source_row_id IN ({marks})
                    ) WHERE rn=1""",
                (task_id, *chunk),
            ).fetchall()
            seen: set[str] = set()
            for row in log_rows:
                source_row_id = str(row["source_row_id"])
                seen.add(source_row_id)
                action = str(row["operation_type"])
                if action in _MATCH_ACTIONS:
                    result[source_row_id] = ("MATCH", str(row["selected_group_code"]) if row["selected_group_code"] is not None else None)
                elif action in _UNMATCH_ACTIONS:
                    result[source_row_id] = ("UNMATCHED", None)
                elif action in _CLEAR_ACTIONS:
                    result[source_row_id] = None
            missing = [source_row_id for source_row_id in chunk if source_row_id not in seen]
            if missing:
                missing_marks = ",".join("?" for _ in missing)
                review_rows = connection.execute(
                    f"""SELECT source_row_id,action,selected_group_code FROM (
                            SELECT source_row_id,action,selected_group_code,
                                   ROW_NUMBER() OVER(PARTITION BY source_row_id ORDER BY created_at DESC,rowid DESC) AS rn
                            FROM reviews WHERE task_id=? AND source_row_id IN ({missing_marks})
                        ) WHERE rn=1""",
                    (task_id, *missing),
                ).fetchall()
                for row in review_rows:
                    source_row_id = str(row["source_row_id"])
                    action = str(row["action"])
                    if action in {"CONFIRM_CANDIDATE", "MATCH", "REMATCH", "IMPORT_MATCH", "CONFIRM_TOP1"}:
                        result[source_row_id] = ("MATCH", str(row["selected_group_code"]) if row["selected_group_code"] is not None else None)
                    elif action in {"REJECT_ALL", "MARK_UNMATCHED"}:
                        result[source_row_id] = ("UNMATCHED", None)
                    elif action in _CLEAR_ACTIONS:
                        result[source_row_id] = None
        return result

    @staticmethod
    def _candidate_map(
        connection: sqlite3.Connection,
        task_id: str,
        source_row_ids: list[str],
    ) -> dict[str, dict[int, dict[str, object]]]:
        result: dict[str, dict[int, dict[str, object]]] = {}
        for chunk in _chunks(source_row_ids):
            marks = ",".join("?" for _ in chunk)
            rows = connection.execute(
                f"""SELECT source_row_id,rank,target_row_number,target_group_code
                    FROM match_candidates WHERE task_id=? AND source_row_id IN ({marks}) AND rank<=5""",
                (task_id, *chunk),
            ).fetchall()
            for row in rows:
                result.setdefault(str(row["source_row_id"]), {})[int(row["rank"])] = dict(row)
        return result

    @staticmethod
    def _target_for_group(candidates: dict[int, dict[str, object]], group_code: str | None) -> int | None:
        if group_code is None:
            return None
        for candidate in candidates.values():
            if str(candidate.get("target_group_code")) == str(group_code):
                value = candidate.get("target_row_number")
                return int(value) if value is not None else None
        return None

    @staticmethod
    def _baseline(row: dict[str, object]) -> tuple[str, str | None]:
        status = str(row.get("original_status") or "REVIEW")
        group = str(row["top1_group_code"]) if status == "MATCHED" and row.get("top1_group_code") is not None else None
        return status, group

    @staticmethod
    def _result_bucket() -> dict[str, object]:
        return {"success": 0, "skipped": 0, "conflicts": 0, "errors": 0, "details": []}

    @staticmethod
    def _detail(result: dict[str, object], kind: str, source_row_id: str, code: str, message: str, **extra: object) -> None:
        key = "conflicts" if kind == "conflict" else "errors" if kind == "error" else "skipped" if kind == "skipped" else "success"
        result[key] = int(result[key]) + 1
        if kind != "success" or extra:
            detail = {"source_row_id": source_row_id, "type": kind, "code": code, "message": message, **extra}
            cast_details = result["details"]
            assert isinstance(cast_details, list)
            cast_details.append(detail)

    def _apply_desired(
        self,
        connection: sqlite3.Connection,
        *,
        task_id: str,
        rows: list[dict[str, object]],
        expected_versions: dict[str, str | None],
        desired_by_id: dict[str, tuple[str, str | None, str]],
        operator: str,
        source: str,
        comment: str,
        batch_operation_id: str,
    ) -> dict[str, object]:
        result = self._result_bucket()
        ids = [str(row["source_row_id"]) for row in rows]
        row_by_id = {str(row["source_row_id"]): row for row in rows}
        candidates = self._candidate_map(connection, task_id, ids)
        manual_states = self._latest_manual_states(connection, task_id, ids)
        now = _now()
        updates: list[tuple[object, ...]] = []
        logs: list[tuple[object, ...]] = []
        reviews: list[tuple[object, ...]] = []
        audits: list[tuple[object, ...]] = []

        for source_row_id, desired in desired_by_id.items():
            row = row_by_id.get(source_row_id)
            if row is None:
                self._detail(result, "error", source_row_id, "MATCH_ITEM_NOT_FOUND", "匹配记录不存在")
                continue
            desired_kind, desired_group, operation_type = desired
            expected = expected_versions.get(source_row_id)
            current_version = str(row.get("updated_at") or "")
            active = manual_states.get(source_row_id)
            effective_desired = (desired_kind, desired_group if desired_kind == "MATCH" else None)
            if active == effective_desired:
                self._detail(result, "skipped", source_row_id, "IDEMPOTENT", "该人工结果已存在", version=current_version)
                continue
            if expected is not None and expected != current_version:
                self._detail(
                    result,
                    "conflict",
                    source_row_id,
                    "VERSION_CONFLICT",
                    "记录已被其他用户修改，请刷新后重试",
                    expected_version=expected,
                    current_version=current_version,
                    current_status=row.get("current_status"),
                    current_group_code=row.get("final_group_code"),
                )
                continue
            if active is not None and operation_type not in _CLEAR_ACTIONS:
                rematch_allowed = operation_type == "REMATCH" and active[0] == "MATCH" and expected is not None
                if not rematch_allowed:
                    self._detail(
                        result,
                        "conflict",
                        source_row_id,
                        "MANUAL_RESULT_CONFLICT",
                        "该记录已存在不同人工结果，未覆盖",
                        existing={"decision": active[0], "group_code": active[1]},
                        incoming={"decision": effective_desired[0], "group_code": effective_desired[1]},
                    )
                    continue

            before_status = str(row["current_status"])
            previous_group = str(row["final_group_code"]) if row.get("final_group_code") is not None else None
            row_candidates = candidates.get(source_row_id, {})
            previous_target = self._target_for_group(row_candidates, previous_group)
            if desired_kind == "MATCH":
                selected_candidate = next(
                    (candidate for candidate in row_candidates.values() if str(candidate.get("target_group_code")) == str(desired_group)),
                    None,
                )
                if selected_candidate is None:
                    self._detail(result, "error", source_row_id, "CANDIDATE_NOT_FOUND", "所选 Top5 候选不存在")
                    continue
                after_status = "CONFIRMED"
                after_group = str(desired_group)
                target_row = selected_candidate.get("target_row_number")
                target_row_number = int(target_row) if target_row is not None else None
            elif desired_kind == "UNMATCHED":
                after_status, after_group, target_row_number = "UNMATCHED", None, None
            elif desired_kind == "BASELINE":
                after_status, after_group = self._baseline(row)
                target_row_number = self._target_for_group(row_candidates, after_group)
            else:
                self._detail(result, "error", source_row_id, "INVALID_DECISION", "人工处理动作不正确")
                continue

            if operation_type == "CANCEL_MATCH" and (active is None or active[0] != "MATCH"):
                self._detail(result, "skipped", source_row_id, "NO_MANUAL_MATCH", "当前记录没有可取消的人工匹配")
                continue
            if operation_type == "CANCEL_UNMATCHED" and (active is None or active[0] != "UNMATCHED"):
                self._detail(result, "skipped", source_row_id, "NO_MANUAL_UNMATCHED", "当前记录没有可取消的人工未匹配")
                continue
            if operation_type == "RESTORE_ALGORITHM" and active is None and before_status == after_status and previous_group == after_group:
                self._detail(result, "skipped", source_row_id, "ALREADY_BASELINE", "记录已经是算法原始结果")
                continue

            operation_id = uuid.uuid4().hex
            updates.append((after_status, after_group, now, task_id, source_row_id, current_version))
            logs.append(
                (
                    operation_id,
                    task_id,
                    source_row_id,
                    row.get("source_row_number"),
                    operator or "system",
                    now,
                    operation_type,
                    before_status,
                    after_status,
                    previous_group,
                    after_group,
                    previous_target,
                    target_row_number,
                    source,
                    comment,
                )
            )
            reviews.append(
                (
                    uuid.uuid4().hex,
                    task_id,
                    source_row_id,
                    before_status,
                    after_group,
                    operation_type,
                    operator or "system",
                    comment,
                    now,
                )
            )
            audits.append(
                (
                    uuid.uuid4().hex,
                    "match_item",
                    f"{task_id}:{source_row_id}",
                    operation_type,
                    _json(
                        {
                            "operator": operator or "system",
                            "source": source,
                            "batch_operation_id": batch_operation_id,
                            "operation_id": operation_id,
                            "before_status": before_status,
                            "after_status": after_status,
                            "previous_group_code": previous_group,
                            "selected_group_code": after_group,
                            "previous_target_row_number": previous_target,
                            "target_row_number": target_row_number,
                            "comment": comment,
                        }
                    ),
                    now,
                )
            )
            self._detail(result, "success", source_row_id, "OK", "操作成功")

        if updates:
            connection.executemany(
                """UPDATE match_items SET current_status=?,final_group_code=?,updated_at=?
                   WHERE task_id=? AND source_row_id=? AND updated_at=?""",
                updates,
            )
            # BEGIN IMMEDIATE protects the selected versions for the whole transaction.
            connection.executemany(
                """INSERT INTO match_operation_logs(
                    operation_id,task_id,source_row_id,source_row_number,operator,operated_at,
                    operation_type,before_status,after_status,previous_group_code,selected_group_code,
                    previous_target_row_number,target_row_number,source,comment
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                logs,
            )
            connection.executemany(
                "INSERT INTO reviews(review_id,task_id,source_row_id,original_status,selected_group_code,action,operator,comment,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                reviews,
            )
            connection.executemany("INSERT INTO audit_events VALUES(?,?,?,?,?,?)", audits)
        result["batch_operation_id"] = batch_operation_id
        result["selected_count"] = len(desired_by_id)
        return result

    def bulk_action(
        self,
        task_id: str,
        *,
        action: str,
        selection: dict[str, object],
        operator: str,
        source: str = "WEB",
        comment: str = "",
        operation_id: str | None = None,
    ) -> dict[str, object]:
        action = action.upper()
        if action not in _ALLOWED_ACTIONS:
            raise DomainError("INVALID_BULK_ACTION", "批量操作类型不正确", status_code=422)
        batch_operation_id = operation_id or uuid.uuid4().hex
        with self.repo.connect() as connection:
            self._ensure_task(connection, task_id)
            connection.execute("BEGIN IMMEDIATE")
            rows, expected, mode = self._selection_rows(connection, task_id, selection)
            ids = [str(row["source_row_id"]) for row in rows]
            candidates = self._candidate_map(connection, task_id, ids) if action == "CONFIRM_TOP1" else {}
            desired: dict[str, tuple[str, str | None, str]] = {}
            missing_top1: list[str] = []
            for row in rows:
                source_row_id = str(row["source_row_id"])
                if action == "CONFIRM_TOP1":
                    candidate = candidates.get(source_row_id, {}).get(1)
                    if candidate is None:
                        missing_top1.append(source_row_id)
                        continue
                    desired[source_row_id] = ("MATCH", str(candidate["target_group_code"]), "CONFIRM_TOP1")
                elif action == "MARK_UNMATCHED":
                    desired[source_row_id] = ("UNMATCHED", None, "MARK_UNMATCHED")
                elif action == "CANCEL_MATCH":
                    desired[source_row_id] = ("BASELINE", None, "CANCEL_MATCH")
                elif action == "CANCEL_UNMATCHED":
                    desired[source_row_id] = ("BASELINE", None, "CANCEL_UNMATCHED")
                else:
                    desired[source_row_id] = ("BASELINE", None, "RESTORE_ALGORITHM")
            result = self._apply_desired(
                connection,
                task_id=task_id,
                rows=rows,
                expected_versions=expected,
                desired_by_id=desired,
                operator=operator,
                source=source,
                comment=comment,
                batch_operation_id=batch_operation_id,
            )
            for source_row_id in missing_top1:
                self._detail(result, "error", source_row_id, "NO_CANDIDATE", "没有可确认的 Top1 候选")
            result["selection_mode"] = mode
            result["matched_filter_count"] = len(rows) if mode == "filter" else None
            return result

    def submit_decisions(
        self,
        task_id: str,
        decisions: list[dict[str, object]],
        *,
        operator: str,
        source: str = "WEB",
        comment: str = "",
        operation_id: str | None = None,
        require_expected_version: bool = True,
    ) -> dict[str, object]:
        if not decisions:
            raise DomainError("INVALID_DECISIONS", "至少需要提交一条人工选择", status_code=422)
        if len(decisions) > 100_000:
            raise DomainError("TOO_MANY_DECISIONS", "单次人工选择超过 100000 条", status_code=413)
        batch_operation_id = operation_id or uuid.uuid4().hex
        seen: dict[str, tuple[str, object]] = {}
        conflicted_ids: set[str] = set()
        expected: dict[str, str | None] = {}
        normalized: dict[str, tuple[str, str | None, str]] = {}
        requested_rank: dict[str, int | None] = {}
        requested_group: dict[str, str | None] = {}
        conflicts_before_db: list[dict[str, object]] = []
        for decision in decisions:
            source_row_id = str(decision.get("source_row_id") or "").strip()
            if not source_row_id:
                raise DomainError("INVALID_DECISIONS", "人工选择缺少 source_row_id", status_code=422)
            expected_value = decision.get("expected_version")
            if require_expected_version and expected_value is None:
                raise DomainError("EXPECTED_VERSION_REQUIRED", "在线批量提交必须携带 expected_version", status_code=422)
            expected[source_row_id] = str(expected_value) if expected_value is not None else None
            group = decision.get("target_group_code")
            rank_value = decision.get("target_rank")
            unmatched = bool(decision.get("unmatched", False))
            signature: tuple[str, object] | tuple[str, object, str]
            if unmatched:
                signature = ("UNMATCHED", "")
                requested_rank[source_row_id] = None
                requested_group[source_row_id] = None
            elif group is not None:
                signature = ("GROUP", str(group))
                requested_rank[source_row_id] = None
                requested_group[source_row_id] = str(group)
            elif rank_value is not None:
                rank = int(rank_value)
                if rank < 1 or rank > 5:
                    raise DomainError("INVALID_CANDIDATE_RANK", "候选序号仅支持 1～5", status_code=422)
                signature = ("RANK", rank)
                requested_rank[source_row_id] = rank
                requested_group[source_row_id] = None
            else:
                raise DomainError("INVALID_DECISIONS", "人工选择必须指定 target_group_code、target_rank 或 unmatched", status_code=422)
            operation = str(decision.get("operation") or "MATCH").upper()
            if operation not in {"MATCH", "REMATCH"}:
                raise DomainError("INVALID_DECISION_OPERATION", "人工选择 operation 仅支持 MATCH 或 REMATCH", status_code=422)
            signature = (signature[0], signature[1], operation)
            previous = seen.get(source_row_id)
            if source_row_id in conflicted_ids:
                continue
            if previous is not None and previous != signature:
                conflicted_ids.add(source_row_id)
                seen.pop(source_row_id, None)
                conflicts_before_db.append(
                    {
                        "source_row_id": source_row_id,
                        "type": "conflict",
                        "code": "DUPLICATE_DECISION_CONFLICT",
                        "message": "同一批次对同一源记录给出了不同结果",
                    }
                )
                continue
            seen[source_row_id] = signature

        ids = list(seen)
        with self.repo.connect() as connection:
            self._ensure_task(connection, task_id)
            connection.execute("BEGIN IMMEDIATE")
            rows, _, _ = self._selection_rows(
                connection,
                task_id,
                {"mode": "explicit", "items": [{"source_row_id": value, "expected_version": expected.get(value)} for value in ids]},
            )
            candidates = self._candidate_map(connection, task_id, ids)
            for source_row_id, signature in seen.items():
                requested_operation = str(signature[2])
                if signature[0] == "UNMATCHED":
                    normalized[source_row_id] = ("UNMATCHED", None, "MARK_UNMATCHED")
                    continue
                row_candidates = candidates.get(source_row_id, {})
                selected: dict[str, object] | None = None
                rank = requested_rank.get(source_row_id)
                group = requested_group.get(source_row_id)
                if rank is not None:
                    selected = row_candidates.get(rank)
                elif group is not None:
                    selected = next((item for item in row_candidates.values() if str(item.get("target_group_code")) == group), None)
                if selected is None:
                    normalized[source_row_id] = ("MATCH", f"__MISSING__:{group or rank}", "IMPORT_MATCH" if source == "EXCEL" else requested_operation)
                else:
                    normalized[source_row_id] = (
                        "MATCH",
                        str(selected["target_group_code"]),
                        "IMPORT_MATCH" if source == "EXCEL" else requested_operation,
                    )
            result = self._apply_desired(
                connection,
                task_id=task_id,
                rows=rows,
                expected_versions=expected,
                desired_by_id=normalized,
                operator=operator,
                source=source,
                comment=comment,
                batch_operation_id=batch_operation_id,
            )
        if conflicts_before_db:
            result["conflicts"] = int(result["conflicts"]) + len(conflicts_before_db)
            details = result["details"]
            assert isinstance(details, list)
            details.extend(conflicts_before_db)
        return result

    @staticmethod
    def _parse_excel_selection(value: object) -> tuple[str, int | None] | None:
        if value is None or str(value).strip() == "":
            return None
        text = str(value).strip().replace(" ", "")
        if text in {"均不匹配", "不匹配", "无匹配", "NONE", "none"}:
            return ("UNMATCHED", None)
        if text.isdigit() and 1 <= int(text) <= 5:
            return ("MATCH", int(text))
        if text.startswith("候选") and text[2:].isdigit() and 1 <= int(text[2:]) <= 5:
            return ("MATCH", int(text[2:]))
        raise DomainError("INVALID_MANUAL_SELECTION", f"人工选择值不正确: {value}", status_code=422)

    def import_workbook(
        self,
        task_id: str,
        stream: BinaryIO,
        *,
        operator: str,
        filename: str = "",
    ) -> dict[str, object]:
        try:
            workbook = load_workbook(stream, data_only=True, read_only=True)
        except Exception as exc:
            raise DomainError("MANUAL_EXCEL_PARSE_FAILED", "人工匹配 Excel 无法解析", status_code=400) from exc
        if "人工匹配" not in workbook.sheetnames:
            raise DomainError("MANUAL_EXCEL_INVALID", "未找到“人工匹配”工作表", status_code=422)
        sheet = workbook["人工匹配"]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = {str(value).strip(): index for index, value in enumerate(header_values) if value is not None}
        required = {"人工选择", "__task_id", "__source_row_id"}
        if not required.issubset(headers):
            raise DomainError("MANUAL_EXCEL_INVALID", "人工匹配 Excel 缺少内部识别列，请使用系统下载的模板", status_code=422)

        decisions: list[dict[str, object]] = []
        pre_details: list[dict[str, object]] = []
        blank_rows = 0
        seen: dict[str, tuple[str, int | None]] = {}
        conflicted_ids: set[str] = set()
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            selection_value = values[headers["人工选择"]] if headers["人工选择"] < len(values) else None
            if selection_value is None or str(selection_value).strip() == "":
                blank_rows += 1
                continue
            source_row_id = str(values[headers["__source_row_id"]] if headers["__source_row_id"] < len(values) else "").strip()
            try:
                parsed = self._parse_excel_selection(selection_value)
                assert parsed is not None
                embedded_task_id = str(values[headers["__task_id"]] if headers["__task_id"] < len(values) else "").strip()
                if embedded_task_id != task_id:
                    raise DomainError("MANUAL_EXCEL_TASK_MISMATCH", "Excel 中的任务与当前任务不一致", status_code=409)
                if not source_row_id:
                    raise DomainError("MANUAL_EXCEL_INVALID", "Excel 行缺少源记录标识", status_code=422)
                previous = seen.get(source_row_id)
                if source_row_id in conflicted_ids:
                    continue
                if previous is not None:
                    if previous == parsed:
                        pre_details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "skipped", "code": "DUPLICATE_IDEMPOTENT", "message": "同一文件中的重复相同结果已忽略"})
                    else:
                        conflicted_ids.add(source_row_id)
                        seen.pop(source_row_id, None)
                        pre_details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "conflict", "code": "DUPLICATE_ROW_CONFLICT", "message": "同一文件对同一源记录给出了不同结果"})
                    continue
                seen[source_row_id] = parsed
            except DomainError as exc:
                pre_details.append({"excel_row": excel_row, "source_row_id": source_row_id, "type": "error", "code": exc.code, "message": exc.message})

        for source_row_id, parsed in seen.items():
            if parsed[0] == "UNMATCHED":
                decisions.append({"source_row_id": source_row_id, "unmatched": True})
            else:
                decisions.append({"source_row_id": source_row_id, "target_rank": int(parsed[1] or 0)})

        if decisions:
            result = self.submit_decisions(
                task_id,
                decisions,
                operator=operator,
                source="EXCEL",
                comment=f"导入文件: {filename}" if filename else "Excel 上传",
                require_expected_version=False,
            )
        else:
            result = self._result_bucket()
            result["batch_operation_id"] = uuid.uuid4().hex
            result["selected_count"] = 0
        result["skipped"] = int(result["skipped"]) + blank_rows + sum(1 for detail in pre_details if detail["type"] == "skipped")
        result["conflicts"] = int(result["conflicts"]) + sum(1 for detail in pre_details if detail["type"] == "conflict")
        result["errors"] = int(result["errors"]) + sum(1 for detail in pre_details if detail["type"] == "error")
        details = result["details"]
        assert isinstance(details, list)
        details.extend(pre_details)
        # Preserve the legacy response keys used by the existing STEP3 UI.
        return {
            "task_id": task_id,
            "success": int(result["success"]),
            "skipped": int(result["skipped"]),
            "conflicts": int(result["conflicts"]),
            "errors": int(result["errors"]),
            "success_count": int(result["success"]),
            "skipped_count": int(result["skipped"]),
            "conflict_count": int(result["conflicts"]),
            "error_count": int(result["errors"]),
            "details": details,
            "operation_id": result["batch_operation_id"],
        }
