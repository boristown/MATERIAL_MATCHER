from __future__ import annotations

import math
import re
from collections import defaultdict
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook


_TOKEN_RE = re.compile(r"[A-Za-z0-9]+|[\u4e00-\u9fff]+")
_NUM_RE = re.compile(r"[-+]?\d+(?:\.\d+)?")


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    if text in {"", "88", "-", "/", "<null>", "null", "n/a"}:
        return ""
    text = text.replace("×", "x").replace("＊", "x").replace("Φ", "φ").replace("Ø", "φ")
    text = re.sub(r"[\s_—–]+", " ", text)
    text = re.sub(r"\s*([/,:;，；：()（）\[\]【】])\s*", r"\1", text)
    return text.strip()


def tokenize(value: Any) -> set[str]:
    text = normalize_text(value)
    if not text:
        return set()
    result: set[str] = set()
    for token in _TOKEN_RE.findall(text):
        if len(token) >= 2 or token.isdigit():
            result.add(token)
        if re.search(r"[\u4e00-\u9fff]", token) and len(token) > 2:
            result.update(token[i : i + 2] for i in range(len(token) - 1))
    return result


def _exact(a: Any, b: Any) -> float:
    na, nb = normalize_text(a), normalize_text(b)
    if not na or not nb:
        return 0.0
    return 1.0 if na == nb else 0.0


def _fuzzy(a: Any, b: Any) -> float:
    na, nb = normalize_text(a), normalize_text(b)
    if not na or not nb:
        return 0.0
    if na == nb:
        return 1.0
    return SequenceMatcher(None, na, nb).ratio()


def _token_score(a: Any, b: Any) -> float:
    aa, bb = tokenize(a), tokenize(b)
    if not aa or not bb:
        return 0.0
    return len(aa & bb) / len(aa | bb)


def _numeric_text(a: Any, b: Any) -> float:
    na, nb = normalize_text(a), normalize_text(b)
    if not na or not nb:
        return 0.0
    nums_a = [float(v) for v in _NUM_RE.findall(na)]
    nums_b = [float(v) for v in _NUM_RE.findall(nb)]
    text_score = _fuzzy(na, nb)
    if not nums_a or not nums_b:
        return text_score
    common = 0
    used: set[int] = set()
    for value in nums_a:
        best_idx = None
        best_delta = math.inf
        for idx, other in enumerate(nums_b):
            if idx in used:
                continue
            delta = abs(value - other) / max(abs(value), abs(other), 1.0)
            if delta < best_delta:
                best_idx, best_delta = idx, delta
        if best_idx is not None and best_delta <= 0.02:
            used.add(best_idx)
            common += 1
    numeric = common / max(len(nums_a), len(nums_b))
    return min(1.0, numeric * 0.75 + text_score * 0.25)


def score_value(a: Any, b: Any, method: str) -> float:
    method = (method or "hybrid").lower()
    if method in {"exact", "normalized_exact"}:
        return _exact(a, b)
    if method in {"numeric", "numeric_text", "normalized_standard"}:
        return _numeric_text(a, b)
    if method in {"token", "token_similarity"}:
        return _token_score(a, b)
    if method == "contains":
        na, nb = normalize_text(a), normalize_text(b)
        return 1.0 if na and nb and (na in nb or nb in na) else 0.0
    if method in {"semantic", "hybrid"}:
        # MVP fallback before the vector/BBQ backend is attached: combine lexical
        # evidence without changing the public matcher contract.
        return max(_fuzzy(a, b), _token_score(a, b))
    return _fuzzy(a, b)


@dataclass
class MappingRule:
    source_header: str
    target_header: str
    weight: float = 1.0
    method: str = "hybrid"
    name: str = ""
    critical: bool = False


@dataclass
class MatchConfig:
    source_sheet: str | None
    source_header_row: int
    target_sheet: str | None
    target_header_row: int
    source_id_column: str
    group_code_column: str
    mappings: list[MappingRule]
    threshold: float = 0.85
    review_threshold: float | None = None
    top_n: int = 5
    candidate_limit: int = 400

    @classmethod
    def from_dict(cls, raw: dict[str, Any]) -> "MatchConfig":
        mappings = [MappingRule(**item) for item in raw.get("mappings", [])]
        if not mappings:
            raise ValueError("至少需要一个字段映射")
        return cls(
            source_sheet=raw.get("source_sheet"),
            source_header_row=int(raw.get("source_header_row", 1)),
            target_sheet=raw.get("target_sheet"),
            target_header_row=int(raw.get("target_header_row", 1)),
            source_id_column=str(raw.get("source_id_column", "")),
            group_code_column=str(raw.get("group_code_column", "")),
            mappings=mappings,
            threshold=float(raw.get("threshold", 0.85)),
            review_threshold=float(raw["review_threshold"]) if raw.get("review_threshold") is not None else None,
            top_n=max(1, min(50, int(raw.get("top_n", 5)))),
            candidate_limit=max(20, min(5000, int(raw.get("candidate_limit", 400)))),
        )


def _worksheet(path: Path, sheet: str | None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet and sheet in workbook.sheetnames:
        return workbook, workbook[sheet]
    return workbook, workbook[workbook.sheetnames[0]]


def read_excel_rows(path: Path, sheet: str | None, header_row: int, max_rows: int | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    workbook, ws = _worksheet(path, sheet)
    headers = ["" if cell.value is None else str(cell.value).strip() for cell in ws[header_row]]
    rows: list[dict[str, Any]] = []
    for idx, cells in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=1):
        row = {headers[col]: value for col, value in enumerate(cells) if col < len(headers) and headers[col]}
        if not any(value is not None and str(value).strip() for value in row.values()):
            continue
        rows.append(row)
        if max_rows is not None and idx >= max_rows:
            break
    workbook.close()
    return headers, rows


def _combined(row: dict[str, Any], header: str) -> Any:
    # Wizard currently produces one-to-one physical mappings. Composite fields
    # can be represented by "A + B" without introducing customer-specific code.
    if " + " not in header:
        return row.get(header)
    values = [normalize_text(row.get(part.strip())) for part in header.split(" + ")]
    return " ".join(value for value in values if value)


def _row_score(source: dict[str, Any], target: dict[str, Any], rules: list[MappingRule]) -> tuple[float, list[dict[str, Any]]]:
    weighted = 0.0
    total_weight = 0.0
    evidence: list[dict[str, Any]] = []
    critical_conflict = False
    for rule in rules:
        source_value = _combined(source, rule.source_header)
        target_value = _combined(target, rule.target_header)
        if not normalize_text(source_value) or not normalize_text(target_value):
            continue
        score = score_value(source_value, target_value, rule.method)
        weight = max(0.0, float(rule.weight))
        weighted += score * weight
        total_weight += weight
        if rule.critical and score < 0.25:
            critical_conflict = True
        evidence.append(
            {
                "name": rule.name or f"{rule.source_header}↔{rule.target_header}",
                "source": "" if source_value is None else str(source_value),
                "target": "" if target_value is None else str(target_value),
                "score": round(score, 4),
                "weight": weight,
            }
        )
    score = weighted / total_weight if total_weight else 0.0
    if critical_conflict:
        score = min(score, 0.79)
    return score, evidence


def _build_candidate_index(target_rows: list[dict[str, Any]], rules: list[MappingRule]) -> dict[str, set[int]]:
    index: dict[str, set[int]] = defaultdict(set)
    for idx, row in enumerate(target_rows):
        for rule in rules:
            for token in tokenize(_combined(row, rule.target_header)):
                index[token].add(idx)
    return index


def _candidate_ids(source: dict[str, Any], rules: list[MappingRule], index: dict[str, set[int]], target_count: int, limit: int) -> list[int]:
    counts: dict[int, int] = defaultdict(int)
    for rule in rules:
        for token in tokenize(_combined(source, rule.source_header)):
            for idx in index.get(token, ()):
                counts[idx] += 1
    if not counts:
        return list(range(min(target_count, limit)))
    return [idx for idx, _ in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[:limit]]


def run_matching(
    source_path: Path,
    target_path: Path,
    raw_config: dict[str, Any],
    *,
    max_source_rows: int | None = None,
    max_target_rows: int | None = None,
) -> dict[str, Any]:
    config = MatchConfig.from_dict(raw_config)
    source_headers, source_rows = read_excel_rows(source_path, config.source_sheet, config.source_header_row, max_source_rows)
    target_headers, target_rows = read_excel_rows(target_path, config.target_sheet, config.target_header_row, max_target_rows)
    if config.group_code_column not in target_headers:
        raise ValueError(f"集团码列不存在：{config.group_code_column}")
    for rule in config.mappings:
        for header, available, label in (
            (rule.source_header, source_headers, "客户物料"),
            (rule.target_header, target_headers, "集团码"),
        ):
            for part in [p.strip() for p in header.split(" + ")]:
                if part not in available:
                    raise ValueError(f"{label}字段不存在：{part}")

    candidate_index = _build_candidate_index(target_rows, config.mappings)
    output_rows: list[dict[str, Any]] = []
    matched = review = unmatched = 0

    for source_idx, source in enumerate(source_rows, start=1):
        candidate_ids = _candidate_ids(source, config.mappings, candidate_index, len(target_rows), config.candidate_limit)
        ranked: list[tuple[float, int, list[dict[str, Any]]]] = []
        for target_idx in candidate_ids:
            score, evidence = _row_score(source, target_rows[target_idx], config.mappings)
            ranked.append((score, target_idx, evidence))
        ranked.sort(key=lambda item: item[0], reverse=True)
        top = ranked[: config.top_n]
        best_score = top[0][0] if top else 0.0
        best_target = target_rows[top[0][1]] if top else {}
        if best_score > config.threshold:
            status = "MATCHED"
            matched += 1
        elif config.review_threshold is not None and best_score > config.review_threshold:
            status = "REVIEW"
            review += 1
        else:
            status = "UNMATCHED"
            unmatched += 1
        group_code = best_target.get(config.group_code_column, "") if status == "MATCHED" else ""
        candidates = []
        for rank, (score, target_idx, evidence) in enumerate(top, start=1):
            target = target_rows[target_idx]
            candidates.append(
                {
                    "rank": rank,
                    "score": round(score, 6),
                    "group_code": target.get(config.group_code_column, ""),
                    "target": target,
                    "evidence": evidence,
                }
            )
        output_rows.append(
            {
                "source_index": source_idx,
                "source": source,
                "source_id": source.get(config.source_id_column, "") if config.source_id_column else "",
                "status": status,
                "matched_group_code": group_code,
                "score": round(best_score, 6),
                "candidates": candidates,
            }
        )

    total = len(output_rows)
    return {
        "summary": {
            "source_rows": total,
            "target_rows": len(target_rows),
            "matched": matched,
            "review": review,
            "unmatched": unmatched,
            "matched_rate": round(matched / total, 6) if total else 0.0,
            "threshold": config.threshold,
            "review_threshold": config.review_threshold,
        },
        "source_headers": source_headers,
        "rows": output_rows,
    }


def write_result_xlsx(result: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "匹配结果"
    source_headers = list(result.get("source_headers", []))
    result_headers = source_headers + ["匹配状态", "集团码", "相似度"]
    ws.append(result_headers)
    for item in result.get("rows", []):
        source = item["source"]
        ws.append([source.get(header, "") for header in source_headers] + [item["status"], item["matched_group_code"], item["score"]])

    topn = wb.create_sheet("TopN")
    topn.append(["源行号", "源物料", "排名", "候选集团码", "相似度", "字段证据"])
    for item in result.get("rows", []):
        for candidate in item.get("candidates", []):
            evidence_text = "; ".join(
                f"{e['name']}={e['score']:.3f}" for e in candidate.get("evidence", [])
            )
            topn.append(
                [
                    item["source_index"],
                    item.get("source_id", ""),
                    candidate["rank"],
                    candidate.get("group_code", ""),
                    candidate["score"],
                    evidence_text,
                ]
            )
    wb.save(path)
